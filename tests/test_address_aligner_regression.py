from __future__ import annotations

from pathlib import Path
from types import SimpleNamespace
from zipfile import ZIP_DEFLATED, ZipFile
import contextlib
import importlib.util
import io
import json
import os
import sys
import tempfile
import unittest

from openpyxl import load_workbook
import pandas as pd
import yaml


ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from system_core.Ultimate_GT_Aligner import (  # noqa: E402
    _detect_companion_columns,
    _alignment_scope_match,
    _alignment_scope_pair_blocked,
    _alignment_scope_pair_score,
    _relaxed_house_check,
    parse_address,
    process_file,
)
from system_core.address_engine import AddressNormalizer  # noqa: E402
from system_core.address_engine.address_diagnostics import record_issue_codes  # noqa: E402
from system_core.address_engine.address_slots import build_address_parts  # noqa: E402
from system_core.address_engine.reference_builder import (  # noqa: E402
    DEFAULT_SLOT_ORDER,
    DEFAULT_ENABLED_SLOTS,
    DEFAULT_SELECTED_SLOTS,
    ReferenceOptions,
    build_reference_workbook,
)
from system_core.address_engine.reference_processor import ReferenceProcessingOptions, process_reference_workbook  # noqa: E402
from system_core.address_engine.legacy_office_converter import convert_legacy_office_folder  # noqa: E402
from system_core.address_engine.oktmo_lookup import oktmo_scope_key_profile  # noqa: E402
from system_core.address_engine.oktmo_user_keys import load_current_keys, load_pins, write_current_keys  # noqa: E402
from system_core.address_engine.source_table_assembly import (  # noqa: E402
    AddressAssemblyOptions,
    AddressAppendOptions,
    _CandidateRow,
    _GroupedRows,
    _extract_clean_address_fragments,
    _final_rows_from_groups,
    _normalize_house_slot,
    _normalize_street_slot,
    _oktmo_scope_candidate_match,
    append_collected_address_workbook,
    assemble_address_workbook,
    benchmark_address_collection,
)
from system_core.core.jobs import execute_operation  # noqa: E402
from system_core.core.manifest import Operation  # noqa: E402
from system_core.core.paths import ProjectPaths, ensure_project_dirs  # noqa: E402
from system_core.services.address_aligner_service import (  # noqa: E402
    ADDRESS_RUNTIME_PARAMETER_IDS,
    AddressRuntimeContext,
    OKTMO_SCOPE_PARAMETER_IDS,
    _aligner_runtime_config,
    _address_runtime_summary,
    _artifact_summary,
    oktmo_clear_pin_file,
    oktmo_data_status,
    oktmo_add_scope_keys,
    oktmo_add_candidate_keys,
    oktmo_export_pin_file,
    oktmo_find_settlement_candidates,
    oktmo_import_pin_file,
    oktmo_municipality_scope_options,
    oktmo_pin_scope_value,
    oktmo_region_scope_options,
    oktmo_reset_project_keys,
    oktmo_scope_pin_status,
)
from system_core.address_engine.source_cleanup import delete_office_temp_files  # noqa: E402
from system_core.address_engine.source_intake import SourceCandidate, extract_reference_candidates, extract_source_candidates  # noqa: E402


def write_minimal_odt(path: Path, text: str) -> None:
    content = f"""<?xml version="1.0" encoding="UTF-8"?>
<office:document-content
  xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0"
  xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0"
  xmlns:table="urn:oasis:names:tc:opendocument:xmlns:table:1.0">
  <office:body><office:text><text:p>{text}</text:p></office:text></office:body>
</office:document-content>"""
    with ZipFile(path, "w", ZIP_DEFLATED) as archive:
        archive.writestr("content.xml", content)


def write_minimal_docx(path: Path, text: str) -> None:
    escaped = (
        str(text)
        .replace("&", "&amp;")
        .replace("<", "&lt;")
        .replace(">", "&gt;")
    )
    document = f"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
  <w:body>
    <w:p><w:r><w:t>{escaped}</w:t></w:r></w:p>
  </w:body>
</w:document>"""
    with ZipFile(path, "w", ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", document)


def write_searchable_pdf(path: Path, text: str) -> None:
    try:
        import pymupdf as fitz
    except ImportError:
        import fitz

    document = fitz.open()
    page = document.new_page()
    font_file = Path(os.environ.get("WINDIR", "C:\\Windows")) / "Fonts" / "arial.ttf"
    kwargs = {"fontsize": 11}
    if font_file.exists():
        kwargs.update({"fontfile": str(font_file), "fontname": "Arial"})
    page.insert_text((72, 72), text, **kwargs)
    document.save(path)
    document.close()


def write_oktmo_fixture(data_root: Path) -> None:
    rosstat = data_root / "rosstat"
    rosstat.mkdir(parents=True, exist_ok=True)
    (rosstat / "data-20260601T0000-structure.csv").write_text(
        "\n".join(
            [
                "71;000;000;000;;1;Муниципальные образования Тюменской области;",
                "71;001;001;000;;1;Тюменский муниципальный район;",
                "71;001;001;001;;2;деревня Патрушева;",
                "71;001;001;002;;2;село Мирный;",
                "71;001;001;003;;2;поселок Центральный;",
                "71;001;001;004;;2;село Ключи;",
                "71;800;000;000;;1;Ханты-Мансийский автономный округ - Югра;",
                "71;876;000;000;;1;городской округ Сургут;",
                "71;876;000;001;;2;г Сургут;",
                "71;900;000;000;;1;Ямало-Ненецкий автономный округ;",
            ]
        ),
        encoding="utf-8",
    )


def write_region_fixture(data_root: Path) -> None:
    (data_root / "regions_reference.yaml").write_text(
        "\n".join(
            [
                "regions:",
                "- region_id: tyumenskaya_oblast",
                "  display_name: Тюменская область",
                "  short_name: Тюменская область",
                "  subject_type: область",
                "  aliases:",
                "    - Тюменской области",
                "- region_id: altayskiy_krai",
                "  display_name: Алтайский край",
                "  short_name: Алтайский край",
                "  subject_type: край",
                "  aliases:",
                "    - Алтайского края",
                "- region_id: khanty_mansiyskiy_ao_yugra",
                "  display_name: Ханты-Мансийский автономный округ — Югра",
                "  short_name: ХМАО — Югра",
                "  subject_type: автономный округ",
                "  aliases:",
                "    - Ханты-Мансийский автономный округ - Югра",
                "    - Ханты-Мансийского автономного округа - Югры",
                "    - ХМАО",
                "    - Югра",
            ]
        ),
        encoding="utf-8",
    )


def write_altay_klyuchi_fixture(data_root: Path) -> None:
    rosstat = data_root / "rosstat"
    rosstat.mkdir(parents=True, exist_ok=True)
    (rosstat / "data-20260601T0000-structure.csv").write_text(
        "\n".join(
            [
                "01;000;000;000;;1;Муниципальные образования Алтайского края;",
                "01;649;420;000;;1;Ключевский сельсовет;п Ключи",
                "01;649;420;101;;2;п Ключи;",
            ]
        ),
        encoding="utf-8",
    )


def write_tyumen_scope_conflict_fixture(data_root: Path) -> None:
    rosstat = data_root / "rosstat"
    rosstat.mkdir(parents=True, exist_ok=True)
    (rosstat / "data-20260601T0000-structure.csv").write_text(
        "\n".join(
            [
                '"71";"530";"000";"000";"3";"1";"Казанский муниципальный округ";"с Казанское";;"906";"2";25.12.2025;01.02.2026',
                '"71";"530";"000";"000";"3";"2";"Населенные пункты, входящие в состав Казанского муниципального округа";;;"906";"2";25.12.2025;01.02.2026',
                '"71";"530";"000";"161";"7";"2";"д Дальнетравное";;;"843";"3";30.07.2025;01.01.2026',
                '"71";"701";"000";"000";"2";"1";"город Тюмень";"г Тюмень";;"000";"0";14.06.2013;01.01.2014',
                '"71";"701";"000";"000";"2";"2";"Населенные пункты, входящие в состав городского округа город Тюмень";;;"000";"0";14.06.2013;01.01.2014',
                '"71";"701";"000";"001";"3";"2";"г Тюмень";;;"000";"0";14.06.2013;01.01.2014',
            ]
        ),
        encoding="utf-8",
    )


class AddressParsingTests(unittest.TestCase):
    def test_house_number_after_street_without_comma_or_space(self) -> None:
        self.assertEqual(parse_address("Ленина10")[1:5], ("ленина", frozenset(), "10", ""))
        self.assertEqual(parse_address("ул Ленина10А")[1:5], ("ленина", frozenset(), "10", "а"))
        self.assertEqual(parse_address("ул Ленина 10к2")[1:5], ("ленина", frozenset(), "10", "к2"))
        self.assertEqual(parse_address("ул Ленина, 10 к. 2")[1:5], ("ленина", frozenset(), "10", "к2"))

    def test_settlement_type_variants_from_official_registry(self) -> None:
        self.assertEqual(parse_address("д Патрушева ул Ленина 10"), ("патрушева", "ленина", frozenset(), "10", ""))
        self.assertEqual(
            parse_address("деревня Патрушева, улица Ленина, д. 10"),
            ("патрушева", "ленина", frozenset(), "10", ""),
        )
        self.assertEqual(parse_address("с Залесово ул Советская 5"), ("залесово", "советская", frozenset(), "5", ""))
        self.assertEqual(parse_address("пгт Боровский ул Мира 7"), ("боровский", "мира", frozenset(), "7", ""))
        self.assertEqual(
            parse_address("ж/д ст Эбергард ул Центральная 1"),
            ("эбергард", "центральная", frozenset(), "1", ""),
        )
        self.assertEqual(
            parse_address("п ж/д ст Васюринская ул Садовая 2"),
            ("васюринская", "садовая", frozenset(), "2", ""),
        )
        self.assertEqual(parse_address("г Тюмень ул Ленина 10"), ("", "ленина", frozenset(), "10", ""))

    def test_settlement_type_human_spacing_punctuation_variants(self) -> None:
        expected = ("эбергард", "центральная", frozenset(), "1", "")
        for value in (
            "ж/д ст Эбергард ул Центральная 1",
            "ж-д ст. Эбергард ул Центральная 1",
            "ж.-д. ст. Эбергард ул Центральная 1",
            "ж д ст Эбергард ул Центральная 1",
        ):
            self.assertEqual(parse_address(value), expected)

        self.assertEqual(
            parse_address("п. ст. Перелесок ул Школьная 3"),
            ("перелесок", "школьная", frozenset(), "3", ""),
        )
        self.assertEqual(
            parse_address("п/ст Перелесок ул Школьная 3"),
            ("перелесок", "школьная", frozenset(), "3", ""),
        )
        self.assertEqual(
            parse_address("р.п. Благовещенка ул Ленина 4"),
            ("благовещенка", "ленина", frozenset(), "4", ""),
        )
        self.assertEqual(
            parse_address("ст.ца Анапская ул Садовая 2"),
            ("анапская", "садовая", frozenset(), "2", ""),
        )
        self.assertEqual(
            parse_address("ст ца Анапская ул Садовая 2"),
            ("анапская", "садовая", frozenset(), "2", ""),
        )

    def test_relaxed_house_match_protects_modifiers(self) -> None:
        self.assertFalse(_relaxed_house_check("10", "к2", "10", ""))
        self.assertFalse(_relaxed_house_check("10", "", "10", "к2"))
        self.assertTrue(_relaxed_house_check("10", "к2", "10", "к2"))
        self.assertTrue(_relaxed_house_check("10", "", "10", ""))

    def test_registry_address_anomaly_repairs_are_ported(self) -> None:
        self.assertEqual(
            parse_address("г. Тюмень, ул. Ленина, 10 стр. 2"),
            ("", "ленина", frozenset(), "10", "с2"),
        )
        self.assertEqual(
            parse_address("г. Тюмень, ул. Ленина, 10, офис 2"),
            ("", "ленина", frozenset(), "10", ""),
        )

        rendered = build_address_parts("г. Тюмень, Ленина, 10").rendered
        self.assertEqual(rendered, "г. Тюмень, ул. Ленина, д. 10")

        self.assertEqual(_normalize_street_slot("пр.Геологоразведчиков"), "пр-д Геологоразведчиков")
        self.assertEqual(_normalize_street_slot("ул. Ленина ЖК «Северный»"), "ул. Ленина")
        self.assertEqual(_normalize_street_slot("ул. Пермякова г. Тюмень"), "ул. Пермякова")
        self.assertEqual(_normalize_street_slot("ул. Салтыкова - Щедрина"), "ул. Салтыкова-Щедрина")
        self.assertEqual(_normalize_street_slot("ул. совесткая"), "ул. Советская")
        self.assertEqual(_normalize_street_slot("ул. советска"), "ул. Советская")
        self.assertEqual(_normalize_street_slot("ул. совсетов"), "ул. Советов")
        self.assertEqual(_normalize_street_slot("ул. прт Блюхера"), "пр-кт Блюхера")
        self.assertEqual(_normalize_house_slot("10, офис 2"), "д. 10")
        self.assertEqual(
            parse_address("г. Тюмень, г. Тюмень Червишевский тракт, д. 1"),
            ("", "червишевский", frozenset(), "1", ""),
        )
        self.assertEqual(
            build_address_parts("г. Тюмень, г. Тюмень Червишевский тракт, д. 1").rendered,
            "г. Тюмень, Червишевский тракт, д. 1",
        )
        self.assertEqual(parse_address("г. Тюмень, пр-кт Блюхера, д. 3")[1:4], ("блюхера", frozenset(), "3"))
        self.assertEqual(parse_address("г. Тюмень, ул. прт Блюхера, д. 3")[1:4], ("блюхера", frozenset(), "3"))
        self.assertEqual(parse_address("г. Тюмень, ул. совсетов, д. 4")[1:4], ("советов", frozenset(), "4"))
        self.assertEqual(parse_address("ул. 1905 года"), ("", "года", frozenset({"1905"}), "", ""))
        self.assertEqual(parse_address("ул. имени 1905 года"), ("", "года", frozenset({"1905"}), "", ""))
        self.assertEqual(
            parse_address("ул. имени 1905 года, д. 7"),
            ("", "года", frozenset({"1905"}), "7", ""),
        )
        self.assertEqual(build_address_parts("ул. 1905 года").rendered, "ул. 1905 года")
        self.assertEqual(build_address_parts("ул. имени 1905 года").rendered, "ул. имени 1905 года")
        self.assertEqual(build_address_parts("ул. имени 1905 года, д. 7").rendered, "ул. имени 1905 года, д. 7")
        normalizer = AddressNormalizer(use_oktmo=False)
        self.assertEqual(record_issue_codes(normalizer.normalize("ул. имени 1905 года, д. 7")), ())
        self.assertIn("street_without_house", record_issue_codes(normalizer.normalize("ул. имени 1905 года")))
        self.assertEqual(
            _extract_clean_address_fragments("Переход к ул. 1-я Линия - ул. Ленина"),
            [],
        )
        self.assertEqual(
            _extract_clean_address_fragments("ул. 1-я Линия, д. 7"),
            ["ул. 1-я Линия, д. 7"],
        )


class CompanionColumnTests(unittest.TestCase):
    def test_explicit_companion_modes(self) -> None:
        df = pd.DataFrame(
            {
                0: ["gt"],
                1: ["stable"],
                2: ["left object"],
                3: ["raw address"],
                4: ["right object"],
            }
        )
        leading_cols = [df.columns[0], df.columns[1]]
        target_cols = [df.columns[3]]

        self.assertEqual(_detect_companion_columns(df, leading_cols, target_cols, "none"), ({}, "none"))

        left_map, left_side = _detect_companion_columns(df, leading_cols, target_cols, "left")
        self.assertEqual(left_side, "left")
        self.assertEqual(left_map[df.columns[3]]["cols"], [df.columns[2]])

        right_map, right_side = _detect_companion_columns(df, leading_cols, target_cols, "right")
        self.assertEqual(right_side, "right")
        self.assertEqual(right_map[df.columns[3]]["cols"], [df.columns[4]])

        between_map, between_side = _detect_companion_columns(df, leading_cols, target_cols, "between")
        self.assertEqual(between_side, "left_block")
        self.assertEqual(between_map[df.columns[3]]["cols"], [df.columns[2]])


class AlignmentBehaviorTests(unittest.TestCase):
    def test_between_mode_moves_far_right_payload_with_matched_address(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            tmp_path = Path(tmp)
            input_file = tmp_path / "wide_payload.xlsx"
            output_file = tmp_path / "aligned.xlsx"
            column_count = 64
            gt_index = 7  # H
            payload_indexes = range(8, 63)  # I:BK
            target_index = 63  # BL
            addresses = [
                "Амурская область, город Благовещенск, ул. Воронкова, 26",
                "Амурская область, с. Нижние Бузули, ул. Ленина, д. 39",
                "Амурская область, г. Сковородино, ул. Победы, 10",
            ]
            original_rows: list[list[object]] = []
            for row_number, address in enumerate(addresses):
                row = [f"stable-{row_number}-{column}" for column in range(column_count)]
                row[gt_index] = address
                for column in payload_indexes:
                    row[column] = f"payload-{row_number}-{column}"
                row[target_index] = address
                original_rows.append(row)

            shuffled_rows = [list(row) for row in original_rows]
            for target_row, source_row in enumerate([2, 0, 1]):
                for column in [*payload_indexes, target_index]:
                    shuffled_rows[target_row][column] = original_rows[source_row][column]
            pd.DataFrame(shuffled_rows).to_excel(input_file, index=False, header=False)

            with contextlib.redirect_stdout(io.StringIO()), contextlib.redirect_stderr(io.StringIO()):
                process_file(
                    input_file,
                    output_file,
                    aligner_config={
                        "ground_truth_column": "H",
                        "target_columns": "BL",
                        "companion_mode": "between",
                    },
                )

            result = pd.read_excel(output_file, header=None)
            self.assertEqual(result.shape[0], len(original_rows))
            for row_number, address in enumerate(addresses):
                self.assertEqual(result.iat[row_number, target_index], address)
                self.assertEqual(
                    list(result.iloc[row_number, payload_indexes]),
                    [original_rows[row_number][column] for column in payload_indexes],
                )

    def test_post_match_normalization_keeps_value_when_core_would_degrade(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            tmp_path = Path(tmp)
            input_file = tmp_path / "input.xlsx"
            output_file = tmp_path / "output.xlsx"
            pd.DataFrame(
                [
                    [
                        "625007,Тюменская обл, г. Тюмень, проезд. 9 Мая, д. 2",
                        "обл. Тюменская, г. Тюмень, проезд. 9 Мая, д. 2",
                        "Тюмень Беляева 37 корпус 1",
                    ],
                    [
                        "625040,Тюменская обл, г. Тюмень, ул. Беляева, д. 37, корп. 1",
                        None,
                        "Тюмень Беляева 37 корпус 1",
                    ],
                ]
            ).to_excel(input_file, index=False, header=False)

            with contextlib.redirect_stdout(io.StringIO()), contextlib.redirect_stderr(io.StringIO()):
                process_file(
                    input_file,
                    output_file,
                    aligner_config={
                        "ground_truth_column": "A",
                        "target_columns": "B,C",
                        "companion_mode": "none",
                        "normalize_after_match": True,
                        "use_oktmo": False,
                    },
                )

            result = pd.read_excel(output_file, header=None)
            self.assertEqual(result.iat[0, 1], "обл. Тюменская, г. Тюмень, проезд. 9 Мая, д. 2")
            self.assertEqual(result.iat[1, 2], "Тюмень Беляева 37 корпус 1")

            workbook = load_workbook(output_file, read_only=True, data_only=True)
            try:
                self.assertIn("AUDIT_CONFLICTS", workbook.sheetnames)
                audit_sheet = workbook["AUDIT_CONFLICTS"]
                audit_headers = [cell.value for cell in next(audit_sheet.iter_rows(min_row=1, max_row=1))]
                audit_rows = list(audit_sheet.iter_rows(min_row=2, values_only=True))
                self.assertEqual(len(audit_rows), 3)
                reason_idx = audit_headers.index("Reason")
                action_idx = audit_headers.index("Action")
                rejected_idx = audit_headers.index("Rejected_Normalized_Value")
                self.assertEqual(
                    {row[reason_idx] for row in audit_rows},
                    {"post_normalization_lost_address_core"},
                )
                self.assertEqual({row[action_idx] for row in audit_rows}, {"kept_original_value"})
                self.assertTrue(any("проезд." in row[rejected_idx] for row in audit_rows))
                self.assertTrue(any("д. 37" in row[rejected_idx] for row in audit_rows))
            finally:
                workbook.close()

    def test_modified_gt_does_not_consume_clean_target_in_relaxed_tail(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            tmp_path = Path(tmp)
            input_file = tmp_path / "input.xlsx"
            output_file = tmp_path / "output.xlsx"
            pd.DataFrame(
                [
                    ["ул Ленина 10к2", "Ленина 10"],
                ]
            ).to_excel(input_file, index=False, header=False)

            with contextlib.redirect_stdout(io.StringIO()), contextlib.redirect_stderr(io.StringIO()):
                process_file(
                    input_file,
                    output_file,
                    aligner_config={
                        "ground_truth_column": "A",
                        "target_columns": "B",
                        "companion_mode": "none",
                    },
                )

            result = pd.read_excel(output_file, header=None)
            self.assertTrue(pd.isna(result.iloc[0, 1]))
            self.assertEqual(result.iloc[1, 1], "Ленина 10")

    def test_matcher_can_use_shared_normalizer_and_oktmo_scope(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            tmp_path = Path(tmp)
            data_root = tmp_path / "data"
            write_oktmo_fixture(data_root)
            input_file = tmp_path / "input.xlsx"
            output_file = tmp_path / "output.xlsx"
            pd.DataFrame(
                [
                    ["деревня Патрушева, улица Ленина, д. 10", "д Патрушева ул Ленина 10"],
                ]
            ).to_excel(input_file, index=False, header=False)

            with contextlib.redirect_stdout(io.StringIO()), contextlib.redirect_stderr(io.StringIO()):
                process_file(
                    input_file,
                    output_file,
                    aligner_config={
                        "ground_truth_column": "A",
                        "target_columns": "B",
                        "companion_mode": "none",
                        "normalize_before_match": True,
                        "use_oktmo": True,
                        "oktmo_data_dir": str(data_root),
                        "subject_ter_hint": "71",
                        "municipality_hint": "Тюменский муниципальный район",
                    },
                )

            result = pd.read_excel(output_file, header=None)
            self.assertEqual(result.iloc[0, 1], "д Патрушева ул Ленина 10")

    def test_collect_before_match_aligns_separate_wild_groups(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            tmp_path = Path(tmp)
            input_file = tmp_path / "wild_groups.xlsx"
            output_file = tmp_path / "aligned.xlsx"
            pd.DataFrame(
                [
                    ["ID", "Эталон", "Группа C", "Группа D"],
                    [1, "г. Тюмень, ул. Альфа, д. 2", "объект 77 | ул. Альфа, 2", None],
                    [2, "г. Тюмень, ул. Бета, д. 10", None, "площадка: ул. Бета, дом 10"],
                ]
            ).to_excel(input_file, index=False, header=False)

            with contextlib.redirect_stdout(io.StringIO()), contextlib.redirect_stderr(io.StringIO()):
                process_file(
                    input_file,
                    output_file,
                    aligner_config={
                        "ground_truth_column": "B",
                        "target_columns": "C,D",
                        "collect_before_match": True,
                        "normalize_before_match": True,
                        "use_oktmo": False,
                        "city": "Тюмень",
                        "companion_mode": "none",
                    },
                )

            workbook = load_workbook(output_file, read_only=True, data_only=True)
            try:
                sheet = workbook.worksheets[0]
                headers = [sheet.cell(row=1, column=column).value for column in range(1, sheet.max_column + 1)]
                self.assertEqual(headers[:4], ["ID", "Эталон", "Группа C", "Группа D"])
                self.assertIn("ул. Альфа", str(sheet.cell(row=2, column=3).value or ""))
                self.assertIn("д. 2", str(sheet.cell(row=2, column=3).value or ""))
                self.assertFalse(sheet.cell(row=2, column=4).value)
                self.assertFalse(sheet.cell(row=3, column=3).value)
                self.assertIn("ул. Бета", str(sheet.cell(row=3, column=4).value or ""))
                self.assertIn("д. 10", str(sheet.cell(row=3, column=4).value or ""))
            finally:
                workbook.close()

    def test_collect_before_match_moves_configured_satellite_columns(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            tmp_path = Path(tmp)
            input_file = tmp_path / "wild_satellites.xlsx"
            output_file = tmp_path / "aligned_satellites.xlsx"
            pd.DataFrame(
                [
                    ["ID", "Эталон", "Адрес C", "Объект C", "Статус C", "Адрес F", "Объект F"],
                    [1, "г. Тюмень, ул. Альфа, д. 2", "объект 77 | ул. Альфа, 2", "школа A", "активен", None, None],
                    [2, "г. Тюмень, ул. Бета, д. 10", None, None, None, "площадка: ул. Бета, дом 10", "сад B"],
                ]
            ).to_excel(input_file, index=False, header=False)

            with contextlib.redirect_stdout(io.StringIO()), contextlib.redirect_stderr(io.StringIO()):
                process_file(
                    input_file,
                    output_file,
                    aligner_config={
                        "ground_truth_column": "B",
                        "target_columns": "C,F",
                        "satellite_columns": "C=D:E; F=G",
                        "collect_before_match": True,
                        "normalize_before_match": True,
                        "use_oktmo": False,
                        "city": "Тюмень",
                        "companion_mode": "none",
                    },
                )

            workbook = load_workbook(output_file, read_only=True, data_only=True)
            try:
                sheet = workbook.worksheets[0]
                headers = [sheet.cell(row=1, column=column).value for column in range(1, sheet.max_column + 1)]
                self.assertEqual(headers[:7], ["ID", "Эталон", "Адрес C", "Объект C", "Статус C", "Адрес F", "Объект F"])
                self.assertIn("ул. Альфа", str(sheet.cell(row=2, column=3).value or ""))
                self.assertEqual(sheet.cell(row=2, column=4).value, "школа A")
                self.assertEqual(sheet.cell(row=2, column=5).value, "активен")
                self.assertFalse(sheet.cell(row=2, column=6).value)
                self.assertFalse(sheet.cell(row=2, column=7).value)
                self.assertFalse(sheet.cell(row=3, column=3).value)
                self.assertFalse(sheet.cell(row=3, column=4).value)
                self.assertFalse(sheet.cell(row=3, column=5).value)
                self.assertIn("ул. Бета", str(sheet.cell(row=3, column=6).value or ""))
                self.assertEqual(sheet.cell(row=3, column=7).value, "сад B")
            finally:
                workbook.close()

    def test_collect_before_match_keeps_same_address_with_different_satellites(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            tmp_path = Path(tmp)
            input_file = tmp_path / "duplicate_satellites.xlsx"
            output_file = tmp_path / "aligned_duplicate_satellites.xlsx"
            pd.DataFrame(
                [
                    ["ID", "Эталон", "Адрес C", "Объект C"],
                    [1, "г. Тюмень, ул. Альфа, д. 2", "ул. Альфа, д. 2", "школа A"],
                    [2, "г. Тюмень, ул. Альфа, д. 2", "ул. Альфа, д. 2", "филиал A"],
                    [3, "г. Тюмень, ул. Альфа, д. 2", "ул. Альфа, д. 2", "школа A"],
                ]
            ).to_excel(input_file, index=False, header=False)

            with contextlib.redirect_stdout(io.StringIO()), contextlib.redirect_stderr(io.StringIO()):
                process_file(
                    input_file,
                    output_file,
                    aligner_config={
                        "ground_truth_column": "B",
                        "target_columns": "C",
                        "satellite_columns": "C=D",
                        "collect_before_match": True,
                        "normalize_before_match": True,
                        "use_oktmo": False,
                        "city": "Тюмень",
                        "companion_mode": "none",
                    },
                )

            workbook = load_workbook(output_file, read_only=True, data_only=True)
            try:
                sheet = workbook.worksheets[0]
                rows = [
                    (
                        sheet.cell(row=row, column=3).value,
                        sheet.cell(row=row, column=4).value,
                    )
                    for row in range(2, sheet.max_row + 1)
                ]
                self.assertIn(("ул. Альфа, д. 2", "школа A"), rows)
                self.assertIn(("ул. Альфа, д. 2", "филиал A"), rows)
                self.assertEqual(rows.count(("ул. Альфа, д. 2", "школа A")), 1)
            finally:
                workbook.close()

    def test_collect_before_match_can_match_related_address_from_other_row(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            tmp_path = Path(tmp)
            input_file = tmp_path / "cross_row.xlsx"
            output_file = tmp_path / "aligned_cross_row.xlsx"
            pd.DataFrame(
                [
                    ["ID", "Эталон", "Адрес C", "Объект C"],
                    [1, "г. Тюмень, ул. Альфа, д. 2", None, None],
                    [2, "г. Тюмень, ул. Бета, д. 10", "ул. Альфа, д. 2", "школа A"],
                ]
            ).to_excel(input_file, index=False, header=False)

            with contextlib.redirect_stdout(io.StringIO()), contextlib.redirect_stderr(io.StringIO()):
                process_file(
                    input_file,
                    output_file,
                    aligner_config={
                        "ground_truth_column": "B",
                        "target_columns": "C",
                        "satellite_columns": "C=D",
                        "collect_before_match": True,
                        "normalize_before_match": True,
                        "use_oktmo": False,
                        "city": "Тюмень",
                        "companion_mode": "none",
                    },
                )

            workbook = load_workbook(output_file, read_only=True, data_only=True)
            try:
                sheet = workbook.worksheets[0]
                self.assertIn("ул. Альфа", str(sheet.cell(row=2, column=3).value or ""))
                self.assertEqual(sheet.cell(row=2, column=4).value, "школа A")
                self.assertFalse(sheet.cell(row=3, column=3).value)
                self.assertFalse(sheet.cell(row=3, column=4).value)
            finally:
                workbook.close()

    def test_normalize_after_match_keeps_legacy_matcher_and_cleans_output(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            tmp_path = Path(tmp)
            input_file = tmp_path / "post_normalize.xlsx"
            output_file = tmp_path / "post_normalized.xlsx"
            pd.DataFrame(
                [
                    ["ID", "Эталон", "Адрес C", "Объект C"],
                    [1, "ул Альфа 2", "улица Альфа дом 2", "школа A"],
                ]
            ).to_excel(input_file, index=False, header=False)

            with contextlib.redirect_stdout(io.StringIO()) as stdout, contextlib.redirect_stderr(io.StringIO()):
                process_file(
                    input_file,
                    output_file,
                    aligner_config={
                        "ground_truth_column": "B",
                        "target_columns": "C",
                        "satellite_columns": "C=D",
                        "normalize_after_match": True,
                        "normalize_before_match": False,
                        "use_oktmo": False,
                        "city": "Тюмень",
                        "companion_mode": "none",
                    },
                )

            self.assertIn("Address normalizer for matching: OFF", stdout.getvalue())
            self.assertIn("Address normalizer after matching: ON", stdout.getvalue())
            workbook = load_workbook(output_file, read_only=True, data_only=True)
            try:
                sheet = workbook.worksheets[0]
                self.assertEqual(sheet.cell(row=2, column=3).value, "ул. Альфа, д. 2")
                self.assertEqual(sheet.cell(row=2, column=4).value, "школа A")
            finally:
                workbook.close()

    def test_fuzzy_match_can_place_first_target_row(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            tmp_path = Path(tmp)
            data_root = tmp_path / "data"
            write_oktmo_fixture(data_root)
            input_file = tmp_path / "input.xlsx"
            output_file = tmp_path / "output.xlsx"
            pd.DataFrame(
                [
                    ["деревня Патрушева, улица Ленина, д. 10", "улица Ленина, дом 10"],
                ]
            ).to_excel(input_file, index=False, header=False)

            with contextlib.redirect_stdout(io.StringIO()), contextlib.redirect_stderr(io.StringIO()):
                process_file(
                    input_file,
                    output_file,
                    aligner_config={
                        "ground_truth_column": "A",
                        "target_columns": "B",
                        "companion_mode": "none",
                        "normalize_before_match": True,
                        "use_oktmo": True,
                        "oktmo_data_dir": str(data_root),
                        "subject_ter_hint": "71",
                    },
                )

            result = pd.read_excel(output_file, header=None)
            self.assertEqual(len(result), 1)
            self.assertEqual(result.iloc[0, 1], "улица Ленина, дом 10")


class AlignmentServiceConfigTests(unittest.TestCase):
    def test_runtime_config_accepts_headers_and_oktmo_scope(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            tmp_path = Path(tmp)
            data_root = tmp_path / "data"
            write_oktmo_fixture(data_root)
            input_file = tmp_path / "input.xlsx"
            pd.DataFrame(
                [
                    {
                        "Эталонный адрес": "д Патрушева ул Ленина 10",
                        "Источник": "деревня Патрушева, улица Ленина, д. 10",
                    }
                ]
            ).to_excel(input_file, index=False)
            logs: list[str] = []
            context = SimpleNamespace(
                operation=SimpleNamespace(
                    parameters={
                        "ground_truth_column_header": "Эталонный адрес",
                        "target_column_header": "Источник",
                        "companion_mode": "none",
                        "satellite_columns": "B=C:D",
                        "normalize_before_match": True,
                        "normalize_after_match": True,
                        "collect_before_match": True,
                        "collect_whole_document": True,
                        "use_oktmo": True,
                        "oktmo_scope_enabled": True,
                        "oktmo_scope_region": "71",
                        "oktmo_scope_municipality": "Тюменский муниципальный район",
                        "oktmo_data_dir": str(data_root),
                    }
                ),
                paths=SimpleNamespace(root=tmp_path),
                log=logs.append,
            )

            config = _aligner_runtime_config(context, input_file)

            self.assertEqual(config["ground_truth_column"], "A")
            self.assertEqual(config["target_columns"], "B")
            self.assertEqual(config["satellite_columns"], "B=C:D")
            self.assertTrue(config["normalize_before_match"])
            self.assertTrue(config["normalize_after_match"])
            self.assertTrue(config["collect_before_match"])
            self.assertTrue(config["collect_whole_document"])
            self.assertEqual(config["subject_ter_hint"], "71")
            self.assertEqual(config["municipality_hint"], "Тюменский муниципальный район")
            self.assertTrue(any("Address matching normalization/post-processing runtime: enabled" in message for message in logs))

    def test_alignment_mode_enables_post_match_normalization_only(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            tmp_path = Path(tmp)
            logs: list[str] = []
            context = SimpleNamespace(
                operation=SimpleNamespace(
                    parameters={
                        "alignment_mode": "normalized",
                        "use_oktmo": False,
                    }
                ),
                paths=SimpleNamespace(root=tmp_path),
                log=logs.append,
            )

            config = _aligner_runtime_config(context)

            self.assertEqual(config["alignment_mode"], "normalized")
            self.assertFalse(config["normalize_before_match"])
            self.assertTrue(config["normalize_after_match"])
            self.assertFalse(config["collect_before_match"])
            self.assertTrue(any("Address matching normalization/post-processing runtime: enabled" in message for message in logs))

    def test_alignment_runtime_summary_uses_common_normalizer_report(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            tmp_path = Path(tmp)
            data_root = tmp_path / "data"
            write_oktmo_fixture(data_root)
            normalizer = AddressNormalizer(
                data_dir=data_root,
                use_oktmo=True,
                subject_ter_hint="71",
                municipality_hint="Тюменский муниципальный район",
            )

            summary = _address_runtime_summary(
                AddressRuntimeContext(
                    normalizer=normalizer,
                    subject_ter_hint="71",
                    municipality_hint="Тюменский муниципальный район",
                    oktmo_scope_enabled=True,
                )
            )

            self.assertTrue(summary["use_oktmo"])
            self.assertTrue(summary["oktmo_scope_enabled"])
            self.assertEqual(
                summary["oktmo_scope"],
                {
                    "enabled": True,
                    "use_oktmo": True,
                    "data_dir": str(data_root),
                    "subject_ter_hint": "71",
                    "municipality_hint": "Тюменский муниципальный район",
                },
            )
            self.assertEqual(summary["normalizer"]["data_dir"], str(data_root))
            self.assertEqual(summary["normalizer"]["subject_ter_hint"], "71")
            self.assertEqual(
                summary["normalizer"]["oktmo_key_profile"]["subject_ter"],
                summary["oktmo_key_profile"]["subject_ter"],
            )

    def test_artifact_summary_normalizes_operation_payload_paths(self) -> None:
        artifacts = _artifact_summary(
            outputs=[Path("out.xlsx"), "out.xlsx", None],
            reports=Path("report.json"),
            staging="stage.xlsx",
            inputs={"source": Path("input.xlsx")},
        )

        self.assertEqual(artifacts["outputs"], ["out.xlsx"])
        self.assertEqual(artifacts["reports"], ["report.json"])
        self.assertEqual(artifacts["staging"], ["stage.xlsx"])
        self.assertEqual(artifacts["inputs"], ["input.xlsx"])


class ServiceEndToEndSmokeTests(unittest.TestCase):
    SERVICE = "system_core.services.address_aligner_service"

    def _project_paths(self, root: Path) -> ProjectPaths:
        paths = ProjectPaths(
            root=root,
            input=root / "input",
            output=root / "output",
            logs=root / "logs",
            report=root / "report",
            workspace=root / "workspace",
            config=root / "config",
            release=root / "release",
            system_core=ROOT / "system_core",
        )
        ensure_project_dirs(paths)
        return paths

    def _operation(self, operation_id: str, function_name: str, parameters: dict[str, object]) -> Operation:
        return Operation(
            id=operation_id,
            title=operation_id,
            description=operation_id,
            service=f"{self.SERVICE}:{function_name}",
            parameters=parameters,
        )

    def _run_operation(self, paths: ProjectPaths, operation_id: str, function_name: str, parameters: dict[str, object]):
        logs: list[str] = []
        result = execute_operation(
            paths,
            self._operation(operation_id, function_name, parameters),
            log_callback=logs.append,
        )
        self.assertTrue(result.ok, result.message + "\n" + "\n".join(logs[-30:]))
        self.assertIn("artifacts", result.data)
        self._assert_artifacts_exist(result.data["artifacts"])
        return result

    def _assert_artifacts_exist(self, artifacts: dict[str, object]) -> None:
        for bucket in ("outputs", "reports", "staging"):
            for path_value in artifacts.get(bucket, []):
                self.assertTrue(Path(str(path_value)).exists(), f"Missing {bucket} artifact: {path_value}")

    def test_service_smoke_collects_and_appends_clean_address_column(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            paths = self._project_paths(root)
            (paths.input / "source.txt").write_text(
                "\n".join(
                    [
                        "г. Тюмень, ул. Альфа, д. 2",
                        "Тюмень, улица Альфа, дом 2",
                        "г. Тюмень, ул. Бета",
                    ]
                ),
                encoding="utf-8",
            )

            collect = self._run_operation(
                paths,
                "run_address_collection",
                "run_address_collection",
                {
                    "address_collection_action": "generate",
                    "address_collection_dir": str(root / "address_collection"),
                    "target_column": "A",
                    "use_oktmo": False,
                    "enabled_slots": ["locality", "street", "house", "match_key"],
                    "slot_order": ["locality", "street", "house"],
                    "sanitize_incomplete": True,
                },
            )
            collect_summary = json.loads(Path(str(collect.data["report"])).read_text(encoding="utf-8"))
            self.assertIn("slot_fill", collect_summary)
            self.assertEqual(collect_summary["slot_fill"]["slots"]["street"]["filled"], 1)
            self.assertEqual(collect_summary["slot_fill"]["slots"]["house"]["filled"], 1)
            self.assertEqual(collect.data["artifacts"]["outputs"], [collect.data["output"]])
            self.assertEqual(collect.data["artifacts"]["staging"], [collect.data["staging"]])

            ready_file = root / "address_collection" / "ready.xlsx"
            pd.DataFrame({"Address": ["г. Тюмень, ул. Гамма, д. 3"]}).to_excel(ready_file, index=False)
            append = self._run_operation(
                paths,
                "run_address_collection",
                "run_address_collection",
                {
                    "address_collection_action": "append",
                    "address_collection_dir": str(root / "address_collection"),
                    "collected_file": str(collect.data["output"]),
                    "reference_file": str(ready_file),
                    "collected_column": "A",
                    "target_column": "Address",
                    "use_oktmo": False,
                },
            )
            append_summary = json.loads(Path(str(append.data["report"])).read_text(encoding="utf-8"))
            self.assertEqual(append_summary["source_contract"]["kind"], "clean_address_column")
            self.assertFalse(append_summary["source_contract"]["uses_staging"])
            self.assertEqual(append.data["artifacts"]["outputs"], [append.data["output"]])
            self.assertEqual(append.data["artifacts"]["staging"], [])

    def test_service_smoke_normalizes_processes_and_aligns_addresses(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            paths = self._project_paths(root)
            pd.DataFrame(
                [
                    ["г. Тюмень, ул. Бета, д. 10"],
                    ["г. Тюмень, ул. Альфа, д. 2"],
                    ["Тюмень, улица Альфа, дом 2"],
                ]
            ).to_excel(paths.input / "dirty_reference.xlsx", index=False, header=False)

            normalize = self._run_operation(
                paths,
                "run_reference_normalization",
                "run_reference_normalization",
                {
                    "reference_action": "generate",
                    "use_oktmo": False,
                    "enabled_slots": ["locality", "street", "house", "match_key"],
                    "slot_order": ["locality", "street", "house"],
                },
            )
            normalize_summary = json.loads(Path(str(normalize.data["report"])).read_text(encoding="utf-8"))
            self.assertIn("slot_fill", normalize_summary)
            self.assertGreaterEqual(normalize_summary["slot_fill"]["rows_total"], 2)
            self.assertEqual(normalize.data["artifacts"]["outputs"], [normalize.data["output"]])
            self.assertEqual(normalize.data["artifacts"]["staging"], [normalize.data["staging"]])

            process = self._run_operation(
                paths,
                "run_reference_processing",
                "run_reference_processing",
                {
                    "reference_processing_action": "sort_deconstruct",
                    "reference_file": str(normalize.data["output"]),
                    "reference_columns": "Normalized_Address",
                    "use_oktmo": False,
                    "enabled_slots": ["street", "house"],
                    "slot_order": ["street", "house"],
                    "remove_source_column": True,
                },
            )
            process_summary = json.loads(Path(str(process.data["report"])).read_text(encoding="utf-8"))
            self.assertIn("slot_fill", process_summary)
            self.assertEqual(process.data["artifacts"]["outputs"], [process.data["output"]])
            self.assertEqual(process.data["artifacts"]["reports"], [process.data["report"]])

            for file_path in paths.input.glob("*.xlsx"):
                file_path.unlink()
            pd.DataFrame(
                [
                    ["г. Тюмень, ул. Альфа, д. 2", "Тюмень ул Альфа 2"],
                ]
            ).to_excel(paths.input / "align.xlsx", index=False, header=False)
            align = self._run_operation(
                paths,
                "run_address_aligner",
                "run_address_aligner",
                {
                    "ground_truth_column": "A",
                    "target_columns": "B",
                    "companion_mode": "none",
                    "normalize_before_match": True,
                    "use_oktmo": False,
                },
            )
            align_summary = json.loads(Path(str(align.data["report"])).read_text(encoding="utf-8"))
            self.assertEqual(align_summary["mode"], "address-alignment")
            self.assertIn("normalizer", align_summary["address_runtime"])
            self.assertEqual(align.data["artifacts"]["outputs"], align.data["outputs"])
            self.assertEqual(align.data["artifacts"]["reports"], [align.data["report"]])


class ManifestConsistencyTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        manifest_path = ROOT / "config" / "tool_manifest.yaml"
        manifest = yaml.safe_load(manifest_path.read_text(encoding="utf-8"))
        cls.operations = {}

        def collect(items: list[dict[str, object]]) -> None:
            for item in items:
                if not isinstance(item, dict):
                    continue
                cls.operations[str(item.get("id") or "")] = item
                children = item.get("children", [])
                if isinstance(children, list):
                    collect(children)

        collect(manifest.get("operation_groups", []))

    def _fields(self, operation_id: str) -> dict[str, dict[str, object]]:
        return {field["id"]: field for field in self.operations[operation_id]["fields"]}

    def _field_ids(self, operation_id: str) -> list[str]:
        return [str(field["id"]) for field in self.operations[operation_id]["fields"]]

    def test_slot_controls_match_engine_defaults(self) -> None:
        for operation_id in ("run_address_collection", "run_reference_normalization", "run_reference_processing"):
            with self.subTest(operation_id=operation_id):
                fields = self._fields(operation_id)
                enabled_options = fields["enabled_slots"]["options"]
                self.assertEqual(tuple(option["value"] for option in enabled_options), DEFAULT_ENABLED_SLOTS)
                checked_defaults = tuple(option["value"] for option in enabled_options if option.get("default") is True)
                self.assertEqual(checked_defaults, DEFAULT_SELECTED_SLOTS)
                self.assertFalse(next(option for option in enabled_options if option["value"] == "microdistrict").get("default"))

                order_field = fields["slot_order"]
                self.assertEqual(tuple(order_field["default"].keys()), DEFAULT_SLOT_ORDER)
                self.assertEqual(tuple(option["value"] for option in order_field["options"]), DEFAULT_SLOT_ORDER)
                self.assertEqual(
                    [option["default_order"] for option in order_field["options"]],
                    list(range(1, len(DEFAULT_SLOT_ORDER) + 1)),
                )

    def test_oktmo_controls_are_split_between_address_screens_and_child(self) -> None:
        for operation_id in (
            "run_address_collection",
            "run_address_aligner",
            "run_reference_normalization",
            "run_reference_processing",
        ):
            with self.subTest(operation_id=operation_id):
                fields = self._fields(operation_id)
                self.assertTrue(set(OKTMO_SCOPE_PARAMETER_IDS).issubset(fields))
                self.assertEqual(fields["use_oktmo"].get("default"), True)
                self.assertEqual(fields["use_oktmo"].get("label_ru"), "Подключить ключи ОКТМО")
                self.assertIn("Если файл ключей пустой", str(fields["use_oktmo"].get("hint_ru")))
                self.assertNotIn("oktmo_scope_enabled", fields)
                self.assertNotIn("oktmo_scope_region", fields)
                self.assertNotIn("oktmo_scope_municipality", fields)
                self.assertNotIn("oktmo_key_toolbar", fields)

        action_field_ids = {
            "run_address_collection": "address_collection_action",
            "run_reference_normalization": "reference_action",
            "run_reference_processing": "reference_processing_action",
        }
        for operation_id, action_id in action_field_ids.items():
            with self.subTest(operation_id=operation_id, order="use_oktmo_after_action"):
                field_ids = self._field_ids(operation_id)
                self.assertEqual(field_ids[field_ids.index(action_id) + 1], "use_oktmo")
                self.assertEqual(self._fields(operation_id)["use_oktmo"].get("group"), "parameters")

        align_fields = self._fields("run_address_aligner")
        self.assertTrue(set(ADDRESS_RUNTIME_PARAMETER_IDS).issubset(align_fields))
        self.assertEqual(align_fields["alignment_mode"].get("group"), "alignment_mode")
        self.assertEqual(align_fields["alignment_mode"].get("default"), "fast")
        self.assertEqual(
            [option["value"] for option in align_fields["alignment_mode"]["options"]],
            ["fast", "normalized"],
        )
        self.assertEqual(align_fields["normalize_before_match"].get("default"), False)
        self.assertEqual(align_fields["normalize_after_match"].get("default"), False)
        self.assertEqual(align_fields["collect_before_match"].get("default"), False)
        self.assertEqual(align_fields["collect_whole_document"].get("default"), False)
        self.assertTrue(align_fields["normalize_after_match"].get("advanced"))
        self.assertEqual(align_fields["normalize_after_match"].get("group"), "alignment_normalization")
        self.assertEqual(align_fields["use_oktmo"].get("group"), "alignment_normalization")
        self.assertEqual(align_fields["satellite_columns"].get("group"), "related_columns")

        oktmo_child_fields = self._fields("oktmo_key_console")
        self.assertEqual(
            oktmo_child_fields["oktmo_scope_region"].get("options_source"),
            "system_core.services.address_aligner_service:oktmo_region_scope_options",
        )
        self.assertEqual(oktmo_child_fields["oktmo_scope_region"].get("refresh_on_change"), True)
        self.assertEqual(
            oktmo_child_fields["oktmo_scope_municipality"].get("options_source"),
            "system_core.services.address_aligner_service:oktmo_municipality_scope_options",
        )
        self.assertEqual(oktmo_child_fields["oktmo_scope_municipality"].get("depends_on"), ["oktmo_scope_region"])
        self.assertEqual(oktmo_child_fields["oktmo_key_toolbar"].get("type"), "oktmo_key_toolbar")
        self.assertEqual(self.operations["oktmo_tools"].get("auto_open_child"), "oktmo_key_console")
        self.assertIn("update_rosstat_oktmo", self.operations)
        self.assertTrue(self.operations["update_rosstat_oktmo"].get("gui_hidden"))

    def test_clean_output_workbook_checkbox_is_available_for_generation_modes(self) -> None:
        for operation_id in ("run_address_collection", "run_reference_normalization"):
            with self.subTest(operation_id=operation_id):
                fields = self._fields(operation_id)
                self.assertIn("clean_output_workbook", fields)
                self.assertEqual(fields["clean_output_workbook"].get("type"), "checkbox")
                self.assertEqual(fields["clean_output_workbook"].get("group"), "options")
                self.assertEqual(fields["clean_output_workbook"].get("default"), False)


class RosstatUpdateTests(unittest.TestCase):
    def test_configured_rosstat_url_prefers_explicit_then_audion_env(self) -> None:
        from system_core.address_engine import rosstat_update  # noqa: WPS433

        env_names = rosstat_update.OKTMO_URL_ENV_NAMES
        previous = {name: os.environ.get(name) for name in env_names}
        try:
            for name in env_names:
                os.environ.pop(name, None)
            self.assertEqual(rosstat_update._configured_source_url(None), rosstat_update.DEFAULT_OKTMO_URL)

            os.environ["AUDION_ROSSTAT_OKTMO_URL"] = "https://example.test/env-data.csv"
            self.assertEqual(rosstat_update._configured_source_url(None), "https://example.test/env-data.csv")
            self.assertEqual(
                rosstat_update._configured_source_url("https://example.test/direct-data.csv"),
                "https://example.test/direct-data.csv",
            )
        finally:
            for name, value in previous.items():
                if value is None:
                    os.environ.pop(name, None)
                else:
                    os.environ[name] = value

    def test_rosstat_service_lets_env_override_manifest_default_url(self) -> None:
        from system_core.services import address_aligner_service as service  # noqa: WPS433

        calls: list[dict[str, object]] = []
        clear_calls: list[bool] = []
        warm_calls: list[Path] = []
        original_update = service.update_oktmo_snapshot
        original_clear = service.clear_oktmo_lookup_caches
        original_warm = service.warm_oktmo_cache

        def fake_update(**kwargs: object) -> dict[str, object]:
            calls.append(kwargs)
            return {
                "requested_url": kwargs.get("source_url") or "default/env",
                "latest_data_file": "data-20260101T0000-structure-20260101T0000.csv",
                "summary_path": str(Path(str(kwargs["report_dir"])) / "rosstat_oktmo_update_summary.json"),
            }

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            context = SimpleNamespace(
                operation=SimpleNamespace(parameters={"rosstat_oktmo_url": service.DEFAULT_OKTMO_URL}),
                paths=SimpleNamespace(root=root, report=root / "report"),
                log=lambda _message: None,
                progress=lambda _value: None,
                cancelled=lambda: False,
            )
            service.update_oktmo_snapshot = fake_update
            service.clear_oktmo_lookup_caches = lambda: clear_calls.append(True)
            service.warm_oktmo_cache = lambda data_dir: warm_calls.append(Path(data_dir)) or {"regions": 89}
            try:
                result = service.update_rosstat_oktmo_data(context)
            finally:
                service.update_oktmo_snapshot = original_update
                service.clear_oktmo_lookup_caches = original_clear
                service.warm_oktmo_cache = original_warm

        self.assertEqual(calls[0]["source_url"], None)
        self.assertEqual(clear_calls, [True])
        self.assertEqual(warm_calls, [root / "data"])
        self.assertTrue(result["caches_cleared"])
        self.assertEqual(result["cache"], {"regions": 89})

    def test_rosstat_service_keeps_explicit_non_default_url(self) -> None:
        from system_core.services import address_aligner_service as service  # noqa: WPS433

        calls: list[dict[str, object]] = []
        clear_calls: list[bool] = []
        warm_calls: list[Path] = []
        original_update = service.update_oktmo_snapshot
        original_clear = service.clear_oktmo_lookup_caches
        original_warm = service.warm_oktmo_cache

        def fake_update(**kwargs: object) -> dict[str, object]:
            calls.append(kwargs)
            return {
                "requested_url": kwargs.get("source_url"),
                "latest_data_file": "data-20260101T0000-structure-20260101T0000.csv",
                "summary_path": str(Path(str(kwargs["report_dir"])) / "rosstat_oktmo_update_summary.json"),
            }

        explicit_url = "https://example.test/custom-data.csv"
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            context = SimpleNamespace(
                operation=SimpleNamespace(parameters={"rosstat_oktmo_url": explicit_url}),
                paths=SimpleNamespace(root=root, report=root / "report"),
                log=lambda _message: None,
                progress=lambda _value: None,
                cancelled=lambda: False,
            )
            service.update_oktmo_snapshot = fake_update
            service.clear_oktmo_lookup_caches = lambda: clear_calls.append(True)
            service.warm_oktmo_cache = lambda data_dir: warm_calls.append(Path(data_dir)) or {"regions": 89}
            try:
                service.update_rosstat_oktmo_data(context)
            finally:
                service.update_oktmo_snapshot = original_update
                service.clear_oktmo_lookup_caches = original_clear
                service.warm_oktmo_cache = original_warm

        self.assertEqual(calls[0]["source_url"], explicit_url)
        self.assertEqual(clear_calls, [True])
        self.assertEqual(warm_calls, [root / "data"])

    def test_clear_oktmo_caches_drops_parsed_and_option_state(self) -> None:
        from system_core.address_engine import oktmo_lookup  # noqa: WPS433

        oktmo_lookup._INDEX = {"sample": [("1", "name", "71")]}
        oktmo_lookup._INDEX_KEYS = ["sample"]
        oktmo_lookup._MUNICIPALITY_BY_CODE = {"1": "municipality"}
        oktmo_lookup._MUNICIPALITY_BY_LOCALITY_KEY = {"sample": [("municipality", "71")]}
        oktmo_lookup._CONTEXT_BY_CODE = {}
        oktmo_lookup._SUBJECT_TER_BY_REGION_KEY = {"sample": "71"}
        oktmo_lookup._INDEX_SOURCE_KEY = ("data.csv", 1.0)
        oktmo_lookup._LOOKUP_OKTMO_CACHE[("a", "b", "c", "d")] = ("1", "name")
        oktmo_lookup._LOOKUP_MUNICIPALITY_CACHE[("a", "b", "c", "d", "e")] = "municipality"
        oktmo_lookup._LOOKUP_CONTEXT_CACHE[("a", "b", "c", "d")] = None
        oktmo_lookup._REGION_SCOPE_OPTIONS_CACHE = (("data.csv", 1.0), [{"value": "71"}])
        oktmo_lookup._MUNICIPALITY_SCOPE_OPTIONS_CACHE[((("data.csv", 1.0)), "71")] = [{"value": "x"}]
        oktmo_lookup._SCOPE_KEY_PROFILE_CACHE[((("data.csv", 1.0)), "71", "")] = {"items": []}

        oktmo_lookup.clear_oktmo_caches()

        self.assertIsNone(oktmo_lookup._INDEX)
        self.assertIsNone(oktmo_lookup._INDEX_KEYS)
        self.assertIsNone(oktmo_lookup._MUNICIPALITY_BY_CODE)
        self.assertIsNone(oktmo_lookup._MUNICIPALITY_BY_LOCALITY_KEY)
        self.assertIsNone(oktmo_lookup._CONTEXT_BY_CODE)
        self.assertIsNone(oktmo_lookup._SUBJECT_TER_BY_REGION_KEY)
        self.assertIsNone(oktmo_lookup._INDEX_SOURCE_KEY)
        self.assertIsNone(oktmo_lookup._REGION_SCOPE_OPTIONS_CACHE)
        self.assertEqual(oktmo_lookup._LOOKUP_OKTMO_CACHE, {})
        self.assertEqual(oktmo_lookup._LOOKUP_MUNICIPALITY_CACHE, {})
        self.assertEqual(oktmo_lookup._LOOKUP_CONTEXT_CACHE, {})
        self.assertEqual(oktmo_lookup._MUNICIPALITY_SCOPE_OPTIONS_CACHE, {})
        self.assertEqual(oktmo_lookup._SCOPE_KEY_PROFILE_CACHE, {})

    def test_warm_oktmo_cache_serializes_parallel_first_load(self) -> None:
        from concurrent.futures import ThreadPoolExecutor  # noqa: WPS433

        from system_core.address_engine import oktmo_lookup  # noqa: WPS433

        with tempfile.TemporaryDirectory() as tmp:
            data_root = Path(tmp) / "data"
            write_oktmo_fixture(data_root)
            oktmo_lookup.clear_oktmo_caches(data_root)
            original_build = oktmo_lookup._build_index
            build_calls: list[bool] = []

            def counted_build(*args: object, **kwargs: object) -> object:
                build_calls.append(True)
                return original_build(*args, **kwargs)

            oktmo_lookup._build_index = counted_build
            try:
                with ThreadPoolExecutor(max_workers=8) as executor:
                    statuses = list(executor.map(lambda _index: oktmo_lookup.warm_oktmo_cache(data_root), range(8)))
            finally:
                oktmo_lookup._build_index = original_build
                oktmo_lookup.clear_oktmo_caches(data_root)

        self.assertEqual(len(build_calls), 1)
        self.assertTrue(all(status == statuses[0] for status in statuses))
        self.assertTrue(statuses[0]["cache_exists"])
        self.assertGreater(statuses[0]["cache_bytes"], 0)
        self.assertGreater(statuses[0]["regions"], 0)


class ReferenceNormalizationTests(unittest.TestCase):
    def test_postal_index_is_slot_and_oktmo_enriches_parts(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            data_root = Path(tmp) / "data"
            write_oktmo_fixture(data_root)
            normalizer = AddressNormalizer(data_dir=data_root, subject_ter_hint="71")

            record = normalizer.normalize("625000, деревня Патрушева, улица Ленина, д. 10")

            self.assertEqual(record.postal_index, "625000")
            self.assertEqual(record.oktmo_code, "71001001001")
            self.assertEqual(record.municipality, "Тюменский муниципальный район")
            self.assertEqual(record.parts.postal_index, "625000")
            self.assertEqual(record.parts.locality, "д. Патрушева")
            self.assertIn("Ленина", record.normalized or "")

    def test_oktmo_scope_options_feed_region_and_municipality_hints(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            tmp_path = Path(tmp)
            data_root = tmp_path / "data"
            missing_status = oktmo_data_status(tmp_path, {"oktmo_data_dir": str(data_root)}, {})
            self.assertFalse(missing_status["ok"])

            write_oktmo_fixture(data_root)
            ready_status = oktmo_data_status(tmp_path, {"oktmo_data_dir": str(data_root)}, {})
            self.assertTrue(ready_status["ok"])
            self.assertEqual(ready_status["files"], 1)

            regions = oktmo_region_scope_options(tmp_path, {"oktmo_data_dir": str(data_root)}, {})
            self.assertTrue(any(option.get("value") == "71" for option in regions))
            self.assertTrue(any(option.get("value") == "71:800-899" for option in regions))
            self.assertTrue(any(option.get("value") == "71:900-999" for option in regions))

            pending_municipalities = oktmo_municipality_scope_options(
                tmp_path,
                {"oktmo_data_dir": str(data_root)},
                {},
            )
            self.assertEqual(pending_municipalities[0].get("value"), "")
            self.assertIn("Сначала выберите регион", str(pending_municipalities[0].get("label_ru") or ""))
            self.assertFalse(any(option.get("name") for option in pending_municipalities))

            municipalities = oktmo_municipality_scope_options(
                tmp_path,
                {"oktmo_data_dir": str(data_root), "oktmo_scope_region": "71"},
                {},
            )
            self.assertTrue(any(option.get("name") == "Тюменский муниципальный район" for option in municipalities))
            self.assertTrue(any(option.get("value") == "71001001" for option in municipalities))
            self.assertFalse(any("Сургут" in str(option.get("label_ru") or option.get("label") or "") for option in municipalities))

            hmao_municipalities = oktmo_municipality_scope_options(
                tmp_path,
                {"oktmo_data_dir": str(data_root), "oktmo_scope_region": "71:800-899"},
                {},
            )
            self.assertTrue(any(option.get("name") == "городской округ Сургут" for option in hmao_municipalities))

            tyumen_surgut = oktmo_find_settlement_candidates(
                tmp_path,
                {"oktmo_data_dir": str(data_root), "oktmo_scope_region": "71", "oktmo_settlement_query": "Сургут"},
            )
            self.assertEqual(tyumen_surgut["count"], 0)
            hmao_surgut = oktmo_find_settlement_candidates(
                tmp_path,
                {"oktmo_data_dir": str(data_root), "oktmo_scope_region": "71:800-899", "oktmo_settlement_query": "Сургут"},
            )
            self.assertGreaterEqual(hmao_surgut["count"], 1)

            normalizer = AddressNormalizer(
                data_dir=data_root,
                subject_ter_hint="71",
                municipality_hint="Тюменский муниципальный район",
            )
            record = normalizer.normalize("д Патрушева ул Ленина 10")
            self.assertEqual(record.oktmo_code, "71001001001")
            self.assertEqual(record.municipality, "Тюменский муниципальный район")

    def test_oktmo_child_settlement_candidate_updates_current_keys(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            tmp_path = Path(tmp)
            data_root = tmp_path / "data"
            write_oktmo_fixture(data_root)
            values = {
                "oktmo_data_dir": str(data_root),
                "oktmo_scope_region": "71",
                "oktmo_settlement_query": "Патрушева",
            }

            found = oktmo_find_settlement_candidates(tmp_path, values)
            updates = found["field_updates"]
            self.assertGreaterEqual(found["count"], 1)
            self.assertTrue(updates["oktmo_candidate_options"])
            self.assertEqual(updates["oktmo_candidate_ref"], "71001001001")

            values.update(updates)
            added = oktmo_add_candidate_keys(tmp_path, values)
            self.assertGreater(added["added"], 0)
            self.assertEqual(added["field_updates"]["oktmo_selected_settlement"], "Патрушева")
            keys = (tmp_path / "config" / "oktmo_current_keys.txt").read_text(encoding="utf-8")
            self.assertIn("Патрушева", keys)

    def test_oktmo_municipality_scope_plus_adds_mo_and_all_settlement_keys(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            tmp_path = Path(tmp)
            data_root = tmp_path / "data"
            write_oktmo_fixture(data_root)
            values = {
                "oktmo_data_dir": str(data_root),
                "oktmo_scope_region": "71",
                "oktmo_scope_municipality": "71001001",
            }

            added = oktmo_add_scope_keys(tmp_path, values, "municipality")

            self.assertGreater(added["added"], 0)
            keys = (tmp_path / "config" / "oktmo_current_keys.txt").read_text(encoding="utf-8")
            self.assertIn("Тюменский муниципальный район", keys)
            self.assertIn("Патрушева", keys)
            self.assertIn("Мирный", keys)
            self.assertIn("Центральный", keys)
            self.assertIn("Ключи", keys)

    def test_oktmo_pins_export_import_and_clear_roundtrip(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            tmp_path = Path(tmp)
            data_root = tmp_path / "data"
            write_oktmo_fixture(data_root)
            values = {"oktmo_data_dir": str(data_root), "oktmo_scope_region": "71"}

            region_result = oktmo_pin_scope_value(tmp_path, values, "region")
            self.assertTrue(region_result["pinned"])
            self.assertTrue(oktmo_scope_pin_status(tmp_path, values, "region")["pinned"])

            municipalities = [
                option
                for option in oktmo_municipality_scope_options(tmp_path, values, {})
                if option.get("value")
            ]
            self.assertTrue(municipalities)
            values["oktmo_scope_municipality"] = municipalities[0]["value"]

            municipality_result = oktmo_pin_scope_value(tmp_path, values, "municipality")
            self.assertTrue(municipality_result["pinned"])
            self.assertTrue(oktmo_scope_pin_status(tmp_path, values, "municipality")["pinned"])
            self.assertEqual(load_pins(tmp_path, "municipality")[0]["scope"], "71")

            decorated = oktmo_municipality_scope_options(tmp_path, values, {})
            pinned_values = {
                str(option.get("value"))
                for option in decorated
                if str(option.get("pinned", "")).lower() == "true"
            }
            self.assertIn(str(values["oktmo_scope_municipality"]), pinned_values)

            exported = oktmo_export_pin_file(tmp_path)
            export_path = Path(str(exported["path"]))
            self.assertTrue(export_path.exists())
            self.assertIn("output", export_path.parts)
            self.assertIn("oktmo_pins", export_path.parts)
            payload = json.loads(export_path.read_text(encoding="utf-8"))
            self.assertEqual(payload["files"]["regions"], "config\\oktmo_region_pins.json")
            self.assertNotIn(str(tmp_path), json.dumps(payload.get("files", {}), ensure_ascii=False))

            cleared = oktmo_clear_pin_file(tmp_path)
            self.assertEqual(cleared["regions"], 0)
            self.assertEqual(cleared["municipalities"], 0)
            self.assertEqual(load_pins(tmp_path, "region"), [])
            self.assertEqual(load_pins(tmp_path, "municipality"), [])

            restored = oktmo_import_pin_file(tmp_path)
            self.assertEqual(restored["regions"], 1)
            self.assertEqual(restored["municipalities"], 1)
            self.assertTrue(load_pins(tmp_path, "region"))
            self.assertTrue(load_pins(tmp_path, "municipality"))

    def test_oktmo_reset_project_keys_clears_selection_but_preserves_pins(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            tmp_path = Path(tmp)
            data_root = tmp_path / "data"
            write_oktmo_fixture(data_root)
            values = {"oktmo_data_dir": str(data_root), "oktmo_scope_region": "71"}
            oktmo_pin_scope_value(tmp_path, values, "region")
            write_current_keys(tmp_path, ["Тюменский", "Патрушева"])

            result = oktmo_reset_project_keys(tmp_path)

            self.assertEqual(load_current_keys(tmp_path), [])
            self.assertTrue(load_pins(tmp_path, "region"))
            self.assertTrue(result["pins_preserved"])
            updates = result["field_updates"]
            self.assertIsNone(updates["oktmo_scope_region"])
            self.assertIsNone(updates["oktmo_scope_municipality"])
            self.assertEqual(updates["oktmo_candidate_options"], [])
            self.assertEqual(updates["oktmo_candidate_preview"], "")

    def test_oktmo_select_pin_slot_uses_pinned_indexes(self) -> None:
        from system_core.ui_nicegui import app as gui_app  # noqa: WPS433

        fake_select = SimpleNamespace(_values=["a", "b", "c"], slots={})

        def add_slot(name: str, template: str) -> None:
            fake_select.slots[name] = SimpleNamespace(template=template)

        fake_select.add_slot = add_slot
        field = {
            "id": "fake_oktmo_scope",
            "options": [
                {"value": "a", "label": "A", "pinned": "true"},
                {"value": "b", "label": "B"},
                {"value": "c", "label": "C", "pinned": "true"},
            ],
        }

        gui_app.decorate_oktmo_select_option_props(fake_select, field)

        template = fake_select.slots["option"].template
        self.assertIn("[0, 2].includes(props.opt.value)", template)
        self.assertIn("audion-select-option-pinned", template)
        self.assertIn("push_pin", template)

    def test_workspace_path_history_remembers_unpinned_and_clear_keeps_pins(self) -> None:
        from system_core.ui_nicegui.workbench import WorkbenchConfig, WorkbenchHistory  # noqa: WPS433

        with tempfile.TemporaryDirectory() as tmp:
            tmp_path = Path(tmp)
            history_file = tmp_path / "path_history.json"
            history = WorkbenchHistory(
                WorkbenchConfig(
                    root=tmp_path,
                    input_path=tmp_path / "input",
                    output_path=tmp_path / "output",
                    history_path=history_file,
                )
            )
            pinned_source = tmp_path / "source_pinned"
            recent_source = tmp_path / "source_recent"
            target_path = tmp_path / "target_recent"

            history.remember("source", str(pinned_source))
            history.remember("source", str(recent_source))
            history.remember("target", str(target_path))
            history.set_pinned("source", str(pinned_source), True, required_message="path required")

            data = json.loads(history_file.read_text(encoding="utf-8"))
            sources = {item["path"]: item for item in data["sources"]}
            self.assertIn(str(recent_source), sources)
            self.assertFalse(sources[str(recent_source)]["pinned"])
            self.assertTrue(sources[str(pinned_source)]["pinned"])
            self.assertEqual(data["targets"][0]["path"], str(target_path))

            result = history.clear_cache_keep_pins()

            data = json.loads(history_file.read_text(encoding="utf-8"))
            self.assertEqual(result["removed_sources"], 1)
            self.assertEqual(result["removed_targets"], 1)
            self.assertEqual(result["kept_pins"], 1)
            self.assertEqual([item["path"] for item in data["sources"]], [str(pinned_source)])
            self.assertTrue(data["sources"][0]["pinned"])
            self.assertEqual(data["targets"], [])

    def test_workspace_startup_always_returns_project_input_output(self) -> None:
        from system_core.ui_nicegui import app as gui_app  # noqa: WPS433

        with tempfile.TemporaryDirectory() as tmp:
            history_file = Path(tmp) / "path_history.json"
            configured = Path(tmp) / "configured_source"
            pinned = Path(tmp) / "pinned_source"
            default_path = Path(tmp) / "project_input"
            history_file.write_text(
                json.dumps(
                    {
                        "sources": [
                            {"path": str(configured), "count": 1, "last_used": "2026-01-01T00:00:00", "pinned": True},
                            {"path": str(pinned), "count": 1, "last_used": "2026-01-02T00:00:00", "pinned": True},
                        ],
                        "targets": [],
                    },
                    ensure_ascii=False,
                ),
                encoding="utf-8",
            )
            original_history_file = gui_app._workspace_history_file
            gui_app._workspace_history_file = lambda: history_file
            try:
                self.assertEqual(
                    gui_app._startup_workspace_path("source", str(configured), str(pinned), default_path),
                    str(default_path),
                )
                self.assertEqual(gui_app._workspace_setting_for_disk("source", str(configured), default_path), "")
            finally:
                gui_app._workspace_history_file = original_history_file

    def test_workspace_path_select_pin_slot_uses_compact_pinned_indexes(self) -> None:
        from system_core.ui_nicegui.workbench import (  # noqa: WPS433
            WorkbenchAdapter,
            WorkbenchConfig,
            WorkbenchRenderer,
        )

        with tempfile.TemporaryDirectory() as tmp:
            tmp_path = Path(tmp)
            history_file = tmp_path / "path_history.json"
            pinned_source = tmp_path / "source_pinned"
            recent_source = tmp_path / "source_recent"
            history_file.write_text(
                json.dumps(
                    {
                        "sources": [
                            {"path": str(recent_source), "count": 8, "last_used": "2026-01-02T00:00:00", "pinned": False},
                            {"path": str(pinned_source), "count": 1, "last_used": "2026-01-01T00:00:00", "pinned": True},
                        ],
                        "targets": [],
                    },
                    ensure_ascii=False,
                ),
                encoding="utf-8",
            )
            adapter = WorkbenchAdapter(
                config=WorkbenchConfig(
                    root=tmp_path,
                    input_path=tmp_path / "input",
                    output_path=tmp_path / "output",
                    history_path=history_file,
                ),
                current_path_callback=lambda role: recent_source if role == "source" else tmp_path / "output",
                save_path_callback=lambda role, value: None,
                language_callback=lambda: "ru",
                translate_callback=lambda key, **kwargs: key,
                log_callback=lambda message: None,
                notify_callback=lambda message, level: None,
                reload_callback=lambda delay_ms: None,
                busy_callback=lambda: False,
                feedback_callback=lambda: {},
                set_feedback_callback=lambda role, action: None,
                clear_feedback_callback=lambda: None,
            )
            renderer = WorkbenchRenderer(
                adapter=adapter,
                handlers=SimpleNamespace(),
                display_path_callback=str,
            )
            options = renderer.path_select_options("source", recent_source)
            self.assertIn(str(pinned_source), options)
            self.assertIn(str(recent_source), options)
            self.assertNotIn("PIN |", "".join(options.values()))

            fake_select = SimpleNamespace(_values=list(options.keys()), slots={})

            def add_slot(name: str, template: str) -> None:
                fake_select.slots[name] = SimpleNamespace(template=template)

            fake_select.add_slot = add_slot

            renderer.decorate_path_select_options(fake_select, "source", recent_source)

            template = fake_select.slots["option"].template
            self.assertIn("[0].includes(Number(props.opt.value))", template)
            self.assertIn("audion-path-option-pinned", template)
            self.assertIn("audion-path-option-pin-cell", template)
            self.assertIn("push_pin", template)

    def test_oktmo_scope_extra_keys_are_weighted_and_context_aware(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            data_root = Path(tmp) / "data"
            write_oktmo_fixture(data_root)

            profile = oktmo_scope_key_profile(
                data_root,
                subject_ter_hint="71",
                municipality_hint="Тюменский муниципальный район",
            )

            keys = profile["weighted_extra_keys"]
            subject_keys = [item for item in keys if item["source"] == "subject"]
            self.assertTrue(subject_keys)
            self.assertTrue(all(item["role"] == "constraint" and item["weight"] == 0 for item in subject_keys))
            self.assertNotIn("Тюменской области", profile["extra_keys"])

            municipality_keys = [item for item in keys if item["source"] == "municipality"]
            self.assertTrue(any(item["role"] == "medium" and item["weight"] == 45 for item in municipality_keys))
            self.assertTrue(any(item["value"].lower() == "тюменский" for item in municipality_keys))

            locality_keys = [item for item in keys if item["source"] == "locality"]
            patrusheva_keys = [item for item in locality_keys if item["code"] == "71001001001"]
            self.assertTrue(patrusheva_keys)
            self.assertTrue(all(item["role"] == "strong" and item["weight"] == 90 for item in patrusheva_keys))
            mirny_keys = [item for item in locality_keys if item["code"] == "71001001002"]
            central_keys = [item for item in locality_keys if item["code"] == "71001001003"]
            klyuchi_keys = [item for item in locality_keys if item["code"] == "71001001004"]
            self.assertTrue(mirny_keys)
            self.assertTrue(central_keys)
            self.assertTrue(klyuchi_keys)
            self.assertTrue(
                all(item["context_required"] and item["role"] == "contextual" and item["weight"] == 20 for item in mirny_keys)
            )
            self.assertTrue(
                all(item["context_required"] and item["role"] == "contextual" and item["weight"] == 20 for item in central_keys)
            )
            self.assertTrue(
                all(item["context_required"] and item["role"] == "contextual" and item["weight"] == 20 for item in klyuchi_keys)
            )
            klyuchi_values = {item["value"] for item in klyuchi_keys}
            self.assertIn("Ключи", klyuchi_values)
            self.assertIn("Ключей", klyuchi_values)
            self.assertNotIn("село Ключи", klyuchi_values)
            self.assertNotIn("ключ", klyuchi_values)

    def test_matcher_scope_scoring_blocks_conflicting_oktmo_localities(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            data_root = Path(tmp) / "data"
            write_oktmo_fixture(data_root)
            profile = oktmo_scope_key_profile(
                data_root,
                subject_ter_hint="71",
                municipality_hint="Тюменский муниципальный район",
            )
            normalizer = AddressNormalizer(
                data_dir=data_root,
                subject_ter_hint="71",
                municipality_hint="Тюменский муниципальный район",
            )

            gt = normalizer.normalize("с. Ключи, ул. Ленина, д. 7")
            same_place = normalizer.normalize("село Ключи, улица Ленина, дом 7")
            other_place = normalizer.normalize("село Мирный, улица Ленина, дом 7")

            gt_scope = _alignment_scope_match(gt.raw, gt, profile)
            same_scope = _alignment_scope_match(same_place.raw, same_place, profile)
            other_scope = _alignment_scope_match(other_place.raw, other_place, profile)

            self.assertFalse(_alignment_scope_pair_blocked(gt, same_place))
            self.assertGreater(_alignment_scope_pair_score(gt_scope, same_scope, gt, same_place), 0)
            self.assertTrue(_alignment_scope_pair_blocked(gt, other_place))
            self.assertLess(_alignment_scope_pair_score(gt_scope, other_scope, gt, other_place), 0)

    def test_oktmo_canonical_locality_and_region_slots_win_over_dirty_type(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            data_root = Path(tmp) / "data"
            data_root.mkdir()
            write_region_fixture(data_root)
            write_altay_klyuchi_fixture(data_root)
            normalizer = AddressNormalizer(data_dir=data_root, subject_ter_hint="01")

            record = normalizer.normalize("Алтайский край, д Ключи, ул Ленина, д. 10")

            self.assertEqual(record.oktmo_name, "п Ключи")
            self.assertEqual(record.parts.locality, "п. Ключи")
            self.assertEqual(record.subject, "Алтайский край")
            self.assertEqual(record.federal_district, "Сибирский федеральный округ")
            self.assertIn("п. Ключи", record.normalized or "")

    def test_microdistrict_is_separate_slot_before_street(self) -> None:
        record = AddressNormalizer(use_oktmo=False).normalize(
            "625019, Тюменская обл, г. Тюмень, мкр. Лесной, ул. Центральная, д. 1"
        )

        self.assertEqual(record.parts.microdistrict, "мкр. Лесной")
        self.assertIsNone(record.parts.territory)
        self.assertEqual(record.parts.street, "ул. Центральная")
        self.assertIn("мкр. Лесной, ул. Центральная", record.normalized or "")

    def test_reference_can_omit_microdistrict_from_address_when_street_exists(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            tmp_path = Path(tmp)
            input_dir = tmp_path / "input"
            output_dir = tmp_path / "output"
            input_dir.mkdir()
            output_dir.mkdir()
            pd.DataFrame(
                [
                    ["Тюменская обл, г. Тюмень, мкр. Лесной, ул. Центральная, д. 1"],
                    ["Тюменская обл, г. Тюмень, ул. Центральная, д. 1"],
                ]
            ).to_excel(input_dir / "dirty.xlsx", index=False, header=False)

            result = build_reference_workbook(
                input_dir=input_dir,
                output_file=output_dir / "AddressReference_microdistrict.xlsx",
                normalizer=AddressNormalizer(use_oktmo=False),
                options=ReferenceOptions(enabled_slots=DEFAULT_ENABLED_SLOTS),
            )

            reference = pd.read_excel(result.output_file)
            self.assertEqual(len(reference), 1)
            self.assertIn("Microdistrict", reference.columns)
            self.assertEqual(reference.iloc[0]["Microdistrict"], "мкр. Лесной")
            self.assertNotIn("мкр. Лесной", reference.iloc[0]["Normalized_Address"])
            self.assertEqual(int(reference.iloc[0]["Source_Count"]), 2)

    def test_reference_generation_uses_enabled_slots(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            tmp_path = Path(tmp)
            data_root = tmp_path / "data"
            input_dir = tmp_path / "input"
            output_dir = tmp_path / "output"
            input_dir.mkdir()
            output_dir.mkdir()
            write_oktmo_fixture(data_root)
            pd.DataFrame(
                [
                    ["625000, деревня Патрушева, улица Ленина, д. 10"],
                    ["д Патрушева ул Ленина 10"],
                    ["не адрес"],
                ]
            ).to_excel(input_dir / "dirty.xlsx", index=False, header=False)

            result = build_reference_workbook(
                input_dir=input_dir,
                output_file=output_dir / "AddressReference_test.xlsx",
                normalizer=AddressNormalizer(data_dir=data_root, subject_ter_hint="71"),
                options=ReferenceOptions(enabled_slots=DEFAULT_ENABLED_SLOTS),
                staging_file=output_dir / "AddressReference_test.stage.xlsx",
                report_file=output_dir / "AddressReference_test.summary.json",
            )

            reference = pd.read_excel(result.output_file)
            self.assertEqual(len(reference), 1)
            self.assertIn("Postal_Index", reference.columns)
            self.assertIn("Microdistrict", reference.columns)
            self.assertEqual(str(reference.iloc[0]["Postal_Index"]), "625000")
            self.assertTrue(str(reference.iloc[0]["Normalized_Address"]).startswith("625000, "))
            self.assertEqual(int(reference.iloc[0]["Source_Count"]), 2)
            self.assertTrue(result.staging_file.exists())
            staging = load_workbook(result.staging_file, read_only=True, data_only=True)
            try:
                self.assertIn("Reference_Candidates", staging.sheetnames)
                self.assertIn("Reference_Final", staging.sheetnames)
                candidates_sheet = staging["Reference_Candidates"]
                headers = [candidates_sheet.cell(row=1, column=column).value for column in range(1, candidates_sheet.max_column + 1)]
                status_column = headers.index("Reference_Status") + 1
                statuses = {
                    str(candidates_sheet.cell(row=row, column=status_column).value)
                    for row in range(2, candidates_sheet.max_row + 1)
                }
                self.assertEqual(statuses, {"ready", "duplicate"})
                self.assertEqual(staging["Reference_Final"].max_row, 2)
            finally:
                staging.close()

            no_index_output = output_dir / "AddressReference_no_index.xlsx"
            no_index_slots = tuple(slot for slot in DEFAULT_ENABLED_SLOTS if slot != "postal_index")
            build_reference_workbook(
                input_dir=input_dir,
                output_file=no_index_output,
                normalizer=AddressNormalizer(data_dir=data_root, subject_ter_hint="71"),
                options=ReferenceOptions(enabled_slots=no_index_slots),
            )
            reference_without_index = pd.read_excel(no_index_output)
            self.assertNotIn("Postal_Index", reference_without_index.columns)
            self.assertFalse(str(reference_without_index.iloc[0]["Normalized_Address"]).startswith("625000"))
            summary = json.loads(result.report_file.read_text(encoding="utf-8"))
            self.assertEqual(summary["mode"], "address-reference-normalization")
            self.assertEqual(summary["staging_file"], str(result.staging_file))
            self.assertEqual(summary["normalizer"]["oktmo_key_profile"]["subject_ter"], "71")
            self.assertEqual(summary["options"]["enabled_slots"], list(DEFAULT_ENABLED_SLOTS))
            self.assertEqual(summary["address_diagnostics"]["total"], 1)
            self.assertEqual(summary["address_diagnostics"]["issue_rows"], 0)
            self.assertEqual(summary["slot_fill"]["rows_total"], 1)
            self.assertIn("postal_index", summary["slot_fill"]["selected_slots"])
            self.assertIn("locality", summary["slot_fill"]["selected_slots"])
            self.assertIn("street", summary["slot_fill"]["selected_slots"])
            self.assertIn("house", summary["slot_fill"]["selected_slots"])
            self.assertEqual(summary["slot_fill"]["slots"]["postal_index"]["filled"], 1)
            self.assertEqual(summary["slot_fill"]["slots"]["locality"]["filled"], 1)
            self.assertEqual(summary["slot_fill"]["slots"]["street"]["filled"], 1)
            self.assertEqual(summary["slot_fill"]["slots"]["house"]["filled"], 1)

    def test_reference_generation_can_write_compact_clean_workbook(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            tmp_path = Path(tmp)
            input_dir = tmp_path / "input"
            output_dir = tmp_path / "output"
            input_dir.mkdir()
            output_dir.mkdir()
            pd.DataFrame([["г. Тюмень, ул. Ленина, д. 10"]]).to_excel(
                input_dir / "dirty.xlsx",
                index=False,
                header=False,
            )

            result = build_reference_workbook(
                input_dir=input_dir,
                output_file=output_dir / "AddressReference_clean.xlsx",
                normalizer=AddressNormalizer(use_oktmo=False),
                options=ReferenceOptions(
                    enabled_slots=("locality", "street", "house", "match_key"),
                    slot_order=("locality", "street", "house"),
                    clean_output_workbook=True,
                ),
                report_file=output_dir / "AddressReference_clean.summary.json",
            )

            workbook = load_workbook(result.output_file)
            try:
                sheet = workbook.worksheets[0]
                self.assertEqual(sheet.freeze_panes, "A2")
                self.assertEqual(sheet.auto_filter.ref, sheet.dimensions)
                self.assertEqual(float(sheet.cell(row=1, column=1).font.sz), 9.0)
                self.assertEqual(float(sheet.cell(row=2, column=1).font.sz), 9.0)
                self.assertEqual(sheet.cell(row=1, column=1).value, "Normalized_Address")
                self.assertIn("ул. Ленина", str(sheet.cell(row=2, column=1).value))
            finally:
                workbook.close()

            summary = json.loads(result.report_file.read_text(encoding="utf-8"))
            self.assertTrue(summary["options"]["clean_output_workbook"])

    def test_reference_generation_respects_slot_order(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            tmp_path = Path(tmp)
            data_root = tmp_path / "data"
            input_dir = tmp_path / "input"
            output_dir = tmp_path / "output"
            input_dir.mkdir()
            output_dir.mkdir()
            write_oktmo_fixture(data_root)
            pd.DataFrame([["625000, деревня Патрушева, улица Ленина, д. 10"]]).to_excel(
                input_dir / "dirty.xlsx",
                index=False,
                header=False,
            )

            result = build_reference_workbook(
                input_dir=input_dir,
                output_file=output_dir / "AddressReference_order.xlsx",
                normalizer=AddressNormalizer(data_dir=data_root, subject_ter_hint="71"),
                options=ReferenceOptions(
                    enabled_slots=("postal_index", "locality", "street", "house", "match_key"),
                    slot_order=("locality", "street", "house", "postal_index"),
                ),
            )

            reference = pd.read_excel(result.output_file)
            self.assertEqual(
                reference.iloc[0]["Normalized_Address"],
                "д. Патрушева, ул. Ленина, д. 10, 625000",
            )

    def test_existing_reference_auto_detects_named_reference_column(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            input_dir = Path(tmp)
            workbook = input_dir / "AddressReference_existing.xlsx"
            pd.DataFrame(
                [
                    ["id", "Normalized_Address", "OKTMO_Code"],
                    [1, "625000, деревня Патрушева, улица Ленина, д. 10", "71001001001"],
                ]
            ).to_excel(workbook, index=False, header=False)

            candidates, issues = extract_reference_candidates(workbook)

            self.assertFalse(issues)
            self.assertEqual([candidate.value for candidate in candidates], ["625000, деревня Патрушева, улица Ленина, д. 10"])

    def test_existing_reference_manual_header_name_is_supported(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            input_dir = Path(tmp)
            workbook = input_dir / "AddressReference_existing.xlsx"
            pd.DataFrame(
                [
                    ["ID", "Эталонный адрес", "OKTMO_Code"],
                    [1, "деревня Патрушева, улица Ленина, д. 10", "71001001001"],
                ]
            ).to_excel(workbook, index=False, header=False)

            candidates, issues = extract_reference_candidates(workbook, reference_columns="Эталонный адрес")

            self.assertFalse(issues)
            self.assertEqual([candidate.value for candidate in candidates], ["деревня Патрушева, улица Ленина, д. 10"])

    def test_reference_generation_reads_text_markdown_and_odt_sources(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            tmp_path = Path(tmp)
            data_root = tmp_path / "data"
            input_dir = tmp_path / "input"
            output_dir = tmp_path / "output"
            input_dir.mkdir()
            output_dir.mkdir()
            write_oktmo_fixture(data_root)

            address = "625000, деревня Патрушева, улица Ленина, д. 10"
            (input_dir / "source.txt").write_text(f"{address}\n", encoding="utf-8")
            (input_dir / "source.md").write_text(
                "\n".join(["| Адрес |", "|---|", f"| {address} |"]),
                encoding="utf-8",
            )
            write_minimal_odt(input_dir / "source.odt", address)

            result = build_reference_workbook(
                input_dir=input_dir,
                output_file=output_dir / "AddressReference_sources.xlsx",
                normalizer=AddressNormalizer(data_dir=data_root, subject_ter_hint="71"),
                options=ReferenceOptions(enabled_slots=DEFAULT_ENABLED_SLOTS),
            )

            reference = pd.read_excel(result.output_file)
            self.assertEqual(len(reference), 1)
            self.assertGreaterEqual(int(reference.iloc[0]["Source_Count"]), 3)
            sources = str(reference.iloc[0]["Sources"])
            self.assertIn("source.txt", sources)
            self.assertIn("source.md", sources)
            self.assertIn("source.odt", sources)
            self.assertFalse(result.issues)

    def test_reference_generation_reads_supported_source_matrix(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            tmp_path = Path(tmp)
            data_root = tmp_path / "data"
            input_dir = tmp_path / "input"
            output_dir = tmp_path / "output"
            input_dir.mkdir()
            output_dir.mkdir()
            write_oktmo_fixture(data_root)

            address = "625000, деревня Патрушева, улица Ленина, д. 10"
            pd.DataFrame([["Местоположение"], [address]]).to_excel(input_dir / "source.xlsx", index=False, header=False)
            (input_dir / "source.txt").write_text(f"{address}\n", encoding="utf-8")
            (input_dir / "source.csv").write_text(f"id;Местоположение\n1;{address}\n", encoding="utf-8")
            (input_dir / "source.md").write_text(
                "\n".join(["| Адрес |", "|---|", f"| {address} |"]),
                encoding="utf-8",
            )
            write_minimal_docx(input_dir / "source.docx", address)
            write_minimal_odt(input_dir / "source.odt", address)
            expected_sources = {
                "source.xlsx",
                "source.txt",
                "source.csv",
                "source.md",
                "source.docx",
                "source.odt",
            }
            if importlib.util.find_spec("pymupdf") is not None or importlib.util.find_spec("fitz") is not None:
                write_searchable_pdf(input_dir / "source.pdf", address)
                expected_sources.add("source.pdf")

            result = build_reference_workbook(
                input_dir=input_dir,
                output_file=output_dir / "AddressReference_source_matrix.xlsx",
                normalizer=AddressNormalizer(data_dir=data_root, subject_ter_hint="71"),
                options=ReferenceOptions(enabled_slots=DEFAULT_ENABLED_SLOTS),
            )

            reference = pd.read_excel(result.output_file)
            self.assertEqual(len(reference), 1)
            self.assertEqual(str(reference.iloc[0]["Postal_Index"]), "625000")
            self.assertEqual(reference.iloc[0]["Locality"], "д. Патрушева")
            self.assertEqual(reference.iloc[0]["Street"], "ул. Ленина")
            self.assertEqual(reference.iloc[0]["House"], "д. 10")
            self.assertGreaterEqual(int(reference.iloc[0]["Source_Count"]), len(expected_sources))
            sources = str(reference.iloc[0]["Sources"])
            for source_name in expected_sources:
                self.assertIn(source_name, sources)
            self.assertFalse(result.issues)

    def test_reference_processing_sorts_by_slots_and_deconstructs(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            tmp_path = Path(tmp)
            input_file = tmp_path / "AddressReference_existing.xlsx"
            output_file = tmp_path / "AddressReference_processed.xlsx"
            pd.DataFrame(
                [
                    ["ID", "Normalized_Address", "Comment"],
                    [1, "г. Тюмень, ул. Бета, д. 10", "b10"],
                    [2, "г. Тюмень, ул. Альфа, д. 10", "a10"],
                    [3, "г. Тюмень, ул. Альфа, д. 2", "a2"],
                ]
            ).to_excel(input_file, index=False, header=False)

            result = process_reference_workbook(
                input_file=input_file,
                output_file=output_file,
                normalizer=AddressNormalizer(use_oktmo=False),
                options=ReferenceProcessingOptions(
                    action="sort_deconstruct",
                    reference_columns="Normalized_Address",
                    enabled_slots=("street", "house"),
                    slot_order=("street", "house"),
                    remove_source_column=True,
                ),
                report_file=tmp_path / "AddressReference_processed.summary.json",
            )

            processed = pd.read_excel(result.output_file)
            self.assertEqual(list(processed.columns), ["ID", "Street", "House", "Comment"])
            self.assertEqual(list(processed["Comment"]), ["a2", "a10", "b10"])
            self.assertEqual(list(processed["Street"]), ["ул. Альфа", "ул. Альфа", "ул. Бета"])
            self.assertEqual(list(processed["House"]), ["д. 2", "д. 10", "д. 10"])
            self.assertEqual(result.sorted_rows, 3)
            self.assertEqual(result.deconstructed_rows, 3)
            summary = json.loads(result.report_file.read_text(encoding="utf-8"))
            self.assertEqual(summary["mode"], "address-reference-processing")
            self.assertEqual(summary["action"], "sort_deconstruct")
            self.assertEqual(summary["normalizer"]["oktmo_key_profile"]["morphology"]["skipped"], "no_scope")
            self.assertEqual(summary["address_diagnostics"]["total"], 3)
            self.assertEqual(summary["address_diagnostics"]["issue_rows"], 0)
            self.assertEqual(summary["slot_fill"]["records_total"], 3)
            self.assertEqual(summary["slot_fill"]["empty_rows"], 0)
            self.assertEqual(summary["slot_fill"]["selected_slots"], ["street", "house"])
            self.assertEqual(summary["slot_fill"]["complete_rows"], 3)
            self.assertEqual(summary["slot_fill"]["incomplete_rows"], 0)
            self.assertEqual(summary["slot_fill"]["slots"]["street"]["filled"], 3)
            self.assertEqual(summary["slot_fill"]["slots"]["street"]["fill_rate"], 1.0)
            self.assertEqual(summary["slot_fill"]["slots"]["house"]["filled"], 3)
            self.assertEqual(summary["slot_fill"]["slots"]["house"]["fill_rate"], 1.0)

            short_output = tmp_path / "AddressReference_processed_short.xlsx"
            short_result = process_reference_workbook(
                input_file=input_file,
                output_file=short_output,
                normalizer=AddressNormalizer(use_oktmo=False),
                options=ReferenceProcessingOptions(
                    action="deconstruct",
                    reference_columns="Normalized_Address",
                    enabled_slots=("microdistrict", "house"),
                    slot_order=("microdistrict", "house"),
                    remove_source_column=True,
                    slot_header_mode="short",
                ),
            )
            short_processed = pd.read_excel(short_result.output_file)
            self.assertEqual(list(short_processed.columns), ["ID", "Mkr_Qtr", "House", "Comment"])

            no_header_input = tmp_path / "AddressReference_without_header.xlsx"
            no_header_output = tmp_path / "AddressReference_without_header_processed.xlsx"
            pd.DataFrame(
                [
                    [1, "г. Тюмень, ул. Бета, д. 10", "b10"],
                    [2, "г. Тюмень, ул. Альфа, д. 2", "a2"],
                ]
            ).to_excel(no_header_input, index=False, header=False, sheet_name="Адреса")
            no_header_result = process_reference_workbook(
                input_file=no_header_input,
                output_file=no_header_output,
                normalizer=AddressNormalizer(use_oktmo=False),
                options=ReferenceProcessingOptions(
                    action="sort_deconstruct",
                    reference_columns="B",
                    enabled_slots=("street", "house"),
                    slot_order=("street", "house"),
                    remove_source_column=True,
                    slot_header_mode="short",
                ),
            )

            workbook = load_workbook(no_header_result.output_file, read_only=True, data_only=True)
            try:
                sheet = workbook.worksheets[0]
                self.assertEqual(no_header_result.sheet_name, "Адреса")
                self.assertEqual(sheet.max_row, 3)
                self.assertEqual(sheet.cell(row=1, column=2).value, "Street")
                self.assertEqual(sheet.cell(row=1, column=3).value, "House")
                self.assertEqual(sheet.cell(row=2, column=2).value, "ул. Альфа")
                self.assertEqual(sheet.cell(row=2, column=3).value, "д. 2")
                self.assertEqual(sheet.cell(row=2, column=4).value, "a2")
                self.assertEqual(sheet.cell(row=3, column=2).value, "ул. Бета")
                self.assertEqual(sheet.cell(row=3, column=3).value, "д. 10")
                self.assertEqual(sheet.cell(row=3, column=4).value, "b10")
            finally:
                workbook.close()

    def test_address_collection_creates_table_and_sanitizes_incomplete(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            tmp_path = Path(tmp)
            input_dir = tmp_path / "input"
            collection_dir = tmp_path / "address_collection"
            input_dir.mkdir()
            (input_dir / "source.txt").write_text(
                "\n".join(
                    [
                        "г. Тюмень, ул. Альфа, д. 2",
                        "Тюмень, улица Альфа, дом 2",
                        "г. Тюмень, ул. Бета",
                        "д. 99",
                    ]
                ),
                encoding="utf-8",
            )
            output_file = collection_dir / "AddressCollection_new.xlsx"

            result = assemble_address_workbook(
                input_dir=input_dir,
                output_file=output_file,
                normalizer=AddressNormalizer(use_oktmo=False),
                options=AddressAssemblyOptions(
                    action="generate",
                    target_dir=collection_dir,
                    target_column="B",
                    enabled_slots=("locality", "street", "house", "match_key"),
                    slot_order=("locality", "street", "house"),
                ),
            )

            workbook = load_workbook(result.output_file, read_only=True, data_only=True)
            try:
                sheet = workbook.worksheets[0]
                self.assertEqual(result.appended_rows, 1)
                self.assertEqual(result.target_column, "B")
                self.assertEqual(result.target_start_row, 1)
                self.assertIn("ул. Альфа", str(sheet.cell(row=1, column=2).value))
                self.assertIn("д. 2", str(sheet.cell(row=1, column=2).value))
                self.assertNotIn("Rejected", workbook.sheetnames)
            finally:
                workbook.close()

            staging = load_workbook(result.staging_file, read_only=True, data_only=True)
            try:
                self.assertIn("Slot_Candidates", staging.sheetnames)
                self.assertIn("Sanitized", staging.sheetnames)
                self.assertIn("Rejected", staging.sheetnames)
            finally:
                staging.close()

            summary = json.loads(result.report_file.read_text(encoding="utf-8"))
            self.assertEqual(summary["normalizer"]["use_oktmo"], False)
            self.assertEqual(summary["options"]["enabled_slots"], ["locality", "street", "house", "match_key"])
            self.assertEqual(summary["options"]["slot_order"], ["locality", "street", "house"])
            self.assertEqual(summary["options"]["sanitize_incomplete"], True)
            self.assertEqual(summary["address_diagnostics"]["normalized"]["total"], 1)
            self.assertGreaterEqual(summary["address_diagnostics"]["rejected"]["issue_rows"], 1)
            self.assertEqual(summary["slot_fill"]["rows_total"], 1)
            self.assertEqual(summary["slot_fill"]["selected_slots"], ["locality", "street", "house", "match_key"])
            self.assertEqual(summary["slot_fill"]["complete_rows"], 1)
            self.assertEqual(summary["slot_fill"]["incomplete_rows"], 0)
            self.assertEqual(summary["slot_fill"]["slots"]["street"]["filled"], 1)
            self.assertEqual(summary["slot_fill"]["slots"]["house"]["filled"], 1)

    def test_address_collection_appends_after_last_meaningful_cell(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            tmp_path = Path(tmp)
            input_dir = tmp_path / "input"
            collection_dir = tmp_path / "address_collection"
            input_dir.mkdir()
            collection_dir.mkdir()
            (input_dir / "source.txt").write_text(
                "г. Тюмень, ул. Бета, д. 10\nг. Тюмень, ул. Альфа, д. 2",
                encoding="utf-8",
            )
            ready_file = collection_dir / "ready.xlsx"
            pd.DataFrame({"Адрес": ["г. Тюмень, ул. Альфа, д. 2"]}).to_excel(ready_file, index=False)
            output_file = collection_dir / "ready_added.xlsx"

            result = assemble_address_workbook(
                input_dir=input_dir,
                output_file=output_file,
                normalizer=AddressNormalizer(use_oktmo=False),
                options=AddressAssemblyOptions(
                    action="update",
                    target_dir=collection_dir,
                    target_file=ready_file,
                    target_column="Адрес",
                    enabled_slots=("locality", "street", "house", "match_key"),
                    slot_order=("locality", "street", "house"),
                ),
            )

            workbook = load_workbook(result.output_file, read_only=True, data_only=True)
            try:
                sheet = workbook.worksheets[0]
                self.assertEqual(result.appended_rows, 1)
                self.assertEqual(result.target_column, "A")
                self.assertEqual(result.target_start_row, 3)
                self.assertEqual(sheet.cell(row=1, column=1).value, "Адрес")
                self.assertIn("ул. Альфа", str(sheet.cell(row=2, column=1).value))
                self.assertIn("ул. Бета", str(sheet.cell(row=3, column=1).value))
            finally:
                workbook.close()

            staging = load_workbook(result.staging_file, read_only=True, data_only=True)
            try:
                sheet = staging["Slot_Candidates"]
                statuses = [str(sheet.cell(row=row_index, column=1).value) for row_index in range(2, sheet.max_row + 1)]
                self.assertIn("existing_reference", statuses)
                self.assertIn("ready", statuses)
            finally:
                staging.close()

    def test_address_collection_scores_weighted_oktmo_scope_keys(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            tmp_path = Path(tmp)
            data_root = tmp_path / "data"
            input_dir = tmp_path / "input"
            collection_dir = tmp_path / "address_collection"
            input_dir.mkdir()
            write_oktmo_fixture(data_root)
            (input_dir / "source.txt").write_text(
                "\n".join(
                    [
                        "с. Ключи, ул. Ленина, д. 7",
                        "Ключи",
                    ]
                ),
                encoding="utf-8",
            )

            result = assemble_address_workbook(
                input_dir=input_dir,
                output_file=collection_dir / "AddressCollection_scope.xlsx",
                normalizer=AddressNormalizer(
                    data_dir=data_root,
                    subject_ter_hint="71",
                    municipality_hint="Тюменский муниципальный район",
                ),
                options=AddressAssemblyOptions(
                    action="generate",
                    target_dir=collection_dir,
                    target_column="A",
                    enabled_slots=("locality", "street", "house", "match_key"),
                    slot_order=("locality", "street", "house"),
                ),
            )

            self.assertEqual(result.appended_rows, 1)
            staging = load_workbook(result.staging_file, read_only=True, data_only=True)
            try:
                sheet = staging["Slot_Candidates"]
                headers = [sheet.cell(row=1, column=column).value for column in range(1, sheet.max_column + 1)]
                match_column = headers.index("OKTMO_Scope_Match") + 1
                weight_column = headers.index("OKTMO_Scope_Weight") + 1
                valid_column = headers.index("OKTMO_Scope_Valid") + 1
                context_column = headers.index("OKTMO_Scope_Context") + 1
                match_text = str(sheet.cell(row=2, column=match_column).value)
                self.assertIn("medium", match_text)
                self.assertIn("Ключ", match_text)
                self.assertEqual(float(sheet.cell(row=2, column=weight_column).value), 45.0)
                self.assertTrue(sheet.cell(row=2, column=valid_column).value)
                self.assertTrue(sheet.cell(row=2, column=context_column).value)
            finally:
                staging.close()

            summary = json.loads(result.report_file.read_text(encoding="utf-8"))
            self.assertGreaterEqual(summary["oktmo_scope_matches"]["valid"], 1)

            profile = oktmo_scope_key_profile(
                data_root,
                subject_ter_hint="71",
                municipality_hint="Тюменский муниципальный район",
            )
            context_record = AddressNormalizer(use_oktmo=False).normalize("с. Ключи, ул. Ленина, д. 7")
            context_match = _oktmo_scope_candidate_match(
                {"Normalized_Address": "с. Ключи, ул. Ленина, д. 7", "Street": "ул. Ленина", "House": "д. 7"},
                context_record,
                SourceCandidate(value="с. Ключи, ул. Ленина, д. 7", source_file=input_dir / "source.txt", source_kind="text", locator="L1"),
                profile,
            )
            self.assertTrue(context_match["valid"])
            self.assertTrue(
                any(item["value"] == "Ключи" and item["weight"] == 20 and item["valid"] for item in context_match["matches"])
            )
            weak_record = AddressNormalizer(use_oktmo=False).normalize("Ключи")
            weak_match = _oktmo_scope_candidate_match(
                {"Normalized_Address": "Ключи"},
                weak_record,
                SourceCandidate(value="Ключи", source_file=input_dir / "source.txt", source_kind="text", locator="L2"),
                profile,
            )
            self.assertTrue(weak_match)
            self.assertFalse(weak_match["valid"])

    def test_address_collection_rejects_oktmo_scope_conflict(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            tmp_path = Path(tmp)
            data_root = tmp_path / "data"
            input_dir = tmp_path / "input"
            collection_dir = tmp_path / "address_collection"
            input_dir.mkdir()
            write_tyumen_scope_conflict_fixture(data_root)
            (input_dir / "source.txt").write_text(
                "\n".join(
                    [
                        "Тюменская обл, г. Тюмень, ул. Ленина, д. 10",
                        "Тюменская обл, г. Тюмень, ул. Щорса, д. 2, д. Дальнетравное",
                    ]
                ),
                encoding="utf-8",
            )

            result = assemble_address_workbook(
                input_dir=input_dir,
                output_file=collection_dir / "AddressCollection_scope_conflict.xlsx",
                normalizer=AddressNormalizer(
                    data_dir=data_root,
                    subject_ter_hint="71",
                    municipality_hint="город Тюмень",
                ),
                options=AddressAssemblyOptions(
                    action="generate",
                    target_dir=collection_dir,
                    target_column="A",
                    enabled_slots=("municipality", "locality", "street", "house", "oktmo_code", "match_key"),
                    slot_order=("municipality", "locality", "street", "house"),
                ),
            )

            self.assertEqual(result.rejected_rows, 1)
            self.assertEqual(result.conflict_rows, 1)
            staging = load_workbook(result.staging_file, read_only=True, data_only=True)
            try:
                rejected = staging["Rejected"]
                headers = [rejected.cell(row=1, column=column).value for column in range(1, rejected.max_column + 1)]
                row = {header: rejected.cell(row=2, column=index + 1).value for index, header in enumerate(headers)}
                self.assertEqual(row["Drop_Reason"], "oktmo_scope_conflict")
                self.assertEqual(row["Disposition"], "held_out_of_clean_stack")
                self.assertEqual(row["Review_Action"], "manual_scope_review")
                self.assertEqual(row["OKTMO_Code"], "71530000161")
                self.assertEqual(row["OKTMO_Scope_Code"], "71701000001")
                self.assertEqual(row["Municipality"], "Казанский муниципальный округ")

                self.assertIn("OKTMO_Conflicts", staging.sheetnames)
                conflicts = staging["OKTMO_Conflicts"]
                conflict_headers = [conflicts.cell(row=1, column=column).value for column in range(1, conflicts.max_column + 1)]
                conflict = {
                    header: conflicts.cell(row=2, column=index + 1).value
                    for index, header in enumerate(conflict_headers)
                }
                self.assertEqual(conflict["Drop_Reason"], "oktmo_scope_conflict")
                self.assertEqual(conflict["Disposition"], "held_out_of_clean_stack")
                self.assertEqual(conflict["Review_Action"], "manual_scope_review")
                self.assertIn("Дальнетравное", str(conflict["Raw_Address"]))
                self.assertEqual(conflict["Rejected_OKTMO_Code"], "71530000161")
                self.assertEqual(conflict["OKTMO_Scope_Code"], "71701000001")
            finally:
                staging.close()

            workbook = load_workbook(result.output_file, read_only=True, data_only=True)
            try:
                sheet = workbook.worksheets[0]
                output_values = "\n".join(
                    str(sheet.cell(row=row_index, column=1).value or "")
                    for row_index in range(1, sheet.max_row + 1)
                )
                self.assertNotIn("Казанский муниципальный округ", output_values)
                self.assertNotIn("Дальнетравное", output_values)
                self.assertIn("ул. Щорса", output_values)
            finally:
                workbook.close()

    def test_clean_address_append_deduplicates_ready_column_without_staging(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            tmp_path = Path(tmp)
            source_file = tmp_path / "AddressCollection_clean.xlsx"
            target_file = tmp_path / "ready.xlsx"
            first_output = tmp_path / "ready_added.xlsx"
            second_output = tmp_path / "ready_added_again.xlsx"
            pd.DataFrame(
                {
                    "Адрес": [
                        "г. Тюмень, ул. Альфа, д. 2",
                        "г. Тюмень, ул. Бета, д. 10",
                        "Тюмень, улица Бета, дом 10",
                    ]
                }
            ).to_excel(source_file, index=False)
            pd.DataFrame({"Адрес": ["город Тюмень, улица Альфа, дом 2"]}).to_excel(target_file, index=False)

            result = append_collected_address_workbook(
                output_file=first_output,
                normalizer=AddressNormalizer(use_oktmo=False),
                options=AddressAppendOptions(
                    source_file=source_file,
                    target_file=target_file,
                    source_column="Адрес",
                    target_column="Адрес",
                ),
            )

            self.assertEqual(result.source_rows, 3)
            self.assertEqual(result.appended_rows, 1)
            self.assertEqual(result.duplicate_rows, 2)
            self.assertEqual(result.target_column, "A")
            self.assertEqual(result.target_start_row, 3)
            summary = json.loads(result.report_file.read_text(encoding="utf-8"))
            self.assertEqual(summary["normalizer"]["use_oktmo"], False)
            self.assertEqual(summary["options"]["clean_duplicates"], True)
            self.assertEqual(summary["mode"], "address-clean-column-append")
            self.assertEqual(summary["source_contract"]["kind"], "clean_address_column")
            self.assertFalse(summary["source_contract"]["uses_staging"])
            self.assertEqual(summary["source_contract"]["dedupe_basis"], "stable_address_core")
            self.assertEqual(summary["address_diagnostics"]["source"]["total"], 3)
            self.assertEqual(summary["address_diagnostics"]["accepted"]["total"], 1)
            self.assertEqual(summary["address_diagnostics"]["accepted"]["issue_rows"], 0)

            workbook = load_workbook(first_output, read_only=True, data_only=True)
            try:
                sheet = workbook.worksheets[0]
                values = [sheet.cell(row=row, column=1).value for row in range(1, sheet.max_row + 1)]
                self.assertEqual(values[0], "Адрес")
                self.assertEqual(len([value for value in values if value]), 3)
                self.assertIn("ул. Бета", str(values[2]))
            finally:
                workbook.close()

            repeat = append_collected_address_workbook(
                output_file=second_output,
                normalizer=AddressNormalizer(use_oktmo=False),
                options=AddressAppendOptions(
                    source_file=source_file,
                    target_file=first_output,
                    source_column="Адрес",
                    target_column="Адрес",
                ),
            )

            self.assertEqual(repeat.appended_rows, 0)
            self.assertEqual(repeat.duplicate_rows, 3)
            workbook = load_workbook(second_output, read_only=True, data_only=True)
            try:
                sheet = workbook.worksheets[0]
                values = [sheet.cell(row=row, column=1).value for row in range(1, sheet.max_row + 1)]
                self.assertEqual(len([value for value in values if value]), 3)
            finally:
                workbook.close()

    def test_clean_address_append_can_write_compact_new_workbook(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            tmp_path = Path(tmp)
            source_file = tmp_path / "AddressCollection_clean.xlsx"
            target_file = tmp_path / "ready.xlsx"
            output_file = tmp_path / "AddressCollection_clean_output.xlsx"
            pd.DataFrame({"Адрес": ["г. Тюмень, ул. Альфа, д. 2", "г. Тюмень, ул. Бета, д. 10"]}).to_excel(
                source_file,
                index=False,
            )
            pd.DataFrame(
                {
                    "Адрес": ["город Тюмень, улица Альфа, дом 2"],
                    "Старая лишняя колонка": ["оставить только адресный столбец в чистом выводе"],
                }
            ).to_excel(target_file, index=False, sheet_name="Dirty")

            result = append_collected_address_workbook(
                output_file=output_file,
                normalizer=AddressNormalizer(use_oktmo=False),
                options=AddressAppendOptions(
                    source_file=source_file,
                    target_file=target_file,
                    source_column="Адрес",
                    target_column="Адрес",
                    clean_output_workbook=True,
                ),
            )

            self.assertEqual(result.appended_rows, 1)
            self.assertEqual(result.target_column, "A")
            self.assertEqual(result.target_start_row, 2)
            workbook = load_workbook(output_file)
            try:
                self.assertEqual(workbook.sheetnames, ["Addresses"])
                sheet = workbook.worksheets[0]
                self.assertEqual(sheet.cell(row=1, column=1).value, "Normalized_Address")
                self.assertEqual(sheet.max_column, 1)
                values = [sheet.cell(row=row, column=1).value for row in range(1, sheet.max_row + 1)]
                self.assertEqual(len([value for value in values if value]), 3)
                self.assertIn("ул. Альфа", str(values[1]))
                self.assertIn("ул. Бета", str(values[2]))
                self.assertEqual(sheet.freeze_panes, "A2")
                self.assertEqual(float(sheet.cell(row=2, column=1).font.sz), 9.0)
            finally:
                workbook.close()

            summary = json.loads(result.report_file.read_text(encoding="utf-8"))
            self.assertTrue(summary["options"]["clean_output_workbook"])
            self.assertEqual(summary["accepted_rows"], 1)
            self.assertEqual(summary["output_existing_rows"], 1)
            self.assertEqual(summary["output_rows"], 2)

    def test_address_collection_builds_structured_street_house_pair(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            tmp_path = Path(tmp)
            input_dir = tmp_path / "input"
            collection_dir = tmp_path / "address_collection"
            input_dir.mkdir()
            pd.DataFrame(
                [
                    ["г. Тюмень", None, None],
                    ["Улица", "Дом", "Комментарий"],
                    ["Ленина", "10", "школа"],
                ]
            ).to_excel(input_dir / "split_address.xlsx", index=False, header=False)

            result = assemble_address_workbook(
                input_dir=input_dir,
                output_file=collection_dir / "AddressCollection_new.xlsx",
                normalizer=AddressNormalizer(city_name="Тюмень", use_oktmo=False),
                options=AddressAssemblyOptions(
                    action="generate",
                    target_dir=collection_dir,
                    target_column="A",
                    enabled_slots=("locality", "street", "house", "match_key"),
                    slot_order=("locality", "street", "house"),
                ),
            )

            workbook = load_workbook(result.output_file, read_only=True, data_only=True)
            try:
                value = str(workbook.worksheets[0].cell(row=1, column=1).value)
                self.assertIn("ул. Ленина", value)
                self.assertIn("д. 10", value)
            finally:
                workbook.close()

            staging = load_workbook(result.staging_file, read_only=True, data_only=True)
            try:
                sheet = staging["Slot_Candidates"]
                headers = [sheet.cell(row=1, column=column).value for column in range(1, sheet.max_column + 1)]
                status_column = headers.index("Assembly_Status") + 1
                street_candidates_column = headers.index("Street_Candidates") + 1
                self.assertEqual(sheet.cell(row=2, column=status_column).value, "ready")
                self.assertIn("ул. Ленина", str(sheet.cell(row=2, column=street_candidates_column).value))
            finally:
                staging.close()

    def test_address_collection_source_columns_keep_xlsx_groups_separate(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            tmp_path = Path(tmp)
            input_dir = tmp_path / "input"
            collection_dir = tmp_path / "address_collection"
            input_dir.mkdir()
            pd.DataFrame(
                [
                    ["ID", "Местоположение", "Примечание"],
                    [1, "г. Тюмень, ул. Бета, д. 10", "г. Тюмень, ул. Альфа, д. 2"],
                ]
            ).to_excel(input_dir / "groups.xlsx", index=False, header=False)

            result = assemble_address_workbook(
                input_dir=input_dir,
                output_file=collection_dir / "AddressCollection_c.xlsx",
                normalizer=AddressNormalizer(city_name="Тюмень", use_oktmo=False),
                options=AddressAssemblyOptions(
                    action="generate",
                    source_columns="C",
                    target_dir=collection_dir,
                    target_column="A",
                    enabled_slots=("locality", "street", "house", "match_key"),
                    slot_order=("locality", "street", "house"),
                ),
            )

            workbook = load_workbook(result.output_file, read_only=True, data_only=True)
            try:
                values = [
                    str(workbook.worksheets[0].cell(row=row, column=1).value or "")
                    for row in range(1, workbook.worksheets[0].max_row + 1)
                ]
                joined = "\n".join(values)
                self.assertIn("ул. Альфа", joined)
                self.assertNotIn("ул. Бета", joined)
            finally:
                workbook.close()

    def test_address_collection_source_columns_do_not_shift_after_empty_leading_column(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            tmp_path = Path(tmp)
            input_dir = tmp_path / "input"
            collection_dir = tmp_path / "address_collection"
            input_dir.mkdir()
            pd.DataFrame(
                [
                    [None, "Эталон", "Группа C", "Группа D"],
                    [None, "г. Тюмень, ул. Омега, д. 9", "г. Тюмень, ул. Альфа, д. 2", "г. Тюмень, ул. Бета, д. 10"],
                ]
            ).to_excel(input_dir / "empty_leading.xlsx", index=False, header=False)

            result = assemble_address_workbook(
                input_dir=input_dir,
                output_file=collection_dir / "AddressCollection_c.xlsx",
                normalizer=AddressNormalizer(city_name="Тюмень", use_oktmo=False),
                options=AddressAssemblyOptions(
                    action="generate",
                    source_columns="C",
                    target_dir=collection_dir,
                    target_column="A",
                    enabled_slots=("locality", "street", "house", "match_key"),
                    slot_order=("locality", "street", "house"),
                ),
            )

            workbook = load_workbook(result.output_file, read_only=True, data_only=True)
            try:
                values = [
                    str(workbook.worksheets[0].cell(row=row, column=1).value or "")
                    for row in range(1, workbook.worksheets[0].max_row + 1)
                ]
                joined = "\n".join(values)
                self.assertIn("ул. Альфа", joined)
                self.assertNotIn("ул. Бета", joined)
                self.assertNotIn("ул. Омега", joined)
            finally:
                workbook.close()

    def test_address_collection_final_output_deduplicates_same_rendered_address(self) -> None:
        normalizer = AddressNormalizer(city_name="Тюмень", use_oktmo=False)
        raw = "г. Тюмень, ул. Пархоменко, д. 60"
        record = normalizer.normalize(raw)

        def grouped(key: str, source_name: str) -> _GroupedRows:
            candidate = SourceCandidate(
                value=raw,
                source_file=Path(source_name),
                source_kind="source",
                locator="A1",
            )
            return _GroupedRows(
                key,
                (
                    _CandidateRow(
                        row={
                            "Normalized_Address": raw,
                            "Raw_Address": raw,
                            "Locality": "г. Тюмень",
                            "Street": "ул. Пархоменко",
                            "House": "д. 60",
                            "Match_Key": key,
                            "Source_Count": 1,
                            "Sources": source_name,
                            "First_Source": source_name,
                        },
                        record=record,
                        candidate=candidate,
                        source_kind="source",
                        quality_score=10.0,
                    ),
                ),
            )

        normalized_rows, rejected_rows, duplicates, staged_rows = _final_rows_from_groups(
            [grouped("different-key-1", "one.xlsx"), grouped("different-key-2", "two.xlsx")],
            enabled_slots=("locality", "street", "house", "match_key"),
            slot_order=("locality", "street", "house"),
            omit_microdistrict_when_street_found=True,
            include_slots=True,
            sanitize_incomplete=True,
        )

        self.assertFalse(rejected_rows)
        self.assertEqual(len(normalized_rows), 1)
        self.assertEqual(duplicates, 1)
        self.assertEqual(int(normalized_rows[0]["Source_Count"]), 2)
        self.assertTrue(any(row["Assembly_Status"] == "duplicate" and row["Drop_Reason"] == "final_output_duplicate" for row in staged_rows))

    def test_address_collection_strict_sanitizer_rejects_route_street_and_repairs_duplicate_locator(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            tmp_path = Path(tmp)
            input_dir = tmp_path / "input"
            collection_dir = tmp_path / "address_collection"
            route_dir = input_dir / "Приложение 6 Подвоз маршруты"
            route_dir.mkdir(parents=True)
            (input_dir / "clean.txt").write_text("г. Тюмень, ул. ул. Елизарова, д. 74", encoding="utf-8")
            (input_dir / "compact_street.txt").write_text("г. Тюмень, ул.Пермякова, д. 25, к. 1", encoding="utf-8")
            (input_dir / "spaced_marker_dot.txt").write_text("г. Тюмень, ул . Ленина, д. 10", encoding="utf-8")
            (input_dir / "dirty_admin.txt").write_text(
                "МАОУ гимназия № 49 города Тюмени № 20 район улицы Червишевский тракт, "
                "г. Тюмень, ул. Пархоменко, д. 60",
                encoding="utf-8",
            )
            (route_dir / "routes.txt").write_text(
                "1. МАОУ СОШ № 45 (ул. Ишимская, 11) - с/о Степное - ДК Орфей, д. 25",
                encoding="utf-8",
            )

            result = assemble_address_workbook(
                input_dir=input_dir,
                output_file=collection_dir / "AddressCollection_strict.xlsx",
                normalizer=AddressNormalizer(city_name="Тюмень", use_oktmo=False),
                options=AddressAssemblyOptions(
                    action="generate",
                    target_dir=collection_dir,
                    target_column="A",
                    enabled_slots=("locality", "street", "house", "match_key"),
                    slot_order=("locality", "street", "house"),
                ),
            )

            workbook = load_workbook(result.output_file, read_only=True, data_only=True)
            try:
                values = [str(workbook.worksheets[0].cell(row=row, column=1).value or "") for row in range(1, workbook.worksheets[0].max_row + 1)]
                joined = "\n".join(values)
                self.assertEqual(len(values), 4)
                self.assertIn("ул. Елизарова", joined)
                self.assertIn("ул. Ленина", joined)
                self.assertIn("ул. Пархоменко", joined)
                self.assertIn("ул. Пермякова", joined)
                self.assertNotIn("ул .", joined)
                self.assertNotIn("ул. ул.", joined)
                self.assertNotIn("ул.Пермякова", joined)
                self.assertNotIn("пер. мякова", joined.lower())
                self.assertFalse(any("МАОУ" in value for value in values))
            finally:
                workbook.close()

            staging = load_workbook(result.staging_file, read_only=True, data_only=True)
            try:
                sheet = staging["Rejected"]
                headers = [sheet.cell(row=1, column=column).value for column in range(1, sheet.max_column + 1)]
                reason_column = headers.index("Drop_Reason") + 1
                reasons = [sheet.cell(row=row, column=reason_column).value for row in range(2, sheet.max_row + 1)]
                self.assertTrue({"route_source_fragment", "house_without_street_or_territory"}.intersection(reasons))
            finally:
                staging.close()

    def test_address_collection_benchmark_matches_normalized_core(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            tmp_path = Path(tmp)
            collected_file = tmp_path / "collected.xlsx"
            benchmark_file = tmp_path / "benchmark.xlsx"
            pd.DataFrame({"Address": ["ул. Ленина, д. 10"]}).to_excel(collected_file, index=False)
            pd.DataFrame(
                {
                    "Местоположение": [
                        "г. Тюмень, улица Ленина, дом 10",
                        "г. Тюмень, ул. Республики, д. 5",
                    ]
                }
            ).to_excel(benchmark_file, index=False)

            result = benchmark_address_collection(
                collected_file=collected_file,
                benchmark_file=benchmark_file,
                normalizer=AddressNormalizer(city_name="Тюмень", use_oktmo=False),
                benchmark_column="Местоположение",
                collected_column="Address",
                report_file=tmp_path / "benchmark.summary.json",
                enabled_slots=("street", "house", "match_key"),
                slot_order=("street", "house"),
            )

            self.assertEqual(result["benchmark_unique_key_count"], 2)
            self.assertEqual(result["collected_unique_key_count"], 1)
            self.assertEqual(result["matched_unique_key_count"], 1)
            self.assertEqual(result["missing_unique_key_count"], 1)
            self.assertAlmostEqual(result["coverage_ratio"], 0.5)
            self.assertEqual(result["quality"]["coverage_percent"], 50.0)
            self.assertEqual(result["quality"]["precision_percent"], 100.0)
            self.assertEqual(result["address_diagnostics"]["benchmark"]["total"], 2)
            self.assertEqual(result["address_diagnostics"]["collected"]["total"], 1)
            self.assertEqual(result["slot_fill"]["benchmark"]["slots"]["street"]["filled"], 2)
            self.assertEqual(result["slot_fill"]["collected"]["slots"]["house"]["filled"], 1)
            self.assertEqual(result["artifacts"]["reports"], [str(tmp_path / "benchmark.summary.json")])
            persisted = json.loads((tmp_path / "benchmark.summary.json").read_text(encoding="utf-8"))
            self.assertEqual(persisted["quality"], result["quality"])


class SourceCleanupTests(unittest.TestCase):
    def test_reads_searchable_pdf_text_layer(self) -> None:
        if importlib.util.find_spec("pymupdf") is None and importlib.util.find_spec("fitz") is None:
            self.skipTest("PyMuPDF is not installed")

        with tempfile.TemporaryDirectory() as tmp:
            input_dir = Path(tmp)
            pdf_path = input_dir / "searchable.pdf"
            write_searchable_pdf(pdf_path, "Searchable PDF text layer, Lenina 10")

            candidates, issues = extract_source_candidates(pdf_path)

            self.assertFalse(issues)
            self.assertTrue(any("Searchable PDF text layer" in item.value for item in candidates))

    def test_deletes_microsoft_office_temp_files_only(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            input_dir = Path(tmp)
            keep = input_dir / "keep.xlsx"
            keep.write_bytes(b"real")
            temp_files = [
                input_dir / "~$keep.xlsx",
                input_dir / ".~lock.keep.xlsx#",
                input_dir / "WordRecovery.wbk",
                input_dir / "WordAutoSave.asd",
                input_dir / "ExcelBackup.xlk",
                input_dir / "OfficeTemp.tmp",
            ]
            for path in temp_files:
                path.write_bytes(b"temp")

            removed = delete_office_temp_files(input_dir)

            self.assertEqual(len(removed), len(temp_files))
            self.assertTrue(keep.exists())
            self.assertTrue(all(not path.exists() for path in temp_files))

    def test_conversion_cleanup_deletes_same_name_old_pdf_txt_md_without_com(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            input_dir = Path(tmp)
            keeper = input_dir / "source.docx"
            keeper.write_bytes(b"docx")
            duplicates = [
                input_dir / "source.doc",
                input_dir / "source.pdf",
                input_dir / "source.txt",
                input_dir / "source.csv",
                input_dir / "source.md",
                input_dir / "~$source.docx",
            ]
            for path in duplicates:
                path.write_bytes(b"duplicate")

            result = convert_legacy_office_folder(input_dir)

            self.assertTrue(keeper.exists())
            self.assertTrue(all(not path.exists() for path in duplicates))
            self.assertFalse(result.failed)
            self.assertGreaterEqual(len(result.deleted), len(duplicates))


if __name__ == "__main__":
    unittest.main()
