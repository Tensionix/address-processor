from __future__ import annotations

from pathlib import Path
from tempfile import TemporaryDirectory
import sys
import unittest

from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill


ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from system_core.excel_layout import format_result_workbooks, format_workbook, percentile


class ExcelLayoutTests(unittest.TestCase):
    def test_percentile_uses_ninety_percent_coverage(self) -> None:
        self.assertEqual(percentile([10] * 9 + [200]), 10)

    def test_format_workbook_wraps_text_and_ignores_one_width_outlier(self) -> None:
        with TemporaryDirectory() as tmp:
            path = Path(tmp) / "result.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Result"
            sheet["A1"] = "Address"
            sheet["A1"].fill = PatternFill("solid", fgColor="112233")
            for row in range(2, 11):
                sheet.cell(row, 1, "Короткий текст")
            sheet["A11"] = "Очень длинная строка " * 20
            sheet["B1"] = "Formula"
            sheet["B2"] = "=1+1"
            workbook.save(path)
            workbook.close()

            result = format_workbook(path)

            self.assertEqual(result.path, path.resolve())
            formatted = load_workbook(path, data_only=False)
            try:
                output = formatted["Result"]
                self.assertLess(output.column_dimensions["A"].width, 30)
                self.assertTrue(output["A2"].alignment.wrap_text)
                self.assertGreater(output.row_dimensions[11].height, 18)
                self.assertLessEqual(output.row_dimensions[11].height, 180)
                self.assertEqual(output["A1"].fill.fgColor.rgb, "00112233")
                self.assertEqual(output["B2"].value, "=1+1")
            finally:
                formatted.close()

    def test_result_formatter_formats_outputs_and_staging_but_not_inputs(self) -> None:
        with TemporaryDirectory() as tmp:
            root = Path(tmp)
            output = root / "output.xlsx"
            staging = root / "staging.xlsx"
            source = root / "source.xlsx"
            for path in (output, staging, source):
                workbook = Workbook()
                workbook.active["A1"] = "A long header that needs a readable column"
                workbook.save(path)
                workbook.close()

            formatted = format_result_workbooks(
                {
                    "artifacts": {
                        "outputs": [str(output)],
                        "staging": [str(staging)],
                        "inputs": [str(source)],
                    }
                },
                project_root=root,
            )

            self.assertEqual({item.path for item in formatted}, {output.resolve(), staging.resolve()})
            source_book = load_workbook(source)
            output_book = load_workbook(output)
            try:
                self.assertIsNone(source_book.active["A1"].alignment.wrap_text)
                self.assertTrue(output_book.active["A1"].alignment.wrap_text)
            finally:
                source_book.close()
                output_book.close()


if __name__ == "__main__":
    unittest.main()
