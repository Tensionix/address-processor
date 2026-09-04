from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile
import contextlib
import json
import os
import shutil
import time

from system_core.address_engine import AddressNormalizer
from system_core.address_engine.legacy_office_converter import convert_legacy_office_folder
from system_core.address_engine.oktmo_lookup import (
    clear_oktmo_caches as clear_oktmo_lookup_caches,
    municipality_scope_options,
    region_scope_options,
    resolve_municipality_scope_hint,
    resolve_region_scope_id,
    resolve_region_scope_ter,
    warm_oktmo_cache,
)
from system_core.address_engine.oktmo_user_keys import (
    add_current_keys,
    candidate_key_terms,
    candidate_options_from_values,
    clear_all_pins,
    clear_current_keys,
    current_keys_path,
    decorate_pinned_options,
    export_pins_bundle,
    import_pins_bundle,
    is_pinned,
    load_current_keys,
    municipality_key_terms,
    pin_file_status,
    pin_value,
    pins_bundle_path,
    region_key_terms,
    settlement_candidate_options,
    unpin_value,
    write_current_keys,
)
from system_core.address_engine.normalizer_runtime import normalizer_oktmo_key_profile, normalizer_runtime_report
from system_core.address_engine.reference_builder import (
    DEFAULT_ENABLED_SLOTS,
    DEFAULT_SELECTED_SLOTS,
    DEFAULT_SLOT_ORDER,
    ReferenceOptions,
    build_reference_workbook,
)
from system_core.address_engine.reference_processor import ReferenceProcessingOptions, process_reference_workbook
from system_core.address_engine.rosstat_update import DEFAULT_OKTMO_URL, update_rosstat_oktmo_data as update_oktmo_snapshot
from system_core.address_engine.source_cleanup import CleanupRecord, delete_office_temp_files
from system_core.address_engine.source_intake import excel_column_letter
from system_core.address_engine.source_table_assembly import (
    AddressAssemblyOptions,
    AddressAppendOptions,
    append_collected_address_workbook,
    assemble_address_workbook,
    benchmark_address_collection,
)
from system_core.core.jobs import JobContext
from system_core.Parser_Diagnostic import run_diagnostic
from system_core.remove_empty_rows import clean_workbook, default_output_path
from system_core.safe_table_join import ComparisonOptions, JoinOptions, build_comparison_workbook, parse_column_spec, safe_join_workbooks
from system_core.Ultimate_GT_Aligner import CONFIG, process_file


@dataclass(frozen=True)
class AddressRuntimeContext:
    normalizer: AddressNormalizer
    subject_ter_hint: str | None
    municipality_hint: str | None
    oktmo_scope_enabled: bool


OKTMO_SCOPE_PARAMETER_IDS: tuple[str, ...] = (
    "use_oktmo",
)

ADDRESS_RUNTIME_PARAMETER_IDS: tuple[str, ...] = (
    "city",
    *OKTMO_SCOPE_PARAMETER_IDS,
)

OKTMO_RESET_FIELD_UPDATES: dict[str, object] = {
    "oktmo_scope_region": None,
    "oktmo_scope_municipality": None,
    "oktmo_settlement_query": "",
    "oktmo_candidate_ref": "",
    "oktmo_candidate_options": [],
    "oktmo_candidate_preview": "",
    "oktmo_selected_municipality": "",
    "oktmo_selected_settlement": "",
    "oktmo_selected_code": "",
}

ALIGNER_RUNTIME_PARAMETER_IDS: tuple[str, ...] = (
    "ground_truth_column",
    "ground_truth_column_header",
    "target_columns",
    "target_column_header",
    "alignment_mode",
    "companion_mode",
    "satellite_columns",
    "normalize_before_match",
    "normalize_after_match",
    "collect_before_match",
    "collect_whole_document",
    *ADDRESS_RUNTIME_PARAMETER_IDS,
)


def _xlsx_files(folder: Path) -> list[Path]:
    return sorted(path for path in folder.glob("*.xlsx") if path.is_file() and not path.name.startswith("~$"))


def input_column_options(root: Path | str | None = None) -> list[dict[str, object]]:
    project_root = Path(root or Path(__file__).resolve().parents[2]).resolve()
    options: list[dict[str, object]] = [
        {
            "value": "",
            "label": "Auto/manual",
            "label_ru": "Авто/ручной ввод",
        }
    ]
    input_dir = project_root / "input"
    files = _xlsx_files(input_dir) if input_dir.exists() else []
    if not files:
        options[0]["label"] = "Auto/manual (no input workbook found)"
        options[0]["label_ru"] = "Авто/ручной ввод (в input пока нет XLSX)"
        return options
    try:
        from openpyxl import load_workbook

        workbook = files[0]
        book = load_workbook(workbook, read_only=True, data_only=True)
        try:
            sheet = book.worksheets[0]
            first_row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
            options.extend(_reference_header_options(first_row, sheet.title))
        finally:
            book.close()
    except Exception as exc:
        options.append(
            {
                "value": "",
                "label": f"Could not read headers: {exc}",
                "label_ru": f"Не удалось прочитать заголовки: {exc}",
            }
        )
    return options


def reference_column_options(root: Path | str | None = None) -> list[dict[str, object]]:
    project_root = Path(root or Path(__file__).resolve().parents[2]).resolve()
    options: list[dict[str, object]] = [
        {
            "value": "",
            "label": "Auto-detect",
            "label_ru": "Авто-поиск",
        }
    ]
    workbook = _latest_reference_workbook(project_root)
    if not workbook:
        options[0]["label"] = "Auto-detect (no AddressReference workbook found)"
        options[0]["label_ru"] = "Авто-поиск (эталонная книга пока не найдена)"
        return options
    try:
        from openpyxl import load_workbook

        book = load_workbook(workbook, read_only=True, data_only=True)
        try:
            for sheet in book.worksheets:
                first_row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
                row_options = _reference_header_options(first_row, sheet.title)
                if row_options:
                    options.extend(row_options)
                    break
        finally:
            book.close()
    except Exception as exc:
        options.append(
            {
                "value": "",
                "label": f"Could not read headers: {exc}",
                "label_ru": f"Не удалось прочитать заголовки: {exc}",
            }
        )
    return options


def _comparison_workbook(folder: Path) -> Path | None:
    files = sorted(
        (path for path in folder.glob("*.xls*") if path.is_file() and not path.name.startswith("~$")),
        key=lambda path: (path.stat().st_mtime_ns, path.name.lower()),
        reverse=True,
    )
    return files[0] if files else None


def _selected_workbook_path(
    root: Path | str | None,
    values: dict[str, object] | None,
    field_id: str,
) -> Path | None:
    project_root = Path(root or Path(__file__).resolve().parents[2]).resolve()
    raw = str((values or {}).get(field_id) or "").strip().strip('"')
    if not raw:
        return None
    workbook = Path(os.path.expandvars(raw)).expanduser()
    if not workbook.is_absolute():
        workbook = project_root / workbook
    return workbook.resolve()


def _workbook_sheet_options(workbook: Path | None, missing_text: str) -> list[dict[str, object]]:
    if workbook is None:
        return [{"value": "", "label": missing_text, "label_ru": missing_text}]
    if not workbook.exists():
        return [{"value": "", "label": f"Workbook not found: {workbook.name}", "label_ru": f"Книга не найдена: {workbook.name}"}]
    try:
        from openpyxl import load_workbook

        book = load_workbook(workbook, read_only=True, data_only=True)
        try:
            return [{"value": sheet.title, "label": sheet.title, "label_ru": sheet.title} for sheet in book.worksheets]
        finally:
            book.close()
    except Exception as exc:
        return [{"value": "", "label": f"Could not read sheets: {exc}", "label_ru": f"Не удалось прочитать листы: {exc}"}]


def comparison_sheet_a_options(
    root: Path | str | None = None,
    values: dict[str, object] | None = None,
    _field: dict[str, object] | None = None,
) -> list[dict[str, object]]:
    workbook = _selected_workbook_path(root, values, "file_a")
    return _workbook_sheet_options(workbook, "Сначала выберите таблицу A")


def comparison_sheet_b_options(
    root: Path | str | None = None,
    values: dict[str, object] | None = None,
    _field: dict[str, object] | None = None,
) -> list[dict[str, object]]:
    workbook = _selected_workbook_path(root, values, "file_b")
    return _workbook_sheet_options(workbook, "Сначала выберите таблицу B")


def safe_join_primary_sheet_options(
    root: Path | str | None = None,
    values: dict[str, object] | None = None,
    _field: dict[str, object] | None = None,
) -> list[dict[str, object]]:
    workbook = _selected_workbook_path(root, values, "primary_file")
    return _workbook_sheet_options(workbook, "Сначала выберите основную таблицу")


def safe_join_source_sheet_options(
    root: Path | str | None = None,
    values: dict[str, object] | None = None,
    _field: dict[str, object] | None = None,
) -> list[dict[str, object]]:
    workbook = _selected_workbook_path(root, values, "source_file")
    return _workbook_sheet_options(workbook, "Сначала выберите таблицу-источник")


def address_collection_column_options(root: Path | str | None = None) -> list[dict[str, object]]:
    project_root = Path(root or Path(__file__).resolve().parents[2]).resolve()
    options: list[dict[str, object]] = [
        {
            "value": "",
            "label": "A / manual",
            "label_ru": "A / ручной ввод",
        }
    ]
    workbook = _latest_xlsx_in_folder(project_root / "address_collection") or _latest_reference_workbook(project_root)
    if not workbook:
        options[0]["label"] = "A / manual (no address workbook found)"
        options[0]["label_ru"] = "A / ручной ввод (адресная книга пока не найдена)"
        return options
    try:
        from openpyxl import load_workbook

        book = load_workbook(workbook, read_only=True, data_only=True)
        try:
            sheet = book.worksheets[0]
            first_row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
            options.extend(_reference_header_options(first_row, sheet.title))
        finally:
            book.close()
    except Exception as exc:
        options.append(
            {
                "value": "",
                "label": f"Could not read headers: {exc}",
                "label_ru": f"Не удалось прочитать заголовки: {exc}",
            }
        )
    return options


def oktmo_region_scope_options(
    root: Path | str | None = None,
    values: dict[str, object] | None = None,
    field: dict[str, object] | None = None,
) -> list[dict[str, object]]:
    project_root = Path(root or Path(__file__).resolve().parents[2]).resolve()
    data_dir = _option_data_dir(root, values)
    options: list[dict[str, object]] = [
        {
            "value": "",
            "label": "All OKTMO regions",
            "label_ru": "Все регионы ОКТМО",
        }
    ]
    options.extend(region_scope_options(data_dir))
    return decorate_pinned_options(project_root, "region", options)


def oktmo_municipality_scope_options(
    root: Path | str | None = None,
    values: dict[str, object] | None = None,
    field: dict[str, object] | None = None,
) -> list[dict[str, object]]:
    project_root = Path(root or Path(__file__).resolve().parents[2]).resolve()
    data_dir = _option_data_dir(root, values)
    region = str((values or {}).get("oktmo_scope_region") or (values or {}).get("subject_ter_hint") or "").strip()
    if not region:
        return [
            {
                "value": "",
                "label": "Select an OKTMO region first",
                "label_ru": "Сначала выберите регион ОКТМО",
            }
        ]
    subject_scope = resolve_region_scope_id(data_dir, region) or resolve_region_scope_ter(data_dir, region) or region
    options: list[dict[str, object]] = [
        {
            "value": "",
            "label": "All municipalities in scope",
            "label_ru": "Все муниципалитеты в области поиска",
        }
    ]
    options.extend(municipality_scope_options(data_dir, subject_scope))
    return decorate_pinned_options(project_root, "municipality", options, scope=subject_scope)


def oktmo_settlement_candidate_options(
    root: Path | str | None = None,
    values: dict[str, object] | None = None,
    field: dict[str, object] | None = None,
) -> list[dict[str, str]]:
    return candidate_options_from_values(dict(values or {}))


def oktmo_data_status(
    root: Path | str | None = None,
    values: dict[str, object] | None = None,
    field: dict[str, object] | None = None,
) -> dict[str, object]:
    data_dir = _option_data_dir(root, values)
    rosstat_dir = data_dir / "rosstat"
    files = sorted((path for path in rosstat_dir.glob("data-*.csv") if path.is_file()), key=lambda path: path.name.lower())
    latest = max(files, key=lambda path: (path.stat().st_mtime, path.name.lower())) if files else None
    return {
        "ok": latest is not None,
        "data_dir": str(data_dir),
        "rosstat_dir": str(rosstat_dir),
        "latest": str(latest) if latest else "",
        "latest_name": latest.name if latest else "",
        "files": len(files),
    }


def _option_data_dir(root: Path | str | None, values: dict[str, object] | None = None) -> Path:
    project_root = Path(root or Path(__file__).resolve().parents[2]).resolve()
    raw = str((values or {}).get("oktmo_data_dir") or "").strip().strip('"')
    if not raw:
        return project_root / "data"
    path = Path(os.path.expandvars(raw)).expanduser()
    if not path.is_absolute():
        path = project_root / path
    if path.name.lower() == "rosstat":
        return path.parent
    return path


def _clean_ui_value(value: object) -> str:
    return " ".join(str(value or "").replace("\xa0", " ").split()).strip(" ,;")


def _selected_oktmo_scope_value(values: dict[str, object] | None, kind: str) -> str:
    values = values or {}
    key = "oktmo_scope_region" if kind == "region" else "oktmo_scope_municipality"
    return _clean_ui_value(values.get(key))


def _selected_region_scope_ter(root: Path | str | None, values: dict[str, object] | None) -> str:
    data_dir = _option_data_dir(root, values)
    region = _selected_oktmo_scope_value(values, "region") or _clean_ui_value((values or {}).get("subject_ter_hint"))
    return resolve_region_scope_ter(data_dir, region) or region


def _selected_region_scope_id(root: Path | str | None, values: dict[str, object] | None) -> str:
    data_dir = _option_data_dir(root, values)
    region = _selected_oktmo_scope_value(values, "region") or _clean_ui_value((values or {}).get("subject_ter_hint"))
    return resolve_region_scope_id(data_dir, region) or resolve_region_scope_ter(data_dir, region) or region


def _selected_municipality_scope_ter(root: Path | str | None, values: dict[str, object] | None) -> str:
    region_scope = _selected_region_scope_id(root, values)
    if region_scope:
        return region_scope
    municipality = _selected_oktmo_scope_value(values, "municipality")
    digits = "".join(ch for ch in municipality if ch.isdigit())
    return digits[:2] if len(digits) >= 2 else ""


def _selected_oktmo_option(root: Path | str | None, values: dict[str, object] | None, kind: str) -> dict[str, object]:
    selected = _selected_oktmo_scope_value(values, kind)
    if not selected:
        raise ValueError("Выберите регион ОКТМО." if kind == "region" else "Выберите муниципалитет ОКТМО.")
    options = oktmo_region_scope_options(root, values) if kind == "region" else oktmo_municipality_scope_options(root, values)
    for option in options:
        candidates = {
            _clean_ui_value(option.get("value")),
            _clean_ui_value(option.get("label")),
            _clean_ui_value(option.get("label_ru")),
            _clean_ui_value(option.get("code")),
        }
        if selected in candidates:
            return dict(option)
    return {"value": selected, "label": selected, "label_ru": selected, "code": ""}


def oktmo_scope_pin_status(
    root: Path | str | None = None,
    values: dict[str, object] | None = None,
    kind: str = "region",
) -> dict[str, object]:
    project_root = Path(root or Path(__file__).resolve().parents[2]).resolve()
    selected = _selected_oktmo_scope_value(values, kind)
    scope = _selected_municipality_scope_ter(root, values) if kind == "municipality" else ""
    return {
        "kind": kind,
        "value": selected,
        "pinned": is_pinned(project_root, kind, selected, scope=scope) if selected else False,
    }


def oktmo_pin_scope_value(
    root: Path | str | None = None,
    values: dict[str, object] | None = None,
    kind: str = "region",
) -> dict[str, object]:
    project_root = Path(root or Path(__file__).resolve().parents[2]).resolve()
    option = _selected_oktmo_option(root, values, kind)
    scope = _selected_municipality_scope_ter(root, values) if kind == "municipality" else ""
    result = pin_value(
        project_root,
        kind,
        _clean_ui_value(option.get("value")),
        _clean_ui_value(option.get("label_ru") or option.get("label")),
        _clean_ui_value(option.get("code")),
        scope,
    )
    export_pins_bundle(project_root)
    result.update({"kind": kind, "value": option.get("value"), "pinned": True})
    return result


def oktmo_unpin_scope_value(
    root: Path | str | None = None,
    values: dict[str, object] | None = None,
    kind: str = "region",
) -> dict[str, object]:
    project_root = Path(root or Path(__file__).resolve().parents[2]).resolve()
    selected = _selected_oktmo_scope_value(values, kind)
    if not selected:
        raise ValueError("Выберите регион ОКТМО." if kind == "region" else "Выберите муниципалитет ОКТМО.")
    scope = _selected_municipality_scope_ter(root, values) if kind == "municipality" else ""
    result = unpin_value(project_root, kind, selected, scope)
    export_pins_bundle(project_root)
    result.update({"kind": kind, "value": selected, "pinned": False})
    return result


def oktmo_add_scope_keys(
    root: Path | str | None = None,
    values: dict[str, object] | None = None,
    kind: str = "region",
) -> dict[str, object]:
    project_root = Path(root or Path(__file__).resolve().parents[2]).resolve()
    data_dir = _option_data_dir(project_root, values)
    values = values or {}
    if kind == "region":
        region = _selected_oktmo_scope_value(values, "region")
        terms = region_key_terms(data_dir, region)
    else:
        region = _selected_oktmo_scope_value(values, "region")
        municipality = _selected_oktmo_scope_value(values, "municipality")
        terms = municipality_key_terms(data_dir, region, municipality)
    result = add_current_keys(project_root, terms)
    result.update({"kind": kind, "added": len(terms), "keys": len(load_current_keys(project_root))})
    return result


def oktmo_find_settlement_candidates(
    root: Path | str | None = None,
    values: dict[str, object] | None = None,
) -> dict[str, object]:
    project_root = Path(root or Path(__file__).resolve().parents[2]).resolve()
    data_dir = _option_data_dir(project_root, values)
    params = dict(values or {})
    query = _clean_ui_value(params.get("oktmo_settlement_query") or params.get("settlement") or params.get("locality"))
    region = _clean_ui_value(params.get("oktmo_scope_region") or params.get("subject_ter_hint"))
    if not query:
        preview = (
            "Введите населённый пункт в поле поиска, затем нажмите «Найти кандидатов». "
            "Если нужно добавить весь выбранный муниципалитет, нажмите + рядом с полем муниципалитета."
        )
        return {
            "candidates": [],
            "count": 0,
            "field_updates": {
                "oktmo_candidate_preview": preview,
                "oktmo_candidate_ref": "",
                "oktmo_candidate_options": [],
            },
        }

    candidates = settlement_candidate_options(data_dir, params, limit=40)
    if not candidates:
        preview = f"Не найдено в ОКТМО: {query}"
        return {
            "candidates": [],
            "count": 0,
            "field_updates": {
                "oktmo_candidate_preview": preview,
                "oktmo_candidate_ref": "",
                "oktmo_candidate_options": [],
            },
        }

    lines: list[str] = []
    for index, item in enumerate(candidates, start=1):
        ambiguity = "однозначно" if item.get("unambiguous_in_region") else "есть одноимённые варианты"
        lines.append(
            f"{index}. {item.get('name') or '-'} | "
            f"МО: {item.get('municipality') or '-'} | "
            f"регион: {item.get('region') or region or '-'} | "
            f"ОКТМО {item.get('code') or '-'} | {ambiguity}"
        )
    selected = candidates[0] if len(candidates) == 1 or bool(candidates[0].get("unambiguous_in_region")) else None
    options = [
        {
            "value": str(item.get("code") or item.get("value") or ""),
            "label": str(item.get("label_ru") or item.get("label") or item.get("code") or ""),
            "label_ru": str(item.get("label_ru") or item.get("label") or item.get("code") or ""),
        }
        for item in candidates
        if item.get("code") or item.get("value")
    ]
    updates: dict[str, object] = {
        "oktmo_candidate_preview": "\n".join(lines),
        "oktmo_candidate_options": options,
    }
    if selected:
        updates["oktmo_candidate_ref"] = str(selected.get("code") or selected.get("value") or "")
    return {
        "candidates": candidates,
        "count": len(candidates),
        "field_updates": updates,
    }


def oktmo_add_candidate_keys(
    root: Path | str | None = None,
    values: dict[str, object] | None = None,
) -> dict[str, object]:
    project_root = Path(root or Path(__file__).resolve().parents[2]).resolve()
    data_dir = _option_data_dir(project_root, values)
    params = dict(values or {})
    code = _clean_ui_value(params.get("oktmo_candidate_ref") or params.get("candidate_code"))
    profile = candidate_key_terms(data_dir, code)
    terms = [str(item) for item in profile.get("terms") or [] if str(item).strip()]
    result = add_current_keys(project_root, terms)
    context = profile.get("context") if isinstance(profile.get("context"), dict) else {}
    kind = str(profile.get("kind") or context.get("kind") or "")
    field_updates = {
        "oktmo_selected_municipality": context.get("municipality") or "",
        "oktmo_selected_settlement": context.get("clean_name") if kind == "settlement" else "",
        "oktmo_selected_code": context.get("code") or code,
    }
    if context.get("subject_scope") or context.get("subject_ter"):
        field_updates["oktmo_scope_region"] = context.get("subject_scope") or context.get("subject_ter")
    result.update(
        {
            "kind": kind,
            "added": len(terms),
            "keys": len(load_current_keys(project_root)),
            "candidate": context,
            "field_updates": field_updates,
        }
    )
    return result


def oktmo_save_scope_keys(
    root: Path | str | None = None,
    values: dict[str, object] | None = None,
) -> dict[str, object]:
    project_root = Path(root or Path(__file__).resolve().parents[2]).resolve()
    values = values or {}
    terms: list[str] = []
    data_dir = _option_data_dir(project_root, values)
    region = _selected_oktmo_scope_value(values, "region")
    municipality = _selected_oktmo_scope_value(values, "municipality")
    if region:
        terms.extend(region_key_terms(data_dir, region))
    if municipality:
        terms.extend(municipality_key_terms(data_dir, region, municipality))
    if not terms:
        terms = load_current_keys(project_root)
    result = write_current_keys(project_root, terms)
    result.update({"saved": len(terms)})
    return result


def oktmo_load_key_file(root: Path | str | None = None) -> dict[str, object]:
    project_root = Path(root or Path(__file__).resolve().parents[2]).resolve()
    keys = load_current_keys(project_root)
    return {"path": str(current_keys_path(project_root)), "keys": len(keys), "items": keys, "preview": keys[:20]}


def oktmo_clear_key_file(root: Path | str | None = None) -> dict[str, object]:
    project_root = Path(root or Path(__file__).resolve().parents[2]).resolve()
    result = clear_current_keys(project_root)
    result.update({"items": [], "preview": []})
    return result


def oktmo_reset_project_keys(root: Path | str | None = None) -> dict[str, object]:
    project_root = Path(root or Path(__file__).resolve().parents[2]).resolve()
    result = clear_current_keys(project_root)
    result.update(
        {
            "field_updates": dict(OKTMO_RESET_FIELD_UPDATES),
            "pins_preserved": True,
        }
    )
    return result


def oktmo_open_key_file(root: Path | str | None = None) -> dict[str, object]:
    project_root = Path(root or Path(__file__).resolve().parents[2]).resolve()
    path = current_keys_path(project_root)
    path.parent.mkdir(parents=True, exist_ok=True)
    if not path.exists():
        path.write_text("", encoding="utf-8")
    if os.name == "nt":
        os.startfile(str(path))  # type: ignore[attr-defined]
    return {"path": str(path), "keys": len(load_current_keys(project_root))}


def oktmo_pin_file_status(root: Path | str | None = None) -> dict[str, object]:
    project_root = Path(root or Path(__file__).resolve().parents[2]).resolve()
    return pin_file_status(project_root)


def _oktmo_pin_backup_dir(project_root: Path) -> Path:
    return project_root / "output" / "oktmo_pins"


def _latest_oktmo_pin_import_candidate(project_root: Path) -> Path | None:
    bundle_name = pins_bundle_path(project_root).name
    preferred = [
        project_root / "input" / bundle_name,
        project_root / bundle_name,
    ]
    for path in preferred:
        if path.exists() and path.is_file():
            return path
    backup_dir = _oktmo_pin_backup_dir(project_root)
    candidates = [path for path in backup_dir.glob("oktmo_pins_bundle*.json") if path.is_file()]
    if candidates:
        return max(candidates, key=lambda path: path.stat().st_mtime_ns)
    default_path = pins_bundle_path(project_root)
    return default_path if default_path.exists() else None


def oktmo_export_pin_file(root: Path | str | None = None) -> dict[str, object]:
    project_root = Path(root or Path(__file__).resolve().parents[2]).resolve()
    result = export_pins_bundle(project_root)
    export_dir = _oktmo_pin_backup_dir(project_root)
    export_dir.mkdir(parents=True, exist_ok=True)
    export_path = export_dir / f"oktmo_pins_bundle_{time.strftime('%Y%m%d_%H%M%S')}.json"
    export_result = export_pins_bundle(project_root, export_path)
    export_result["bundle_path"] = result.get("path", "")
    return export_result


def oktmo_import_pin_file(root: Path | str | None = None) -> dict[str, object]:
    project_root = Path(root or Path(__file__).resolve().parents[2]).resolve()
    source = _latest_oktmo_pin_import_candidate(project_root)
    return import_pins_bundle(project_root, source, merge=True)


def oktmo_clear_pin_file(root: Path | str | None = None) -> dict[str, object]:
    project_root = Path(root or Path(__file__).resolve().parents[2]).resolve()
    return clear_all_pins(project_root)


def oktmo_open_pin_file(root: Path | str | None = None) -> dict[str, object]:
    project_root = Path(root or Path(__file__).resolve().parents[2]).resolve()
    path = pins_bundle_path(project_root)
    if not path.exists():
        export_pins_bundle(project_root, path)
    if os.name == "nt":
        os.startfile(str(path))  # type: ignore[attr-defined]
    status = pin_file_status(project_root)
    status["path"] = str(path)
    return status


def _latest_reference_workbook(root: Path) -> Path | None:
    candidates: list[Path] = []
    for folder in (root / "input", root / "output"):
        if not folder.exists():
            continue
        candidates.extend(path for path in folder.glob("AddressReference*.xlsx") if path.is_file() and not path.name.startswith("~$"))
    if not candidates:
        return None
    return max(candidates, key=lambda path: path.stat().st_mtime)


def _latest_xlsx_in_folder(folder: Path) -> Path | None:
    if not folder.exists():
        return None
    candidates = sorted(path for path in folder.glob("*.xlsx") if path.is_file() and not path.name.startswith("~$"))
    if not candidates:
        return None
    return max(candidates, key=lambda path: path.stat().st_mtime)


def _latest_address_collection_workbook(folder: Path, *, exclude: tuple[Path | None, ...] = ()) -> Path | None:
    if not folder.exists():
        return None
    excluded: set[str] = set()
    for path in exclude:
        if not path:
            continue
        try:
            excluded.add(str(path.resolve()))
        except OSError:
            excluded.add(str(path))
    candidates = [
        path
        for path in folder.glob("AddressCollection*.xlsx")
        if path.is_file() and not path.name.startswith("~$") and str(path.resolve()) not in excluded
    ]
    if not candidates:
        candidates = [
            path
            for path in folder.glob("*.xlsx")
            if path.is_file() and not path.name.startswith("~$") and str(path.resolve()) not in excluded
        ]
    if not candidates:
        return None
    return max(candidates, key=lambda path: path.stat().st_mtime)


def _reference_header_options(first_row: tuple[object, ...], sheet_name: str) -> list[dict[str, object]]:
    options: list[dict[str, object]] = []
    seen_values: set[str] = {""}
    for index, value in enumerate(first_row):
        text = str(value or "").replace("\xa0", " ").strip()
        if not text or text.lower() == "nan":
            continue
        letter = excel_column_letter(index)
        if text in seen_values:
            continue
        seen_values.add(text)
        label = f"{letter} · {text}"
        options.append(
            {
                "value": text,
                "label": label,
                "label_ru": label,
                "title": f"{sheet_name}!{letter}",
            }
        )
    return options


def _cleanup_input_office_temp(context: JobContext) -> tuple[CleanupRecord, ...]:
    removed = delete_office_temp_files(context.paths.input, log=context.log)
    if removed:
        context.log(f"Office temporary files removed: {len(removed)}")
    return removed


def _unique_path(path: Path) -> Path:
    if not path.exists():
        return path
    for index in range(2, 1000):
        candidate = path.with_name(f"{path.stem}_{index}{path.suffix}")
        if not candidate.exists():
            return candidate
    raise RuntimeError(f"Could not allocate a unique output name for {path.name}")


def _is_inside(child: Path, parent: Path) -> bool:
    try:
        child_resolved = str(child.resolve())
        parent_resolved = str(parent.resolve())
    except OSError:
        return False
    try:
        common = os.path.commonpath([child_resolved, parent_resolved])
    except ValueError:
        return False
    return common == parent_resolved


def _clean_managed_folder(context: JobContext, folder: Path, label: str) -> dict[str, object]:
    root = context.paths.root.resolve()
    if folder.is_symlink():
        raise RuntimeError(f"{label} is a symbolic link. Cleanup blocked for safety.")

    folder.mkdir(parents=True, exist_ok=True)
    folder_resolved = folder.resolve()
    if not _is_inside(folder_resolved, root):
        raise RuntimeError(f"{label} path is outside project root. Cleanup blocked.")

    removed = 0
    skipped: list[str] = []
    for item in folder.iterdir():
        if item.name == ".gitkeep":
            continue
        try:
            if item.is_symlink():
                item.unlink()
            elif item.is_dir():
                if not _is_inside(item, folder_resolved):
                    skipped.append(f"{item.name} (escapes {label})")
                    continue
                shutil.rmtree(item)
            else:
                item.unlink()
            removed += 1
            context.log(f"Removed from {label}: {item.name}")
        except OSError as exc:
            skipped.append(f"{item.name} ({exc})")

    return {"folder": label, "removed_items": removed, "skipped_items": skipped}


class _LogStream:
    def __init__(self, context: JobContext) -> None:
        self.context = context
        self.buffer = ""
        self.last_progress_emit = 0.0

    def writable(self) -> bool:
        return True

    def write(self, text: str) -> int:
        if not text:
            return 0

        for char in text:
            if char == "\n":
                self._emit(self.buffer)
                self.buffer = ""
            elif char == "\r":
                self._emit_progress_snapshot()
                self.buffer = ""
            else:
                self.buffer += char
        return len(text)

    def flush(self) -> None:
        self._emit(self.buffer)
        self.buffer = ""

    def _emit(self, text: str) -> None:
        line = text.rstrip()
        if line:
            self.context.log(line)

    def _emit_progress_snapshot(self) -> None:
        now = time.monotonic()
        if now - self.last_progress_emit < 0.5:
            return
        self.last_progress_emit = now
        self._emit(self.buffer)


@contextlib.contextmanager
def _relay_output(context: JobContext):
    stream = _LogStream(context)
    with contextlib.redirect_stdout(stream), contextlib.redirect_stderr(stream):
        yield
    stream.flush()


def validate_input(context: JobContext) -> dict[str, object]:
    context.paths.input.mkdir(parents=True, exist_ok=True)
    removed = _cleanup_input_office_temp(context)
    files = _xlsx_files(context.paths.input)
    context.log(f"Input folder: {context.paths.input}")
    if not files:
        context.log("No .xlsx files found.")
    else:
        context.log(f"Found {len(files)} .xlsx file(s):")
        for file_path in files:
            context.log(f"  - {file_path.name}")
    context.progress(1.0)
    return {"files": [str(path) for path in files], "count": len(files), "removed_temp_files": len(removed)}


def oktmo_key_console(context: JobContext) -> dict[str, object]:
    """Summarize the project OKTMO key workspace when the child command is run."""
    parameters = context.operation.parameters
    status = oktmo_data_status(context.paths.root, parameters, {})
    keys = load_current_keys(context.paths.root)
    pins = pin_file_status(context.paths.root)
    context.log("OKTMO project keys")
    context.log(f"Registry: {'ready' if status.get('ok') else 'missing'}")
    context.log(f"Data folder: {status.get('rosstat_dir') or '-'}")
    context.log(f"Latest data file: {status.get('latest_name') or '-'}")
    context.log(f"Current keys: {len(keys)}")
    context.log(f"Pin file: {pins.get('path') or '-'}")
    context.log(f"Region pins: {pins.get('regions', 0)}")
    context.log(f"Municipality pins: {pins.get('municipalities', 0)}")
    context.progress(1.0)
    return {
        "status": status,
        "keys": len(keys),
        "key_file": str(current_keys_path(context.paths.root)),
        "pins": pins,
    }


def _header_key(value: object) -> str:
    import re

    text = str(value or "").lower().replace("ё", "е")
    return re.sub(r"[^0-9a-zа-я]+", "", text)


def _column_header_to_letter(workbook: Path, header_name: str) -> str:
    needle = _header_key(header_name)
    if not needle:
        raise ValueError("Column header is empty.")
    from openpyxl import load_workbook

    book = load_workbook(workbook, read_only=True, data_only=True)
    try:
        sheet = book.worksheets[0]
        first_row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
        for index, value in enumerate(first_row):
            key = _header_key(value)
            if key and (key == needle or needle in key or key in needle):
                return excel_column_letter(index)
    finally:
        book.close()
    raise ValueError(f"Header {header_name!r} was not found in the first row of {workbook.name}.")


def _aligner_runtime_config(context: JobContext, file_path: Path | None = None) -> dict[str, object] | None:
    parameters = context.operation.parameters
    relevant_keys = set(ALIGNER_RUNTIME_PARAMETER_IDS)
    if not any(key in parameters for key in relevant_keys):
        return None

    ground_truth_column = str(parameters.get("ground_truth_column", "") or "").strip()
    ground_truth_header = str(parameters.get("ground_truth_column_header", "") or "").strip()
    target_columns = str(parameters.get("target_columns", "") or "").strip()
    target_header = str(parameters.get("target_column_header", "") or "").strip()
    alignment_mode = _normalized_alignment_mode(parameters)
    companion_mode = str(parameters.get("companion_mode", "auto") or "auto").strip()
    satellite_columns = str(parameters.get("satellite_columns", "") or "").strip()
    normalize_after_match = (
        _bool_parameter(parameters, "normalize_after_match", False)
        or _alignment_mode_requests_post_normalization(alignment_mode)
    )
    if file_path:
        if not ground_truth_column and ground_truth_header:
            ground_truth_column = _column_header_to_letter(file_path, ground_truth_header)
        if not target_columns and target_header:
            target_columns = _column_header_to_letter(file_path, target_header)
    runtime_config = {
        "ground_truth_column": ground_truth_column,
        "target_columns": target_columns,
        "alignment_mode": alignment_mode,
        "companion_mode": companion_mode,
        "satellite_columns": satellite_columns,
        "normalize_before_match": _bool_parameter(parameters, "normalize_before_match", False),
        "normalize_after_match": normalize_after_match,
        "collect_before_match": _bool_parameter(parameters, "collect_before_match", False),
        "collect_whole_document": _bool_parameter(parameters, "collect_whole_document", False),
    }
    context.log(
        "Runtime column selection: "
        f"ground_truth_column={ground_truth_column or '(auto)'}, "
        f"target_columns={target_columns or '(auto)'}, "
        f"alignment_mode={alignment_mode}, "
        f"companion_mode={companion_mode or 'auto'}, "
        f"satellite_columns={satellite_columns or '(auto)'}"
    )
    if runtime_config["normalize_before_match"] or runtime_config["normalize_after_match"] or runtime_config["collect_before_match"]:
        address_runtime = _build_address_runtime_context(context, parameters)
        runtime_config.update(
            {
                "city": str(parameters.get("city") or CONFIG.get("city") or "").strip(),
                "use_oktmo": _bool_parameter(parameters, "use_oktmo", True),
                "oktmo_data_dir": str(address_runtime.normalizer.data_dir),
                "subject_ter_hint": address_runtime.subject_ter_hint or "",
                "municipality_hint": address_runtime.municipality_hint or "",
            }
        )
        context.log("Address matching normalization/post-processing runtime: enabled")
        _log_address_runtime_context(context, address_runtime)
    else:
        context.log("Address matching normalization: off")
    return runtime_config


def _normalized_alignment_mode(parameters: dict[str, object]) -> str:
    mode = str(parameters.get("alignment_mode") or "fast").strip().lower()
    aliases = {
        "normalization": "normalized",
        "normalize": "normalized",
        "post_normalize": "normalized",
        "post-normalize": "normalized",
        "alignment_normalization": "normalized",
        "alignment+normalization": "normalized",
    }
    return aliases.get(mode, mode or "fast")


def _alignment_mode_requests_post_normalization(mode: str) -> bool:
    return mode in {"normalized", "normalised"}


def _bool_parameter(parameters: dict[str, object], key: str, default: bool) -> bool:
    value = parameters.get(key, default)
    if isinstance(value, bool):
        return value
    return str(value or "").strip().lower() in {"1", "true", "yes", "on", "да"}


def _oktmo_scope_parameters(context: JobContext, parameters: dict[str, object]) -> tuple[str | None, str | None, bool]:
    data_dir = _optional_project_path(context, parameters.get("oktmo_data_dir")) or (context.paths.root / "data")
    manual_subject_ter = str(parameters.get("subject_ter_hint") or "").strip() or None
    manual_municipality = str(parameters.get("municipality_hint") or "").strip() or None
    if not _bool_parameter(parameters, "oktmo_scope_enabled", False):
        return manual_subject_ter, manual_municipality, False

    scope_region = str(parameters.get("oktmo_scope_region") or "").strip()
    scope_municipality = str(parameters.get("oktmo_scope_municipality") or "").strip()
    subject_ter = resolve_region_scope_id(data_dir, scope_region) or resolve_region_scope_ter(data_dir, scope_region) or manual_subject_ter
    municipality = resolve_municipality_scope_hint(data_dir, scope_municipality) or manual_municipality
    return subject_ter, municipality, True


def _build_address_runtime_context(context: JobContext, parameters: dict[str, object]) -> AddressRuntimeContext:
    oktmo_data_dir = _optional_project_path(context, parameters.get("oktmo_data_dir"))
    subject_ter_hint, municipality_hint, oktmo_scope_enabled = _oktmo_scope_parameters(context, parameters)
    normalizer = AddressNormalizer(
        city_name=str(parameters.get("city") or CONFIG.get("city") or "").strip(),
        data_dir=oktmo_data_dir,
        use_oktmo=_bool_parameter(parameters, "use_oktmo", True),
        subject_ter_hint=subject_ter_hint,
        municipality_hint=municipality_hint,
    )
    return AddressRuntimeContext(
        normalizer=normalizer,
        subject_ter_hint=subject_ter_hint,
        municipality_hint=municipality_hint,
        oktmo_scope_enabled=oktmo_scope_enabled,
    )


def _log_address_runtime_context(context: JobContext, runtime: AddressRuntimeContext) -> None:
    context.log(f"OKTMO data root: {runtime.normalizer.data_dir}")
    context.log(
        "OKTMO search scope: "
        f"{'enabled' if runtime.oktmo_scope_enabled else 'manual/off'}, "
        f"scope={runtime.subject_ter_hint or '-'}, municipality={runtime.municipality_hint or '-'}"
    )


def _oktmo_scope_runtime_summary(runtime: AddressRuntimeContext) -> dict[str, object]:
    return {
        "enabled": runtime.oktmo_scope_enabled,
        "use_oktmo": bool(runtime.normalizer.use_oktmo),
        "data_dir": str(runtime.normalizer.data_dir),
        "subject_ter_hint": runtime.subject_ter_hint,
        "municipality_hint": runtime.municipality_hint,
    }


def _address_runtime_summary(runtime: AddressRuntimeContext) -> dict[str, object]:
    key_profile = normalizer_oktmo_key_profile(runtime.normalizer)
    normalizer_report = normalizer_runtime_report(runtime.normalizer, key_profile)
    return {
        "data_dir": str(runtime.normalizer.data_dir),
        "use_oktmo": bool(runtime.normalizer.use_oktmo),
        "city_name": runtime.normalizer.city_name,
        "oktmo_scope_enabled": runtime.oktmo_scope_enabled,
        "subject_ter_hint": runtime.subject_ter_hint,
        "municipality_hint": runtime.municipality_hint,
        "oktmo_scope": _oktmo_scope_runtime_summary(runtime),
        "oktmo_key_profile": key_profile,
        "normalizer": normalizer_report,
    }


def _slot_runtime_summary(enabled_slots: tuple[str, ...], slot_order: tuple[str, ...]) -> dict[str, object]:
    return {
        "enabled_slots": list(enabled_slots),
        "slot_order": list(slot_order),
    }


def _artifact_summary(
    *,
    outputs: object = None,
    reports: object = None,
    staging: object = None,
    inputs: object = None,
) -> dict[str, list[str]]:
    return {
        "outputs": _artifact_paths(outputs),
        "reports": _artifact_paths(reports),
        "staging": _artifact_paths(staging),
        "inputs": _artifact_paths(inputs),
    }


def _artifact_paths(value: object) -> list[str]:
    if value in (None, ""):
        return []
    if isinstance(value, (str, Path)):
        return [str(value)] if str(value) else []
    if isinstance(value, dict):
        return _artifact_paths(list(value.values()))
    if isinstance(value, (list, tuple, set, frozenset)):
        result: list[str] = []
        for item in value:
            result.extend(_artifact_paths(item))
        return list(dict.fromkeys(result))
    return [str(value)]


def _optional_project_path(context: JobContext, value: object) -> Path | None:
    text = str(value or "").strip().strip('"')
    if not text:
        return None
    path = Path(os.path.expandvars(text)).expanduser()
    if not path.is_absolute():
        path = context.paths.root / path
    return path


def _source_root_parameter(context: JobContext, value: object) -> Path:
    text = str(value or "").strip().strip('"')
    if not text:
        return context.paths.input
    path = Path(os.path.expandvars(text)).expanduser()
    if path.is_absolute():
        resolved = path
    else:
        root_candidate = context.paths.root / path
        input_candidate = context.paths.input / path
        resolved = root_candidate if root_candidate.exists() else input_candidate
    if not resolved.exists():
        raise RuntimeError(f"Source folder/file was not found: {resolved}")
    return resolved


def _project_folder_parameter(context: JobContext, value: object, *, default_name: str) -> Path:
    text = str(value or "").strip().strip('"')
    if not text:
        return context.paths.root / default_name
    path = Path(os.path.expandvars(text)).expanduser()
    if not path.is_absolute():
        path = context.paths.root / path
    return path


def _slot_selection_parameter(parameters: dict[str, object]) -> tuple[str, ...]:
    value = parameters.get("enabled_slots")
    if isinstance(value, (list, tuple)):
        selected = [str(item).strip() for item in value if str(item).strip()]
    else:
        selected = [part.strip() for part in str(value or "").split(",") if part.strip()]
    allowed = set(DEFAULT_ENABLED_SLOTS)
    filtered = tuple(slot for slot in selected if slot in allowed)
    return filtered or DEFAULT_SELECTED_SLOTS


def _coerce_slot_position(value: object) -> int | None:
    try:
        number = int(str(value).strip())
    except (TypeError, ValueError):
        return None
    return number if number > 0 else None


def _slot_order_from_mapping(value: dict[object, object], enabled_slots: tuple[str, ...]) -> tuple[str, ...]:
    enabled = set(enabled_slots)
    positions: dict[str, int] = {}
    for raw_slot, raw_position in value.items():
        slot = str(raw_slot).strip()
        if slot not in DEFAULT_SLOT_ORDER or slot not in enabled:
            continue
        position = _coerce_slot_position(raw_position)
        if position is not None:
            positions[slot] = position
    if not positions:
        return tuple(slot for slot in DEFAULT_SLOT_ORDER if slot in enabled)
    return tuple(
        slot
        for slot in sorted(
            (slot for slot in DEFAULT_SLOT_ORDER if slot in enabled),
            key=lambda item: (positions.get(item, len(DEFAULT_SLOT_ORDER) + DEFAULT_SLOT_ORDER.index(item)), DEFAULT_SLOT_ORDER.index(item)),
        )
    )


def _slot_order_parameter(parameters: dict[str, object], enabled_slots: tuple[str, ...]) -> tuple[str, ...]:
    value = parameters.get("slot_order")
    enabled = set(enabled_slots)
    if isinstance(value, dict):
        return _slot_order_from_mapping(value, enabled_slots)
    if isinstance(value, (list, tuple)):
        ordered = [str(item).strip() for item in value if str(item).strip() in DEFAULT_SLOT_ORDER]
        ordered.extend(slot for slot in DEFAULT_SLOT_ORDER if slot not in ordered)
        return tuple(slot for slot in ordered if slot in enabled)

    text = str(value or "").strip()
    if text:
        if text.startswith("{"):
            try:
                loaded = json.loads(text)
            except json.JSONDecodeError:
                loaded = None
            if isinstance(loaded, dict):
                return _slot_order_from_mapping(loaded, enabled_slots)
        ordered = [part.strip() for part in text.split(",") if part.strip() in DEFAULT_SLOT_ORDER]
        if ordered:
            ordered.extend(slot for slot in DEFAULT_SLOT_ORDER if slot not in ordered)
            return tuple(slot for slot in ordered if slot in enabled)

    per_slot_positions: dict[str, int] = {}
    for slot in DEFAULT_SLOT_ORDER:
        position = _coerce_slot_position(parameters.get(f"slot_order_{slot}"))
        if position is not None:
            per_slot_positions[slot] = position
    if per_slot_positions:
        return _slot_order_from_mapping(per_slot_positions, enabled_slots)
    return tuple(slot for slot in DEFAULT_SLOT_ORDER if slot in enabled)


def run_address_collection(context: JobContext) -> dict[str, object]:
    context.paths.input.mkdir(parents=True, exist_ok=True)
    context.paths.output.mkdir(parents=True, exist_ok=True)
    removed_temp: tuple[CleanupRecord, ...] = ()
    parameters = context.operation.parameters
    action = str(parameters.get("address_collection_action", "generate") or "generate").strip().lower()
    if action == "new":
        action = "generate"
    if action == "append":
        action = "update"
    if action not in {"generate", "update"}:
        raise ValueError(f"Unsupported address collection action: {action}")

    report_file = context.report_dir / "address_collection_summary.json"
    source_root_value = parameters.get("source_root")
    source_root = context.paths.input
    collection_dir = _project_folder_parameter(context, parameters.get("address_collection_dir"), default_name="address_collection")
    collection_dir.mkdir(parents=True, exist_ok=True)
    explicit_target_file = _optional_project_path(context, parameters.get("reference_file"))
    target_file = explicit_target_file if explicit_target_file and explicit_target_file.exists() else None
    if action == "update" and target_file is None:
        target_file = _latest_xlsx_in_folder(collection_dir)
    timestamp = time.strftime("%Y%m%d_%H%M%S")
    clean_output_workbook = _bool_parameter(parameters, "clean_output_workbook", False)
    if clean_output_workbook:
        output_file = _unique_path(collection_dir / f"AddressCollection_clean_{timestamp}.xlsx")
    elif target_file and target_file.exists():
        output_file = _unique_path(collection_dir / f"{target_file.stem}_added_{timestamp}.xlsx")
    else:
        output_file = _unique_path(collection_dir / f"AddressCollection_{timestamp}.xlsx")
    staging_file = _unique_path(context.paths.output / f"AddressAssembly_{action}_{timestamp}.xlsx")
    address_runtime = _build_address_runtime_context(context, parameters)
    normalizer = address_runtime.normalizer
    enabled_slots = _slot_selection_parameter(parameters)
    slot_order = _slot_order_parameter(parameters, enabled_slots)
    target_column = str(parameters.get("target_column") or "").strip()
    target_column_header = str(parameters.get("target_column_header") or "").strip()
    target_column = target_column or target_column_header or "A"
    collected_file = _optional_project_path(context, parameters.get("collected_file") or parameters.get("source_collection_file"))
    collected_column = str(parameters.get("collected_column") or "").strip() or "A"
    if action == "update":
        if collected_file and not collected_file.exists():
            raise RuntimeError(f"Collected address workbook was not found: {collected_file}")
        if not collected_file:
            collected_file = _latest_address_collection_workbook(collection_dir, exclude=(target_file,))
        if not collected_file:
            raise RuntimeError(
                "Clean collected address workbook was not found. "
                "Run 'New table' first or choose a collected XLSX workbook for append."
            )
        if target_file and collected_file.resolve() == target_file.resolve():
            raise RuntimeError("Collected workbook and ready target workbook must be different files for append.")

        append_options = AddressAppendOptions(
            source_file=collected_file,
            target_file=target_file,
            source_column=collected_column,
            target_column=target_column,
            clean_duplicates=_bool_parameter(parameters, "clean_duplicates", True),
            clean_output_workbook=clean_output_workbook,
        )

        context.log("Address collection screen")
        context.log("Action: append clean collected column")
        context.log(f"Collected workbook: {collected_file}")
        context.log(f"Collected column: {append_options.source_column or 'A'}")
        context.log(f"Ready target workbook: {target_file or '(new workbook)'}")
        context.log(f"Target column: {target_column}")
        context.log(f"Clean output workbook: {'yes' if append_options.clean_output_workbook else 'no'}")
        context.log(f"Output workbook: {output_file}")
        context.log(f"Clean duplicates on append: {'yes' if append_options.clean_duplicates else 'no'}")
        _log_address_runtime_context(context, address_runtime)
        result = append_collected_address_workbook(
            output_file=output_file,
            normalizer=normalizer,
            options=append_options,
            report_file=report_file,
            log=context.log,
            progress=context.progress,
            cancelled=context.cancelled,
        )

        benchmark_result: dict[str, object] | None = None
        benchmark_report_path: Path | None = None
        benchmark_file = _optional_project_path(context, parameters.get("benchmark_file"))
        benchmark_column = str(parameters.get("benchmark_column") or "Местоположение").strip() or "Местоположение"
        if benchmark_file:
            if not benchmark_file.exists():
                context.log(f"Benchmark workbook not found: {benchmark_file}")
            else:
                benchmark_report_path = context.report_dir / "address_collection_benchmark.json"
                benchmark_result = benchmark_address_collection(
                    collected_file=result.output_file,
                    benchmark_file=benchmark_file,
                    normalizer=normalizer,
                    benchmark_column=benchmark_column,
                    collected_column=result.target_column,
                    report_file=benchmark_report_path,
                    enabled_slots=enabled_slots,
                    slot_order=slot_order,
                )
                context.log(
                    "Benchmark: "
                    f"{benchmark_result['matched_unique_key_count']}/"
                    f"{benchmark_result['benchmark_unique_key_count']} unique address core(s) matched "
                    f"(coverage {float(benchmark_result['coverage_ratio']) * 100:.1f}%, "
                    f"precision {float(benchmark_result['precision_ratio']) * 100:.1f}%)."
                )
                context.log(f"Benchmark report: {benchmark_report_path}")

        context.progress(1.0)
        context.log(
            "Address append ready: "
            f"{result.appended_rows} row(s) appended to "
            f"{result.target_sheet}!{result.target_column}{result.target_start_row}, "
            f"{result.duplicate_rows} duplicate(s), "
            f"{result.skipped_rows} skipped."
        )
        context.log(f"Output workbook: {result.output_file}")
        context.log(f"Report: {result.report_file}")
        return {
            "output": str(result.output_file),
            "staging": None,
            "report": str(result.report_file),
            "sources": [str(result.source_file)],
            "issues": [],
            "scanned_values": result.source_rows,
            "plausible_values": result.source_rows - result.skipped_rows,
            "accepted_candidates": 0,
            "normalized_rows": result.accepted_rows,
            "duplicate_rows": result.duplicate_rows,
            "rejected_rows": result.skipped_rows,
            "appended_rows": result.appended_rows,
            "target_sheet": result.target_sheet,
            "target_column": result.target_column,
            "target_start_row": result.target_start_row,
            "benchmark": benchmark_result,
            "address_runtime": _address_runtime_summary(address_runtime),
            "removed_temp_files": len(removed_temp),
            "artifacts": _artifact_summary(
                outputs=result.output_file,
                reports=[result.report_file, benchmark_report_path],
                inputs=result.source_file,
            ),
        }

    removed_temp = _cleanup_input_office_temp(context)
    source_root = _source_root_parameter(context, source_root_value)
    options = AddressAssemblyOptions(
        action=action,
        source_root=source_root,
        source_columns=str(parameters.get("source_columns") or "").strip() or None,
        reference_columns=target_column,
        reference_file=target_file,
        target_dir=collection_dir,
        target_file=target_file,
        target_column=target_column,
        staging_file=staging_file,
        enabled_slots=enabled_slots,
        slot_order=slot_order,
        omit_microdistrict_when_street_found=_bool_parameter(parameters, "omit_microdistrict_when_street_found", True),
        include_slots=_bool_parameter(parameters, "include_slots", True),
        clean_output_workbook=clean_output_workbook,
        clean_duplicates=_bool_parameter(parameters, "clean_duplicates", True),
        sanitize_incomplete=_bool_parameter(parameters, "sanitize_incomplete", True),
        write_rejected=_bool_parameter(parameters, "write_rejected", True),
    )

    context.log("Address collection screen")
    context.log(f"Action: {action}")
    context.log(f"Source root: {source_root}")
    context.log(f"Address collection folder: {collection_dir}")
    context.log(f"Target workbook: {target_file or '(new workbook)'}")
    context.log(f"Target column: {target_column}")
    context.log(f"Staging workbook: {staging_file}")
    context.log(f"Output workbook: {output_file}")
    context.log(f"Columns: {options.source_columns or '(auto scan)'}")
    context.log(f"Existing address column: {options.reference_columns or '(auto detect)'}")
    context.log(f"Clean output workbook: {'yes' if options.clean_output_workbook else 'no'}")
    context.log(f"Enabled slots: {', '.join(options.enabled_slots)}")
    context.log(f"Slot order: {', '.join(options.slot_order)}")
    context.log(f"Clean duplicates: {'yes' if options.clean_duplicates else 'no'}")
    context.log(f"Sanitize incomplete addresses: {'yes' if options.sanitize_incomplete else 'no'}")
    _log_address_runtime_context(context, address_runtime)
    result = assemble_address_workbook(
        input_dir=context.paths.input,
        output_file=output_file,
        normalizer=normalizer,
        options=options,
        report_file=report_file,
        log=context.log,
        progress=context.progress,
        cancelled=context.cancelled,
    )
    benchmark_result: dict[str, object] | None = None
    benchmark_report_path: Path | None = None
    benchmark_file = _optional_project_path(context, parameters.get("benchmark_file"))
    benchmark_column = str(parameters.get("benchmark_column") or "Местоположение").strip() or "Местоположение"
    if benchmark_file:
        if not benchmark_file.exists():
            context.log(f"Benchmark workbook not found: {benchmark_file}")
        else:
            benchmark_report_path = context.report_dir / "address_collection_benchmark.json"
            benchmark_result = benchmark_address_collection(
                collected_file=result.output_file,
                benchmark_file=benchmark_file,
                normalizer=normalizer,
                benchmark_column=benchmark_column,
                collected_column=result.target_column,
                report_file=benchmark_report_path,
                enabled_slots=enabled_slots,
                slot_order=slot_order,
            )
            context.log(
                "Benchmark: "
                f"{benchmark_result['matched_unique_key_count']}/"
                f"{benchmark_result['benchmark_unique_key_count']} unique address core(s) matched "
                f"(coverage {float(benchmark_result['coverage_ratio']) * 100:.1f}%, "
                f"precision {float(benchmark_result['precision_ratio']) * 100:.1f}%)."
            )
            context.log(f"Benchmark report: {benchmark_report_path}")
    context.progress(1.0)
    context.log(
        "Address collection ready: "
        f"{result.normalized_rows} row(s), "
        f"{result.appended_rows} appended to {result.target_sheet}!{result.target_column}{result.target_start_row}, "
        f"{result.accepted_candidates} accepted candidate(s), "
        f"{result.duplicate_rows} duplicate(s) merged, "
        f"{result.rejected_rows} rejected, "
        f"{result.conflict_rows} OKTMO conflict feedback row(s)."
    )
    context.log(f"Staging workbook: {result.staging_file}")
    context.log(f"Output workbook: {result.output_file}")
    context.log(f"Report: {result.report_file}")
    if result.issues:
        context.log(f"Source intake issues: {len(result.issues)}")
        for issue in result.issues[:30]:
            context.log(f"  issue: {issue}")
    return {
        "output": str(result.output_file),
        "staging": str(result.staging_file),
        "report": str(result.report_file),
        "sources": [str(path) for path in result.source_files],
        "issues": list(result.issues),
        "scanned_values": result.scanned_values,
        "plausible_values": result.plausible_values,
        "accepted_candidates": result.accepted_candidates,
        "normalized_rows": result.normalized_rows,
        "duplicate_rows": result.duplicate_rows,
        "rejected_rows": result.rejected_rows,
        "conflict_rows": result.conflict_rows,
        "appended_rows": result.appended_rows,
        "target_sheet": result.target_sheet,
        "target_column": result.target_column,
        "target_start_row": result.target_start_row,
        "benchmark": benchmark_result,
        "address_runtime": _address_runtime_summary(address_runtime),
        "slot_runtime": _slot_runtime_summary(enabled_slots, slot_order),
        "removed_temp_files": len(removed_temp),
        "artifacts": _artifact_summary(
            outputs=result.output_file,
            reports=[result.report_file, benchmark_report_path],
            staging=result.staging_file,
            inputs=result.source_files,
        ),
    }


def run_reference_normalization(context: JobContext) -> dict[str, object]:
    context.paths.input.mkdir(parents=True, exist_ok=True)
    context.paths.output.mkdir(parents=True, exist_ok=True)
    removed_temp = _cleanup_input_office_temp(context)
    parameters = context.operation.parameters
    action = str(parameters.get("reference_action", "generate") or "generate").strip().lower()
    if action not in {"generate", "update", "clean"}:
        raise ValueError(f"Unsupported reference action: {action}")

    timestamp = time.strftime("%Y%m%d_%H%M%S")
    output_file = _unique_path(context.paths.output / f"AddressReference_{action}_{timestamp}.xlsx")
    staging_file = _unique_path(context.paths.output / f"AddressReference_{action}_{timestamp}.stage.xlsx")
    report_file = context.report_dir / "reference_normalization_summary.json"
    reference_file = _optional_project_path(context, parameters.get("reference_file"))
    address_runtime = _build_address_runtime_context(context, parameters)
    normalizer = address_runtime.normalizer
    enabled_slots = _slot_selection_parameter(parameters)
    slot_order = _slot_order_parameter(parameters, enabled_slots)
    reference_columns = str(parameters.get("reference_columns") or "").strip()
    reference_column_header = str(parameters.get("reference_column_header") or "").strip()
    options = ReferenceOptions(
        action=action,
        source_columns=str(parameters.get("source_columns") or "").strip() or None,
        reference_columns=reference_columns or reference_column_header or None,
        reference_file=reference_file,
        enabled_slots=enabled_slots,
        slot_order=slot_order,
        omit_microdistrict_when_street_found=_bool_parameter(parameters, "omit_microdistrict_when_street_found", True),
        include_slots=_bool_parameter(parameters, "include_slots", True),
        clean_output_workbook=_bool_parameter(parameters, "clean_output_workbook", False),
        clean_old_values=_bool_parameter(parameters, "clean_old_values", True),
        clean_old_slots=_bool_parameter(parameters, "clean_old_slots", True),
    )

    context.log("Reference normalization screen")
    context.log(f"Action: {action}")
    context.log(f"Input folder: {context.paths.input}")
    context.log(f"Output workbook: {output_file}")
    context.log(f"Staging workbook: {staging_file}")
    context.log(f"Columns: {options.source_columns or '(auto scan)'}")
    context.log(f"Reference column: {options.reference_columns or '(auto detect)'}")
    context.log(f"Clean output workbook: {'yes' if options.clean_output_workbook else 'no'}")
    context.log(f"Enabled slots: {', '.join(options.enabled_slots)}")
    context.log(f"Slot order: {', '.join(options.slot_order)}")
    context.log(
        "Omit microdistrict/quartal when street is found: "
        f"{'yes' if options.omit_microdistrict_when_street_found else 'no'}"
    )
    _log_address_runtime_context(context, address_runtime)
    if reference_file:
        context.log(f"Existing reference: {reference_file}")

    result = build_reference_workbook(
        input_dir=context.paths.input,
        output_file=output_file,
        normalizer=normalizer,
        options=options,
        staging_file=staging_file,
        report_file=report_file,
    )
    context.progress(1.0)
    context.log(
        "Reference workbook ready: "
        f"{result.normalized_rows} row(s), "
        f"{result.raw_candidates} candidate value(s), "
        f"{result.duplicate_rows} duplicate(s) merged."
    )
    if result.issues:
        context.log(f"Source intake issues: {len(result.issues)}")
        for issue in result.issues[:30]:
            context.log(f"  issue: {issue}")
    for source in result.source_files:
        context.log(f"  source: {source}")
    if result.staging_file:
        context.log(f"Staging workbook: {result.staging_file}")
    if result.report_file:
        context.log(f"Report: {result.report_file}")
    return {
        "output": str(result.output_file),
        "staging": str(result.staging_file) if result.staging_file else None,
        "report": str(result.report_file) if result.report_file else None,
        "sources": [str(path) for path in result.source_files],
        "issues": list(result.issues),
        "raw_candidates": result.raw_candidates,
        "normalized_rows": result.normalized_rows,
        "duplicate_rows": result.duplicate_rows,
        "address_runtime": _address_runtime_summary(address_runtime),
        "slot_runtime": _slot_runtime_summary(enabled_slots, slot_order),
        "removed_temp_files": len(removed_temp),
        "artifacts": _artifact_summary(
            outputs=result.output_file,
            reports=result.report_file,
            staging=result.staging_file,
            inputs=result.source_files,
        ),
    }


def _reference_processing_input_file(context: JobContext, value: object) -> Path:
    explicit = _optional_project_path(context, value)
    if explicit:
        if not explicit.exists():
            raise RuntimeError(f"Reference workbook was not found: {explicit}")
        return explicit
    latest_reference = _latest_reference_workbook(context.paths.root)
    if latest_reference:
        return latest_reference
    files = _xlsx_files(context.paths.input)
    if files:
        return files[0]
    raise RuntimeError(f"No .xlsx workbook found in {context.paths.input}, and no AddressReference workbook found in input/output.")


def run_reference_processing(context: JobContext) -> dict[str, object]:
    context.paths.input.mkdir(parents=True, exist_ok=True)
    context.paths.output.mkdir(parents=True, exist_ok=True)
    removed_temp = _cleanup_input_office_temp(context)
    parameters = context.operation.parameters
    action = str(parameters.get("reference_processing_action", "sort") or "sort").strip().lower()
    input_file = _reference_processing_input_file(context, parameters.get("reference_file"))
    timestamp = time.strftime("%Y%m%d_%H%M%S")
    output_file = _unique_path(context.paths.output / f"ReferenceProcessed_{input_file.stem}_{action}_{timestamp}.xlsx")
    report_file = context.report_dir / "reference_processing_summary.json"

    address_runtime = _build_address_runtime_context(context, parameters)
    normalizer = address_runtime.normalizer
    enabled_slots = _slot_selection_parameter(parameters)
    slot_order = _slot_order_parameter(parameters, enabled_slots)
    reference_columns = str(parameters.get("reference_columns") or "").strip()
    reference_column_header = str(parameters.get("reference_column_header") or "").strip()
    options = ReferenceProcessingOptions(
        action=action,
        reference_columns=reference_columns or reference_column_header or None,
        enabled_slots=enabled_slots,
        slot_order=slot_order,
        remove_source_column=_bool_parameter(parameters, "remove_source_column", True),
        slot_header_mode=str(parameters.get("slot_header_mode") or "full").strip().lower() or "full",
    )

    context.log("Reference processing screen")
    context.log(f"Action: {action}")
    context.log(f"Workbook: {input_file}")
    context.log(f"Output workbook: {output_file}")
    context.log(f"Reference column: {options.reference_columns or '(auto detect)'}")
    context.log(f"Enabled slots: {', '.join(options.enabled_slots)}")
    context.log(f"Sort slot order: {', '.join(options.slot_order)}")
    context.log(f"Remove source address column after deconstruction: {'yes' if options.remove_source_column else 'no'}")
    context.log(f"Slot header mode: {options.slot_header_mode}")
    _log_address_runtime_context(context, address_runtime)
    result = process_reference_workbook(
        input_file=input_file,
        output_file=output_file,
        normalizer=normalizer,
        options=options,
        report_file=report_file,
    )
    context.progress(1.0)
    context.log(
        "Reference processing complete: "
        f"rows={result.data_rows}, "
        f"sorted={result.sorted_rows}, "
        f"deconstructed={result.deconstructed_rows}, "
        f"slots={len(result.slot_columns)}."
    )
    if result.report_file:
        context.log(f"Report: {result.report_file}")
    return {
        "output": str(result.output_file),
        "report": str(result.report_file) if result.report_file else None,
        "input": str(result.input_file),
        "data_rows": result.data_rows,
        "sorted_rows": result.sorted_rows,
        "deconstructed_rows": result.deconstructed_rows,
        "slot_columns": list(result.slot_columns),
        "address_runtime": _address_runtime_summary(address_runtime),
        "slot_runtime": _slot_runtime_summary(enabled_slots, slot_order),
        "removed_temp_files": len(removed_temp),
        "artifacts": _artifact_summary(
            outputs=result.output_file,
            reports=result.report_file,
            inputs=result.input_file,
        ),
    }


def update_rosstat_oktmo_data(context: JobContext) -> dict[str, object]:
    parameters = context.operation.parameters
    raw_source_url = str(parameters.get("rosstat_oktmo_url") or "").strip()
    source_url = raw_source_url
    if source_url.rstrip("/") == DEFAULT_OKTMO_URL.rstrip("/"):
        source_url = ""
    data_dir = _optional_project_path(context, parameters.get("oktmo_data_dir")) or (context.paths.root / "data")
    context.log("Updating Rosstat OKTMO snapshot")
    context.log(f"Source URL: {source_url or 'default/env'}")
    context.log(f"Data root: {data_dir}")
    result = update_oktmo_snapshot(
        data_dir=data_dir,
        report_dir=context.paths.report,
        source_url=source_url or None,
        log=context.log,
        cancelled=context.cancelled,
    )
    try:
        clear_oktmo_lookup_caches(data_dir)
    except TypeError:
        clear_oktmo_lookup_caches()
    context.log("Building the OKTMO lookup cache for the updated snapshot")
    cache_status = warm_oktmo_cache(data_dir)
    context.log(
        "OKTMO cache ready: "
        f"regions={cache_status.get('regions', 0)}, "
        f"index_keys={cache_status.get('index_keys', 0)}, "
        f"contexts={cache_status.get('contexts', 0)}, "
        f"bytes={cache_status.get('cache_bytes', 0)}"
    )
    context.progress(1.0)
    context.log(f"Requested URL: {result.get('requested_url')}")
    context.log(f"Rosstat OKTMO data folder: {data_dir / 'rosstat'}")
    context.log(f"Latest data file: {result.get('latest_data_file')}")
    context.log(f"Update summary: {result.get('summary_path')}")
    result["caches_cleared"] = True
    result["cache"] = cache_status
    return result


def convert_legacy_office_sources(context: JobContext) -> dict[str, object]:
    context.paths.input.mkdir(parents=True, exist_ok=True)
    context.paths.output.mkdir(parents=True, exist_ok=True)
    context.paths.report.mkdir(parents=True, exist_ok=True)
    overwrite = _bool_parameter(context.operation.parameters, "overwrite_existing", False)
    delete_legacy = _bool_parameter(context.operation.parameters, "delete_legacy_after_conversion", True)
    context.log("Converting legacy DOC/XLS sources in input.")
    context.log(f"Overwrite existing DOCX/XLSX: {'yes' if overwrite else 'no'}")
    context.log(f"Delete legacy and same-name PDF/TXT/CSV/MD duplicates: {'yes' if delete_legacy else 'no'}")
    result = convert_legacy_office_folder(
        context.paths.input,
        overwrite=overwrite,
        delete_legacy_after_conversion=delete_legacy,
        log=context.log,
        progress=context.progress,
        cancelled=context.cancelled,
    )
    report = {
        "input_root": str(context.paths.input),
        "overwrite_existing": overwrite,
        "delete_legacy_after_conversion": delete_legacy,
        "converted": list(result.converted),
        "deleted": list(result.deleted),
        "skipped": list(result.skipped),
        "failed": list(result.failed),
    }
    report_path = context.paths.report / "legacy_office_conversion_report.json"
    report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    context.log(
        "Legacy conversion complete: "
        f"converted={len(result.converted)}, "
        f"deleted={len(result.deleted)}, "
        f"skipped={len(result.skipped)}, "
        f"failed={len(result.failed)}"
    )
    context.log(f"Report: {report_path}")
    return {
        "converted": len(result.converted),
        "deleted": len(result.deleted),
        "skipped": len(result.skipped),
        "failed": len(result.failed),
        "report": str(report_path),
    }


def run_address_aligner(context: JobContext) -> dict[str, object]:
    context.paths.input.mkdir(parents=True, exist_ok=True)
    context.paths.output.mkdir(parents=True, exist_ok=True)
    removed_temp = _cleanup_input_office_temp(context)
    files = _xlsx_files(context.paths.input)
    if not files:
        raise RuntimeError(f"No .xlsx files found in {context.paths.input}")

    outputs: list[str] = []
    file_reports: list[dict[str, object]] = []
    started = time.perf_counter()
    total = len(files)
    for index, file_path in enumerate(files, start=1):
        if context.cancelled():
            context.log("Operation cancelled by user.")
            return {"cancelled": True, "outputs": outputs}

        output_file = context.paths.output / f"Aligned_{file_path.name}"
        context.log(f"[{index}/{total}] Processing: {file_path.name}")
        runtime_config = _aligner_runtime_config(context, file_path)
        with _relay_output(context):
            process_file(file_path, output_file, aligner_config=runtime_config)
        outputs.append(str(output_file))
        file_reports.append(
            {
                "input": str(file_path),
                "output": str(output_file),
                "runtime_config": runtime_config or {},
            }
        )
        context.progress(index / total)

    elapsed = time.perf_counter() - started
    alignment_runtime = None
    if (
        _bool_parameter(context.operation.parameters, "normalize_before_match", False)
        or _bool_parameter(context.operation.parameters, "normalize_after_match", False)
        or _alignment_mode_requests_post_normalization(_normalized_alignment_mode(context.operation.parameters))
        or _bool_parameter(context.operation.parameters, "collect_before_match", False)
    ):
        alignment_runtime = _address_runtime_summary(_build_address_runtime_context(context, context.operation.parameters))
    report_path = context.report_dir / "address_alignment_summary.json"
    report_path.parent.mkdir(parents=True, exist_ok=True)
    report = {
        "mode": "address-alignment",
        "outputs": outputs,
        "elapsed_seconds": round(elapsed, 2),
        "removed_temp_files": len(removed_temp),
        "files": file_reports,
        "address_runtime": alignment_runtime,
    }
    report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    context.log(f"Address alignment complete: {len(outputs)} file(s), {elapsed:.1f}s")
    context.log(f"Report: {report_path}")
    return {
        "outputs": outputs,
        "report": str(report_path),
        "elapsed_seconds": round(elapsed, 2),
        "removed_temp_files": len(removed_temp),
        "artifacts": _artifact_summary(outputs=outputs, reports=report_path, inputs=files),
    }


def run_safe_table_join(context: JobContext) -> dict[str, object]:
    parameters = context.operation.parameters
    primary_file = resolve_project_path(context, str(parameters.get("primary_file") or ""))
    source_file = resolve_project_path(context, str(parameters.get("source_file") or ""))
    if not primary_file.exists() or not source_file.exists():
        raise RuntimeError("Primary or source workbook was not found.")
    raw_source_sheets = parameters.get("source_sheets") or []
    if isinstance(raw_source_sheets, (list, tuple, set)):
        source_sheets = tuple(str(part).strip() for part in raw_source_sheets if str(part).strip())
    else:
        source_sheets = tuple(part.strip() for part in str(raw_source_sheets).split(",") if part.strip())
    if not source_sheets:
        raise RuntimeError("Select at least one source sheet.")
    output_name = f"Joined_{primary_file.stem}.xlsx"
    output_file = context.paths.output / output_name
    options = JoinOptions(
        primary_file=primary_file,
        source_file=source_file,
        output_file=output_file,
        primary_sheet=str(parameters.get("primary_sheet") or "Таблица").strip(),
        source_sheets=source_sheets,
        primary_start_row=int(parameters.get("primary_start_row") or 5),
        source_start_row=int(parameters.get("source_start_row") or 7),
        primary_address_columns=parse_column_spec(parameters.get("primary_address_columns") or "B,C,D"),
        source_address_columns=parse_column_spec(parameters.get("source_address_columns") or "B,C"),
        source_payload_columns=parse_column_spec(parameters.get("source_payload_columns") or "D,E,F,G,H,I,J,K"),
        output_header_row=int(parameters.get("output_header_row") or 1),
    )
    context.log(f"Primary workbook (copied, never overwritten): {primary_file}")
    context.log(f"Source workbook (read-only): {source_file}")
    context.progress(0.1)
    result = safe_join_workbooks(options, context.log)
    context.progress(0.95)
    report_path = context.report_dir / "safe_table_join_summary.json"
    report_path.write_text(json.dumps({"mode": "safe-table-join", "options": {**parameters}, **result}, ensure_ascii=False, indent=2), encoding="utf-8")
    result["report"] = str(report_path)
    result["artifacts"] = _artifact_summary(outputs=output_file, reports=report_path, inputs=[primary_file, source_file])
    return result


def run_table_comparison(context: JobContext) -> dict[str, object]:
    parameters = context.operation.parameters
    raw_file_a = str(parameters.get("file_a") or "").strip()
    raw_file_b = str(parameters.get("file_b") or "").strip()
    if not raw_file_a or not raw_file_b:
        raise RuntimeError("Select both tables A and B.")
    file_a = resolve_project_path(context, raw_file_a)
    file_b = resolve_project_path(context, raw_file_b)
    if not file_a.exists() or not file_b.exists():
        raise RuntimeError("Table A or B was not found.")
    sheet_a = str(parameters.get("sheet_a") or "").strip()
    sheet_b = str(parameters.get("sheet_b") or "").strip()
    if not sheet_a or not sheet_b:
        raise RuntimeError("Select a worksheet for both A and B.")
    output_file = context.paths.output / f"Comparison_{file_a.stem}__{file_b.stem}.xlsx"
    options = ComparisonOptions(
        file_a=file_a,
        file_b=file_b,
        output_file=output_file,
        sheet_a=sheet_a,
        sheet_b=sheet_b,
        start_row_a=int(parameters.get("start_row_a") or 2),
        start_row_b=int(parameters.get("start_row_b") or 2),
        address_columns_a=parse_column_spec(parameters.get("address_columns_a")) if str(parameters.get("address_columns_a") or "").strip() else (),
        address_columns_b=parse_column_spec(parameters.get("address_columns_b")) if str(parameters.get("address_columns_b") or "").strip() else (),
        sort_addresses=_bool_parameter(parameters, "sort_addresses", False),
    )
    context.log(f"Table A: {file_a} [{sheet_a}], address columns={options.address_columns_a}")
    context.log(f"Table B: {file_b} [{sheet_b}], address columns={options.address_columns_b}")
    context.progress(0.1)
    result = build_comparison_workbook(options, context.log)
    context.progress(0.95)
    report_path = context.report_dir / "table_comparison_summary.json"
    report = {"mode": "table-comparison", "file_a": str(file_a), "file_b": str(file_b), "parameters": dict(parameters), **result}
    report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    result["report"] = str(report_path)
    result["artifacts"] = _artifact_summary(outputs=output_file, reports=report_path, inputs=[file_a, file_b])
    return result


def run_parser_diagnostic(context: JobContext) -> dict[str, object]:
    context.paths.input.mkdir(parents=True, exist_ok=True)
    context.paths.output.mkdir(parents=True, exist_ok=True)
    removed_temp = _cleanup_input_office_temp(context)
    files = _xlsx_files(context.paths.input)
    if not files:
        raise RuntimeError(f"No .xlsx files found in {context.paths.input}")

    outputs: list[str] = []
    total = len(files)
    for index, file_path in enumerate(files, start=1):
        if context.cancelled():
            context.log("Operation cancelled by user.")
            return {"cancelled": True, "outputs": outputs}

        output_file = context.paths.output / f"Diagnostic_{file_path.name}"
        context.log(f"[{index}/{total}] Diagnostic: {file_path.name}")
        with _relay_output(context):
            run_diagnostic(file_path, output_file)
        outputs.append(str(output_file))
        context.progress(index / total)

    context.log(f"Parser diagnostic complete: {len(outputs)} file(s)")
    return {"outputs": outputs, "removed_temp_files": len(removed_temp)}


def package_output(context: JobContext) -> dict[str, object]:
    context.paths.output.mkdir(parents=True, exist_ok=True)
    context.paths.release.mkdir(parents=True, exist_ok=True)

    files = [path for path in context.paths.output.rglob("*") if path.is_file()]
    if not files:
        raise RuntimeError(f"No output files found in {context.paths.output}")

    zip_path = context.paths.release / "address_aligner_output.zip"
    with ZipFile(zip_path, "w", ZIP_DEFLATED) as archive:
        total = len(files)
        for index, file_path in enumerate(files, start=1):
            if context.cancelled():
                context.log("Packaging cancelled by user.")
                return {"cancelled": True, "zip": str(zip_path)}
            archive.write(file_path, file_path.relative_to(context.paths.output).as_posix())
            context.progress(index / total)

    context.log(f"Created: {zip_path}")
    return {"zip": str(zip_path), "files": len(files)}


def remove_empty_rows_from_output(context: JobContext) -> dict[str, object]:
    context.paths.output.mkdir(parents=True, exist_ok=True)
    files = [
        path for path in _xlsx_files(context.paths.output)
        if not path.stem.endswith("_no_empty_rows")
    ]
    if not files:
        raise RuntimeError(f"No .xlsx files found in {context.paths.output}")

    outputs: list[str] = []
    total_removed = 0
    total = len(files)
    for index, file_path in enumerate(files, start=1):
        if context.cancelled():
            context.log("Empty-row cleanup cancelled by user.")
            return {"cancelled": True, "outputs": outputs, "removed_rows": total_removed}

        output_file = _unique_path(default_output_path(file_path))
        context.log(f"[{index}/{total}] Cleaning fully empty rows: {file_path.name}")
        removed_by_sheet = clean_workbook(file_path, output_file, all_sheets=False)
        removed = sum(removed_by_sheet.values())
        total_removed += removed
        outputs.append(str(output_file))
        details = ", ".join(f"{sheet}: {count}" for sheet, count in removed_by_sheet.items())
        context.log(f"  -> {output_file.name}; removed rows: {removed} ({details})")
        context.progress(index / total)

    context.log(f"Empty-row cleanup complete: {len(outputs)} file(s), removed rows: {total_removed}")
    return {"outputs": outputs, "removed_rows": total_removed}


def cleanup_input_output(context: JobContext) -> dict[str, object]:
    context.log("Cleaning managed input/output folders.")
    input_result = _clean_managed_folder(context, context.paths.input, "input")
    context.progress(0.5)
    if context.cancelled():
        context.log("Input/output cleanup cancelled by user after input cleanup.")
        return {"cancelled": True, "input": input_result}

    output_result = _clean_managed_folder(context, context.paths.output, "output")
    context.progress(1.0)
    total_removed = int(input_result["removed_items"]) + int(output_result["removed_items"])
    total_skipped = len(input_result["skipped_items"]) + len(output_result["skipped_items"])
    context.log(f"Input/output cleanup complete. Removed: {total_removed}, skipped: {total_skipped}")
    return {"input": input_result, "output": output_result, "removed_items": total_removed, "skipped_items": total_skipped}
