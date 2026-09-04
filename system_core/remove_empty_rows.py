from __future__ import annotations

import argparse
from pathlib import Path

from openpyxl import load_workbook


def has_meaningful_value(value) -> bool:
    if value is None:
        return False
    if isinstance(value, str):
        return value.strip() != ""
    return True


def row_is_empty(worksheet, row_number: int) -> bool:
    for row in worksheet.iter_rows(
        min_row=row_number,
        max_row=row_number,
        min_col=1,
        max_col=worksheet.max_column,
        values_only=True,
    ):
        return not any(has_meaningful_value(value) for value in row)
    return True


def compact_worksheet(worksheet) -> int:
    max_row = worksheet.max_row
    max_col = worksheet.max_column
    write_row = 1

    for read_row in range(1, max_row + 1):
        values = [
            worksheet.cell(row=read_row, column=column).value
            for column in range(1, max_col + 1)
        ]
        if not any(has_meaningful_value(value) for value in values):
            continue

        if write_row != read_row:
            for column, value in enumerate(values, start=1):
                worksheet.cell(row=write_row, column=column).value = value
        write_row += 1

    removed = max_row - write_row + 1
    if removed > 0:
        worksheet.delete_rows(write_row, removed)
    return removed


def clean_workbook(input_path: Path, output_path: Path, all_sheets: bool) -> dict[str, int]:
    workbook = load_workbook(input_path)
    worksheets = workbook.worksheets if all_sheets else [workbook.active]
    removed_by_sheet: dict[str, int] = {}

    for worksheet in worksheets:
        removed_by_sheet[worksheet.title] = compact_worksheet(worksheet)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    workbook.close()
    return removed_by_sheet


def default_output_path(input_path: Path) -> Path:
    return input_path.with_name(f"{input_path.stem}_no_empty_rows{input_path.suffix}")


def iter_input_files(path: Path) -> list[Path]:
    if path.is_file():
        return [path]
    return sorted(
        file
        for file in path.glob("*.xlsx")
        if file.is_file() and not file.name.startswith("~$")
    )


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Remove rows that have no values in any column of an Excel workbook."
    )
    parser.add_argument("input", type=Path, help="Input .xlsx file or directory with .xlsx files.")
    parser.add_argument(
        "output",
        nargs="?",
        type=Path,
        help="Output file or directory. If omitted, creates *_no_empty_rows.xlsx next to input.",
    )
    parser.add_argument(
        "--all-sheets",
        action="store_true",
        help="Clean all worksheets. By default only the active worksheet is cleaned.",
    )
    parser.add_argument(
        "--overwrite",
        action="store_true",
        help="Allow overwriting the output file if it already exists.",
    )
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    input_path = args.input.resolve()
    if not input_path.exists():
        print(f"[ERROR] Not found: {input_path}")
        return 1

    input_files = iter_input_files(input_path)
    if not input_files:
        print(f"[WARN] No .xlsx files found: {input_path}")
        return 0

    output_arg = args.output.resolve() if args.output else None

    for input_file in input_files:
        if input_file.suffix.lower() != ".xlsx":
            print(f"[SKIP] Not an .xlsx file: {input_file}")
            continue

        if output_arg is None:
            output_file = default_output_path(input_file)
        elif input_path.is_dir() or output_arg.is_dir():
            output_file = output_arg / input_file.name
        else:
            output_file = output_arg

        if output_file.exists() and not args.overwrite:
            print(f"[SKIP] Output exists, use --overwrite: {output_file}")
            continue

        removed_by_sheet = clean_workbook(input_file, output_file, args.all_sheets)
        from system_core.excel_layout import format_workbook

        format_workbook(output_file)
        total_removed = sum(removed_by_sheet.values())
        details = ", ".join(f"{sheet}: {count}" for sheet, count in removed_by_sheet.items())
        print(f"[OK] {input_file.name} -> {output_file.name}; removed rows: {total_removed} ({details})")

    return 0


if __name__ == "__main__":
    raise SystemExit(main())
