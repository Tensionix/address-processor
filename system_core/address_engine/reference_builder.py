from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Any, Iterable
import re
import json

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
import pandas as pd

from .address_diagnostics import diagnostic_summary
from .normalizer_runtime import normalizer_oktmo_key_profile, normalizer_runtime_report
from .slot_fill import build_slot_fill_summary
from .source_intake import SourceCandidate, discover_source_files, extract_reference_candidates, extract_source_candidates
from .universal_address import AddressNormalizer, NormalizedAddress


REFERENCE_PREFIX = "AddressReference"
ADDRESS_CONSTRUCTION_SLOTS = (
    "postal_index",
    "federal_district",
    "parent_subject",
    "subject",
    "autonomous_okrug",
    "municipality",
    "locality",
    "territory",
    "microdistrict",
    "street",
    "house",
    "premise",
)
TECHNICAL_REFERENCE_SLOTS = (
    "oktmo_code",
    "oktmo_name",
    "house_base",
    "house_mods",
    "house_modifiers",
    "street_numbers",
    "territory_key",
    "match_key",
)
DEFAULT_SLOT_ORDER = ADDRESS_CONSTRUCTION_SLOTS
DEFAULT_ENABLED_SLOTS = ADDRESS_CONSTRUCTION_SLOTS + TECHNICAL_REFERENCE_SLOTS
DEFAULT_SELECTED_SLOTS = tuple(slot for slot in DEFAULT_ENABLED_SLOTS if slot != "microdistrict")
SLOT_COLUMN_BY_ID = {
    "postal_index": "Postal_Index",
    "federal_district": "Federal_District",
    "parent_subject": "Parent_Subject",
    "subject": "Subject",
    "autonomous_okrug": "Autonomous_Okrug",
    "oktmo_code": "OKTMO_Code",
    "oktmo_name": "OKTMO_Name",
    "municipality": "Municipality",
    "locality": "Locality",
    "territory": "Territory",
    "microdistrict": "Microdistrict",
    "street": "Street",
    "house": "House",
    "house_base": "House_Base",
    "house_mods": "House_Mods",
    "house_modifiers": "House_Modifiers",
    "premise": "Premise",
    "street_numbers": "Street_Numbers",
    "territory_key": "Territory_Key",
    "match_key": "Match_Key",
}
OLD_SLOT_COLUMNS = {
    "Raw_Address",
    "Normalized_Address",
    "Postal_Index",
    "Federal_District",
    "Parent_Subject",
    "Subject",
    "Autonomous_Okrug",
    "OKTMO_Code",
    "OKTMO_Name",
    "Municipality",
    "Locality",
    "Territory",
    "Microdistrict",
    "Street",
    "House",
    "House_Base",
    "House_Mods",
    "House_Modifiers",
    "Premise",
    "Street_Numbers",
    "Territory_Key",
    "Match_Key",
    "Raw_Examples",
    "Source_Count",
    "Sources",
    "First_Source",
}


@dataclass(frozen=True)
class ReferenceOptions:
    action: str = "generate"
    source_columns: str | None = None
    reference_columns: str | None = None
    reference_file: Path | None = None
    enabled_slots: tuple[str, ...] = DEFAULT_SELECTED_SLOTS
    slot_order: tuple[str, ...] = DEFAULT_SLOT_ORDER
    omit_microdistrict_when_street_found: bool = True
    include_slots: bool = True
    clean_output_workbook: bool = False
    clean_old_values: bool = True
    clean_old_slots: bool = True


@dataclass(frozen=True)
class ReferenceBuildResult:
    output_file: Path
    staging_file: Path | None
    report_file: Path | None
    source_files: tuple[Path, ...]
    issues: tuple[str, ...]
    raw_candidates: int
    normalized_rows: int
    duplicate_rows: int


def build_reference_workbook(
    *,
    input_dir: Path,
    output_file: Path,
    normalizer: AddressNormalizer,
    options: ReferenceOptions,
    staging_file: Path | None = None,
    report_file: Path | None = None,
) -> ReferenceBuildResult:
    input_files = discover_source_files(input_dir)
    reference_file = options.reference_file
    if reference_file and reference_file.exists():
        reference_file = reference_file.resolve()

    rows: list[dict[str, Any]] = []
    sources: list[Path] = []
    issues: list[str] = []
    raw_candidates = 0

    if options.action in {"update", "clean"}:
        existing = reference_file or _discover_existing_reference(input_dir, output_file.parent)
        if existing:
            reference_file = existing.resolve()
            records, count, source_issues = _rows_from_source_file(
                existing,
                normalizer=normalizer,
                source_columns=options.source_columns,
                reference_columns=options.reference_columns,
                enabled_slots=options.enabled_slots,
                slot_order=options.slot_order,
                omit_microdistrict_when_street_found=options.omit_microdistrict_when_street_found,
                source_kind="existing_reference",
            )
            rows.extend(records)
            raw_candidates += count
            issues.extend(source_issues)
            sources.append(existing)
        elif options.action == "clean":
            raise RuntimeError("No existing reference workbook was found for cleanup.")

    if options.action in {"generate", "update"}:
        for file_path in input_files:
            if reference_file and file_path.resolve() == reference_file:
                continue
            records, count, source_issues = _rows_from_source_file(
                file_path,
                normalizer=normalizer,
                source_columns=options.source_columns,
                reference_columns=None,
                enabled_slots=options.enabled_slots,
                slot_order=options.slot_order,
                omit_microdistrict_when_street_found=options.omit_microdistrict_when_street_found,
                source_kind="source",
            )
            rows.extend(records)
            raw_candidates += count
            issues.extend(source_issues)
            sources.append(file_path)

    if not rows:
        raise RuntimeError("No address-like values were found for reference generation.")

    normalized_rows, duplicate_rows, staged_rows = _merge_reference_rows(
        rows,
        include_slots=options.include_slots,
        enabled_slots=options.enabled_slots,
        slot_order=options.slot_order,
        clean_old_values=options.clean_old_values,
        clean_old_slots=options.clean_old_slots,
    )
    output_file.parent.mkdir(parents=True, exist_ok=True)
    pd.DataFrame(normalized_rows).to_excel(output_file, index=False)
    if options.clean_output_workbook:
        _style_clean_reference_workbook(output_file)
    if staging_file:
        _write_reference_staging_workbook(staging_file, staged_rows, normalized_rows)
    if report_file:
        oktmo_key_profile = normalizer_oktmo_key_profile(normalizer)
        report_file.parent.mkdir(parents=True, exist_ok=True)
        summary = {
            "mode": "address-reference-normalization",
            "action": options.action,
            "input_dir": str(input_dir),
            "staging_file": str(staging_file) if staging_file else None,
            "output_file": str(output_file),
            "source_file_count": len(sources),
            "raw_candidates": raw_candidates,
            "normalized_rows": len(normalized_rows),
            "duplicate_rows": duplicate_rows,
            "address_diagnostics": diagnostic_summary(normalized_rows),
            "slot_fill": build_slot_fill_summary(
                normalized_rows,
                enabled_slots=options.enabled_slots,
                slot_order=options.slot_order,
                slot_columns=SLOT_COLUMN_BY_ID,
                default_slot_order=DEFAULT_SLOT_ORDER,
                get_slot_value=lambda row, slot: row.get(SLOT_COLUMN_BY_ID[slot]),
                get_normalized=lambda row: row.get("Normalized_Address"),
            ),
            "normalizer": normalizer_runtime_report(normalizer, oktmo_key_profile),
            "options": {
                "source_columns": options.source_columns,
                "reference_columns": options.reference_columns,
                "reference_file": str(reference_file) if reference_file else None,
                "enabled_slots": list(options.enabled_slots),
                "slot_order": list(options.slot_order),
                "omit_microdistrict_when_street_found": bool(options.omit_microdistrict_when_street_found),
                "include_slots": bool(options.include_slots),
                "clean_output_workbook": bool(options.clean_output_workbook),
                "clean_old_values": bool(options.clean_old_values),
                "clean_old_slots": bool(options.clean_old_slots),
            },
            "issues": issues[:200],
            "sources": [str(path) for path in sources],
        }
        report_file.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    return ReferenceBuildResult(
        output_file=output_file,
        staging_file=staging_file,
        report_file=report_file,
        source_files=tuple(sources),
        issues=tuple(issues),
        raw_candidates=raw_candidates,
        normalized_rows=len(normalized_rows),
        duplicate_rows=duplicate_rows,
    )


def _rows_from_source_file(
    file_path: Path,
    *,
    normalizer: AddressNormalizer,
    source_columns: str | None,
    reference_columns: str | None,
    enabled_slots: tuple[str, ...],
    slot_order: tuple[str, ...],
    omit_microdistrict_when_street_found: bool,
    source_kind: str,
) -> tuple[list[dict[str, Any]], int, list[str]]:
    if source_kind == "existing_reference":
        candidates, issues = extract_reference_candidates(file_path, reference_columns=reference_columns)
    else:
        candidates, issues = extract_source_candidates(file_path, source_columns=source_columns)
    rows: list[dict[str, Any]] = []
    raw_candidates = 0
    for candidate in candidates:
        record = normalizer.normalize(candidate.value)
        if not _looks_like_reference_address(record):
            continue
        raw_candidates += 1
        rows.append(
            _reference_row(
                record,
                enabled_slots=enabled_slots,
                slot_order=slot_order,
                omit_microdistrict_when_street_found=omit_microdistrict_when_street_found,
                candidate=candidate,
                source_kind=source_kind,
            )
        )
    return rows, raw_candidates, issues


def _reference_row(
    record: NormalizedAddress,
    *,
    enabled_slots: tuple[str, ...],
    slot_order: tuple[str, ...],
    omit_microdistrict_when_street_found: bool,
    candidate: SourceCandidate,
    source_kind: str,
) -> dict[str, Any]:
    slots = record.slot_dict()
    row = {
        "Raw_Address": record.raw,
        "Normalized_Address": _construct_address(
            record,
            enabled_slots,
            slot_order,
            omit_microdistrict_when_street_found=omit_microdistrict_when_street_found,
        )
        or record.normalized,
        "Postal_Index": record.postal_index,
        "Federal_District": slots.get("federal_district"),
        "Parent_Subject": slots.get("parent_subject"),
        "Subject": slots.get("subject"),
        "Autonomous_Okrug": slots.get("autonomous_okrug"),
        "OKTMO_Code": record.oktmo_code,
        "OKTMO_Name": record.oktmo_name,
        "Municipality": slots.get("municipality"),
        "Locality": slots.get("locality"),
        "Territory": slots.get("territory"),
        "Microdistrict": slots.get("microdistrict"),
        "Street": slots.get("street"),
        "House": slots.get("house"),
        "House_Base": slots.get("house_base"),
        "House_Mods": slots.get("house_mods"),
        "House_Modifiers": ", ".join(slots.get("house_modifiers") or []),
        "Premise": slots.get("premise"),
        "Street_Numbers": ", ".join(slots.get("street_numbers") or []),
        "Territory_Key": record.territory_key,
        "Match_Key": record.match_key,
        "Raw_Examples": record.raw,
        "Source_Count": 1,
        "Sources": f"{source_kind}:{candidate.source_kind}:{candidate.source_file.name}:{candidate.locator}",
        "First_Source": str(candidate.source_file),
    }
    row["Match_Key"] = (
        _construct_match_key(
            row,
            enabled_slots,
            omit_microdistrict_when_street_found=omit_microdistrict_when_street_found,
        )
        or record.match_key
    )
    return row


def _merge_reference_rows(
    rows: list[dict[str, Any]],
    *,
    include_slots: bool,
    enabled_slots: tuple[str, ...],
    slot_order: tuple[str, ...],
    clean_old_values: bool,
    clean_old_slots: bool,
) -> tuple[list[dict[str, Any]], int, list[dict[str, Any]]]:
    merged: dict[str, dict[str, Any]] = {}
    group_counts: dict[str, int] = {}
    group_examples: dict[str, list[str]] = {}
    row_keys: dict[int, str] = {}
    duplicates = 0
    for row in rows:
        key = str(row.get("Match_Key") or row.get("Normalized_Address") or row.get("Raw_Address") or "").strip()
        if not key:
            continue
        if not clean_old_values:
            key = f"{key}|raw:{len(merged):06d}"
        row_keys[id(row)] = key
        group_counts[key] = group_counts.get(key, 0) + 1
        examples = group_examples.setdefault(key, [])
        raw_example = _clean(row.get("Raw_Address")) or _clean(row.get("Normalized_Address"))
        if raw_example and raw_example not in examples and len(examples) < 8:
            examples.append(raw_example)
        if key not in merged:
            merged[key] = dict(row)
            continue
        duplicates += 1
        current = merged[key]
        current["Source_Count"] = int(current.get("Source_Count") or 1) + int(row.get("Source_Count") or 1)
        current["Raw_Examples"] = _append_unique_examples(current.get("Raw_Examples"), row.get("Raw_Address"))
        current["Sources"] = _append_unique_examples(current.get("Sources"), row.get("Sources"), limit=20, separator="; ")
        for field in (
            "Postal_Index",
            "Federal_District",
            "Parent_Subject",
            "Subject",
            "Autonomous_Okrug",
            "OKTMO_Code",
            "OKTMO_Name",
            "Municipality",
            "Locality",
            "Territory",
            "Microdistrict",
            "Street",
            "House",
        ):
            if not current.get(field) and row.get(field):
                current[field] = row[field]

    result = [
        _project_reference_columns(
            row,
            include_slots=include_slots,
            enabled_slots=enabled_slots,
            slot_order=slot_order,
            clean_old_slots=clean_old_slots,
        )
        for row in merged.values()
    ]
    result.sort(key=lambda item: str(item.get("Normalized_Address") or item.get("Raw_Address") or "").lower())
    staged_rows = _reference_staging_rows(rows, row_keys, group_counts, group_examples, merged)
    return result, duplicates, staged_rows


def _reference_staging_rows(
    rows: list[dict[str, Any]],
    row_keys: dict[int, str],
    group_counts: dict[str, int],
    group_examples: dict[str, list[str]],
    merged: dict[str, dict[str, Any]],
) -> list[dict[str, Any]]:
    first_seen: set[str] = set()
    staged: list[dict[str, Any]] = []
    for row in rows:
        key = row_keys.get(id(row))
        if not key:
            continue
        status = "ready" if key not in first_seen else "duplicate"
        first_seen.add(key)
        merged_row = merged.get(key, {})
        staged_row = {
            "Reference_Status": status,
            "Merge_Key": key,
            "Group_Size": group_counts.get(key, 1),
            "Final_Address": merged_row.get("Normalized_Address") or row.get("Normalized_Address"),
            "Group_Raw_Examples": " | ".join(group_examples.get(key, [])),
        }
        for column in (
            "Normalized_Address",
            "Raw_Address",
            "Postal_Index",
            "Federal_District",
            "Parent_Subject",
            "Subject",
            "Autonomous_Okrug",
            "Municipality",
            "Locality",
            "Territory",
            "Microdistrict",
            "Street",
            "House",
            "Premise",
            "OKTMO_Code",
            "OKTMO_Name",
            "Match_Key",
            "Sources",
            "First_Source",
        ):
            if column in row:
                staged_row[column] = row.get(column)
        staged.append(staged_row)
    return staged


def _write_reference_staging_workbook(
    staging_file: Path,
    staged_rows: list[dict[str, Any]],
    normalized_rows: list[dict[str, Any]],
) -> None:
    staging_file.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(staging_file, engine="openpyxl") as writer:
        pd.DataFrame(staged_rows).to_excel(writer, index=False, sheet_name="Reference_Candidates")
        pd.DataFrame(normalized_rows).to_excel(writer, index=False, sheet_name="Reference_Final")


def _style_clean_reference_workbook(output_file: Path) -> None:
    workbook = load_workbook(output_file)
    try:
        for sheet in workbook.worksheets:
            _style_clean_reference_sheet(sheet)
        workbook.save(output_file)
    finally:
        workbook.close()


def _style_clean_reference_sheet(sheet: Any) -> None:
    header_fill = PatternFill(fill_type="solid", fgColor="1F2933")
    header_font = Font(name="Segoe UI", size=9, bold=True, color="FFFFFF")
    body_font = Font(name="Segoe UI", size=9)
    alignment = Alignment(vertical="top", wrap_text=True)
    max_row = sheet.max_row or 1
    max_column = sheet.max_column or 1
    for column_index in range(1, max_column + 1):
        header_cell = sheet.cell(row=1, column=column_index)
        header_cell.font = header_font
        header_cell.fill = header_fill
        header_cell.alignment = alignment
        max_len = len(str(header_cell.value or ""))
        for row_index in range(2, max_row + 1):
            cell = sheet.cell(row=row_index, column=column_index)
            cell.font = body_font
            cell.alignment = alignment
            max_len = max(max_len, len(str(cell.value or "")))
        sheet.column_dimensions[excel_column_letter(column_index - 1)].width = max(12, min(64, max_len + 2))
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions


def _project_reference_columns(
    row: dict[str, Any],
    *,
    include_slots: bool,
    enabled_slots: tuple[str, ...],
    slot_order: tuple[str, ...],
    clean_old_slots: bool,
) -> dict[str, Any]:
    base_order = [
        "Normalized_Address",
        "Raw_Address",
    ]
    slot_columns = _ordered_slot_columns(enabled_slots, slot_order)
    if "Match_Key" not in slot_columns:
        slot_columns.append("Match_Key")
    tail_order = ["Raw_Examples", "Source_Count", "Sources", "First_Source"]
    ordered = base_order + (slot_columns if include_slots else ["Match_Key"]) + tail_order
    projected = {key: row.get(key) for key in ordered if key in row}
    if not clean_old_slots:
        for key, value in row.items():
            if key not in projected:
                projected[key] = value
    return projected


def _ordered_slot_columns(enabled_slots: tuple[str, ...], slot_order: tuple[str, ...]) -> list[str]:
    ordered_slots = list(_ordered_construction_slots(enabled_slots, slot_order))
    ordered_slots.extend(slot for slot in enabled_slots if slot not in ordered_slots)
    return [SLOT_COLUMN_BY_ID[slot] for slot in ordered_slots if slot in SLOT_COLUMN_BY_ID]


def _ordered_construction_slots(enabled_slots: tuple[str, ...], slot_order: tuple[str, ...]) -> tuple[str, ...]:
    enabled = set(enabled_slots)
    result: list[str] = []
    for slot in slot_order:
        if slot in ADDRESS_CONSTRUCTION_SLOTS and slot in enabled and slot not in result:
            result.append(slot)
    for slot in ADDRESS_CONSTRUCTION_SLOTS:
        if slot in enabled and slot not in result:
            result.append(slot)
    return tuple(result)


def _construct_address(
    record: NormalizedAddress,
    enabled_slots: tuple[str, ...],
    slot_order: tuple[str, ...],
    *,
    omit_microdistrict_when_street_found: bool = False,
) -> str | None:
    parts = record.parts
    values: list[str] = []
    slot_values = {
        "postal_index": record.postal_index,
        "federal_district": record.federal_district or parts.federal_district,
        "parent_subject": record.parent_subject or parts.parent_subject,
        "subject": record.subject or parts.subject,
        "autonomous_okrug": record.autonomous_okrug or parts.autonomous_okrug,
        "municipality": record.municipality or parts.municipality,
        "locality": parts.locality,
        "territory": parts.territory,
        "microdistrict": parts.microdistrict,
        "street": parts.street,
        "house": parts.house,
        "premise": parts.premise,
    }
    for slot in _ordered_construction_slots(enabled_slots, slot_order):
        if slot == "microdistrict" and omit_microdistrict_when_street_found and parts.street:
            continue
        _append_unique(values, slot_values.get(slot))
    return ", ".join(values) or None


def _construct_match_key(
    row: dict[str, Any],
    enabled_slots: tuple[str, ...],
    *,
    omit_microdistrict_when_street_found: bool = False,
) -> str:
    key_parts: list[str] = []
    for slot in enabled_slots:
        column = SLOT_COLUMN_BY_ID.get(slot)
        if not column:
            continue
        if slot in {"postal_index", "oktmo_name", "house_modifiers", "street_numbers", "territory_key", "match_key"}:
            continue
        if slot == "microdistrict" and omit_microdistrict_when_street_found and row.get("Street"):
            continue
        value = row.get(column)
        if value:
            key_parts.append(_norm(value))
    if not key_parts:
        return str(row.get("Match_Key") or "")
    return "|".join(key_parts)


def _append_unique(values: list[str], value: Any) -> None:
    text = _clean(value)
    if not text:
        return
    text_key = _norm(text)
    text_level_key = _level_norm(text)
    if any(_norm(existing) == text_key or _level_norm(existing) == text_level_key for existing in values):
        return
    values.append(text)


def _looks_like_reference_address(record: NormalizedAddress) -> bool:
    if not record.has_address_core:
        return False
    raw = str(record.raw or "")
    if len(raw) > 600:
        return False
    if not (record.house_base or record.parts.street or record.parts.locality or record.oktmo_code):
        return False
    return True


def _selected_column_indices(value: str | None, df: pd.DataFrame) -> set[int] | None:
    tokens = [part.strip() for part in str(value or "").split(",") if part.strip()]
    if not tokens:
        return None
    result: set[int] = set()
    for token in tokens:
        index = excel_column_index(token) if re.fullmatch(r"[A-Za-z]+", token) else int(token) - 1
        if index < 0 or index >= len(df.columns):
            raise ValueError(f"Column {token!r} is outside workbook width.")
        result.add(index)
    return result


def _discover_existing_reference(input_dir: Path, output_dir: Path) -> Path | None:
    candidates = [
        *input_dir.glob(f"{REFERENCE_PREFIX}*.xlsx"),
        *output_dir.glob(f"{REFERENCE_PREFIX}*.xlsx"),
    ]
    candidates = [path for path in candidates if path.is_file() and not path.name.startswith("~$")]
    if not candidates:
        return None
    return max(candidates, key=lambda path: path.stat().st_mtime)


def _append_unique_examples(
    current: Any,
    value: Any,
    *,
    limit: int = 5,
    separator: str = " | ",
) -> str:
    items: list[str] = []
    for part in str(current or "").split(separator):
        cleaned = part.strip()
        if cleaned and cleaned not in items:
            items.append(cleaned)
    cleaned_value = str(value or "").strip()
    if cleaned_value and cleaned_value not in items:
        items.append(cleaned_value)
    return separator.join(items[:limit])


def excel_column_index(value: str) -> int:
    text = str(value).strip().upper()
    if not re.fullmatch(r"[A-Z]+", text):
        raise ValueError(f"Column must be an Excel letter: {value!r}")
    index = 0
    for char in text:
        index = index * 26 + (ord(char) - 64)
    return index - 1


def excel_column_letter(index: int) -> str:
    result = ""
    index = int(index)
    while True:
        result = chr(65 + index % 26) + result
        index = index // 26 - 1
        if index < 0:
            break
    return result


def _clean(value: Any) -> str | None:
    if value is None:
        return None
    text = re.sub(r"\s+", " ", str(value).replace("\xa0", " ")).strip(" ,")
    if not text or text.lower() == "nan":
        return None
    return text


def _norm(value: Any) -> str:
    text = str(value or "").lower().replace("ё", "е")
    text = re.sub(r"[^0-9a-zа-я]+", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def _level_norm(value: Any) -> str:
    text = _norm(value)
    text = re.sub(
        r"\b(?:город|г|село|с|деревня|д|поселок|пос|п|р\s*п|пгт|хутор|х|станица|ст\s*ца)\b",
        " ",
        text,
    )
    return re.sub(r"\s+", " ", text).strip()
