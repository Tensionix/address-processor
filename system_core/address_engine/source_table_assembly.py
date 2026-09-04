from __future__ import annotations

from collections import Counter, OrderedDict
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Callable
import csv
import json
import re
import xml.etree.ElementTree as ET
import zipfile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
import pandas as pd

from .address_diagnostics import diagnostic_summary
from .audion_address_core import explicit_settlement_hint, looks_like_address_text, parse_address_components
from .normalizer_runtime import normalizer_oktmo_key_profile, normalizer_runtime_report
from .reference_builder import (
    DEFAULT_ENABLED_SLOTS,
    DEFAULT_SELECTED_SLOTS,
    DEFAULT_SLOT_ORDER,
    SLOT_COLUMN_BY_ID,
    _append_unique_examples,
    _construct_match_key,
    _ordered_slot_columns,
    _project_reference_columns,
    _reference_row,
)
from .slot_fill import build_slot_fill_summary
from .source_intake import (
    SourceCandidate,
    discover_source_files,
    excel_column_index,
    excel_column_letter,
    extract_reference_candidates,
    extract_source_candidates,
)
from .universal_address import AddressNormalizer, NormalizedAddress


ADDRESS_COLLECTION_PREFIX = "AddressCollection"
ADDRESS_HINT_RE = re.compile(
    r"\b(?:адрес|местополож|ул\.?|улица|пер\.?|переулок|пр-?кт|проспект|шоссе|проезд|"
    r"б-?р|бульвар|тракт|д\.?|дом|стр\.?|строен|корп\.?|корпус|с\.?|село|д\.?|деревня|"
    r"п\.?|пос\.?|поселок|посёлок|пгт|рп|город|г\.?|тер\.?|мкр|микрорайон|квартал|"
    r"снт|днт|тсн)\b",
    re.IGNORECASE,
)

DIAGNOSTIC_COLUMNS = (
    "Drop_Reason",
    "Disposition",
    "Review_Action",
    "Candidate_Count",
    "Best_Slot_Score",
    "OKTMO_Scope_Match",
    "OKTMO_Scope_Code",
    "OKTMO_Scope_Source",
    "OKTMO_Scope_Weight",
    "OKTMO_Scope_Valid",
    "OKTMO_Scope_Context",
)
STAGING_CANDIDATE_COLUMNS = (
    "Postal_Index",
    "Federal_District",
    "Parent_Subject",
    "Subject",
    "Autonomous_Okrug",
    "Municipality",
    "Locality",
    "Territory",
    "Microdistrict",
    "Street",
    "House",
    "Premise",
    "OKTMO_Code",
    "OKTMO_Name",
)
ADMIN_CONTEXT_RE = re.compile(
    r"\b(?:федеральн(?:ый|ого)\s+округ|область|край|республика|автономн\w+\s+округ|"
    r"муниципальн\w+|городской\s+округ|район|р-н|г\.?|город|с\.?|село|д\.?|деревня|"
    r"п\.?|пос\.?|пос[её]лок|пгт|рп)\b",
    re.IGNORECASE,
)
STREET_MARKER_RE = re.compile(
    r"\b(?:ул\.?|улица|пер\.?|переулок|пр-?кт|проспект|шоссе|проезд|пр-д\.?|"
    r"б-?р|бульвар|наб\.?|набережная|пл\.?|площадь|тупик|аллея|линия)\b",
    re.IGNORECASE,
)
HOUSE_MARKER_RE = re.compile(r"\b(?:д\.?|дом|зд\.?|здание|стр\.?|строение)\s*\d", re.IGNORECASE)
WORD_NS = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
ODT_NS = {
    "office": "urn:oasis:names:tc:opendocument:xmlns:office:1.0",
    "table": "urn:oasis:names:tc:opendocument:xmlns:table:1.0",
    "text": "urn:oasis:names:tc:opendocument:xmlns:text:1.0",
}
ODT_REPEAT_ATTR = f"{{{ODT_NS['table']}}}number-columns-repeated"


def _normalize_spaced_address_marker_dots(value: Any) -> str:
    text = str(value or "").replace("\xa0", " ")
    return re.sub(
        r"\b(г|с|д|п|х|у|ул|пер|пр|пр-т|пр-кт|просп|пр-д|б-р|наб|пл)\s+\.",
        r"\1.",
        text,
        flags=re.IGNORECASE,
    )

ADDRESS_HEADER_MARKERS = (
    "местоположение",
    "местонахождение",
    "место нахождения",
    "адрес",
    "адрес полностью",
    "почтовый адрес",
    "фактический адрес",
)
STREET_HEADER_MARKERS = ("улица", "наименование улицы", "street")
HOUSE_HEADER_MARKERS = ("дом", "номер дома", "house", "house number")
LOCALITY_HEADER_MARKERS = ("город", "населенный пункт", "населённый пункт", "locality")
MUNICIPALITY_HEADER_MARKERS = ("муницип", "городской округ", "муниципальное образование")
SUBJECT_HEADER_MARKERS = ("субъект", "регион", "область", "край")
POSTAL_HEADER_MARKERS = ("почтовый индекс", "индекс", "postal")
MICRODISTRICT_HEADER_MARKERS = ("микрорайон", "мкр", "квартал")
TERRITORY_HEADER_MARKERS = ("территория", "тер ")
OKTMO_HEADER_MARKERS = ("октмо",)
ORG_CONTEXT_MARKERS = (
    "маоу",
    "мадоу",
    "мбоу",
    "мбдоу",
    "моу",
    "гбоу",
    "фгбоу",
    "чоу",
    "ано",
    "сош",
    "оош",
    "нош",
    "гимназия",
    "лицей",
    "школа",
    "детский сад",
    "дошкольное образовательное учреждение",
    "образовательное учреждение",
    "колледж",
    "техникум",
    "университет",
    "академия",
)
ROUTE_CONTEXT_MARKERS = (
    "подвоз",
    "маршрут",
    "маршруты",
    "количество детей км",
    "информация о подвозе",
)
BARE_PR_PROEZD_STREET_NAMES = {
    "геологоразведчиков",
    "заречный",
}
STREET_PREFIX_RE = re.compile(
    r"^\s*(?:ул\.?|улица|пер\.?|переулок|пр-?кт\.?|проспект|просп\.?|шоссе|ш\.?|"
    r"проезд|пр-д\.?|б-?р\.?|бульвар|наб\.?|набережная|пл\.?|площадь|"
    r"тупик|аллея|линия)\s+",
    re.IGNORECASE,
)
STREET_SUFFIX_RE = re.compile(
    r"\b(?:тракт|тракта|шоссе|проезд|бульвар|переулок|проспект|набережная|площадь)\s*$",
    re.IGNORECASE,
)


@dataclass(frozen=True)
class AddressAssemblyOptions:
    action: str = "generate"
    source_root: Path | None = None
    source_columns: str | None = None
    reference_columns: str | None = None
    reference_file: Path | None = None
    target_dir: Path | None = None
    target_file: Path | None = None
    target_column: str | None = "A"
    staging_file: Path | None = None
    enabled_slots: tuple[str, ...] = DEFAULT_SELECTED_SLOTS
    slot_order: tuple[str, ...] = DEFAULT_SLOT_ORDER
    omit_microdistrict_when_street_found: bool = True
    include_slots: bool = True
    clean_output_workbook: bool = False
    clean_duplicates: bool = True
    sanitize_incomplete: bool = True
    write_rejected: bool = True


@dataclass(frozen=True)
class AddressAssemblyResult:
    output_file: Path
    staging_file: Path
    report_file: Path
    source_files: tuple[Path, ...]
    issues: tuple[str, ...]
    scanned_values: int
    plausible_values: int
    accepted_candidates: int
    normalized_rows: int
    duplicate_rows: int
    rejected_rows: int
    conflict_rows: int
    appended_rows: int
    target_column: str
    target_start_row: int
    target_sheet: str


@dataclass(frozen=True)
class AddressAppendOptions:
    source_file: Path
    target_file: Path | None = None
    source_column: str | None = "A"
    target_column: str | None = "A"
    clean_duplicates: bool = True
    clean_output_workbook: bool = False


@dataclass(frozen=True)
class AddressAppendResult:
    output_file: Path
    report_file: Path
    source_file: Path
    target_file: Path | None
    source_rows: int
    target_existing_rows: int
    accepted_rows: int
    duplicate_rows: int
    skipped_rows: int
    appended_rows: int
    target_column: str
    target_start_row: int
    target_sheet: str
    duplicate_samples: tuple[str, ...]
    skipped_samples: tuple[str, ...]


def _normalizer_report(normalizer: AddressNormalizer, key_profile: dict[str, Any] | None = None) -> dict[str, Any]:
    return normalizer_runtime_report(normalizer, key_profile)


def append_collected_address_workbook(
    *,
    output_file: Path,
    normalizer: AddressNormalizer,
    options: AddressAppendOptions,
    report_file: Path | None = None,
    log: Callable[[str], None] | None = None,
    progress: Callable[[float], None] | None = None,
    cancelled: Callable[[], bool] | None = None,
) -> AddressAppendResult:
    source_file = options.source_file.resolve()
    if not source_file.exists():
        raise FileNotFoundError(f"Collected address workbook was not found: {source_file}")
    target_file = options.target_file.resolve() if options.target_file and options.target_file.exists() else None

    source_items = _address_append_items(
        source_file,
        normalizer=normalizer,
        column=options.source_column or "A",
        default_column="A",
    )
    if progress:
        progress(0.25)
    if _cancelled(cancelled):
        raise RuntimeError("Address append was cancelled.")

    target_items = (
        _address_append_items(
            target_file,
            normalizer=normalizer,
            column=options.target_column or "A",
            default_column="A",
        )
        if target_file
        else []
    )
    existing_keys = {str(item.get("key") or "") for item in target_items if item.get("key")}
    if progress:
        progress(0.45)

    accepted_rows: list[dict[str, Any]] = []
    accepted_items: list[dict[str, Any]] = []
    seen_new_keys: set[str] = set()
    duplicate_samples: list[str] = []
    skipped_samples: list[str] = []
    duplicate_rows = 0
    skipped_rows = 0
    for item in source_items:
        if _cancelled(cancelled):
            raise RuntimeError("Address append was cancelled.")
        address = _clean(item.get("address"))
        key = str(item.get("key") or "")
        if not address or not key:
            skipped_rows += 1
            if len(skipped_samples) < 25:
                skipped_samples.append(str(item.get("raw") or address or ""))
            continue
        if options.clean_duplicates and (key in existing_keys or key in seen_new_keys):
            duplicate_rows += 1
            if len(duplicate_samples) < 25:
                duplicate_samples.append(address)
            continue
        accepted_rows.append({"Normalized_Address": address})
        accepted_items.append(item)
        seen_new_keys.add(key)

    output_rows = accepted_rows
    output_existing_rows = 0
    if options.clean_output_workbook and target_items:
        output_rows = []
        output_keys: set[str] = set()
        for item in target_items:
            address = _clean(item.get("address"))
            key = str(item.get("key") or "")
            if not address or not key:
                continue
            if options.clean_duplicates and key in output_keys:
                continue
            output_rows.append({"Normalized_Address": address})
            output_keys.add(key)
        output_existing_rows = len(output_rows)
        output_rows.extend(accepted_rows)

    if progress:
        progress(0.75)
    target_info = _write_address_output_workbook(
        output_file,
        output_rows,
        target_file=None if options.clean_output_workbook else target_file,
        target_column=options.target_column,
        clean_output_workbook=options.clean_output_workbook,
    )
    report_path = report_file or output_file.with_suffix(".summary.json")
    report_path.parent.mkdir(parents=True, exist_ok=True)
    summary = {
        "mode": "address-clean-column-append",
        "source_file": str(source_file),
        "source_column": options.source_column or "A",
        "target_file": str(target_file) if target_file else None,
        "target_column": target_info["column_letter"],
        "output_file": str(output_file),
        "source_contract": {
            "kind": "clean_address_column",
            "uses_staging": False,
            "dedupe_basis": "stable_address_core",
        },
        "source_rows": len(source_items),
        "target_existing_rows": len(target_items),
        "output_existing_rows": output_existing_rows,
        "accepted_rows": len(accepted_rows),
        "duplicate_rows": duplicate_rows,
        "skipped_rows": skipped_rows,
        "appended_rows": len(accepted_rows),
        "output_rows": int(target_info["written_rows"]),
        "target_sheet": target_info["sheet_name"],
        "target_start_row": target_info["start_row"],
        "normalizer": _normalizer_report(normalizer),
        "address_diagnostics": {
            "source": diagnostic_summary(item["record"] for item in source_items if item.get("record")),
            "accepted": diagnostic_summary(item["record"] for item in accepted_items if item.get("record")),
        },
        "options": {
            "clean_duplicates": bool(options.clean_duplicates),
            "clean_output_workbook": bool(options.clean_output_workbook),
        },
        "duplicate_samples": duplicate_samples,
        "skipped_samples": skipped_samples,
    }
    report_path.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    if progress:
        progress(1.0)
    _log(
        log,
        "Clean address column appended: "
        f"{len(accepted_rows)} new row(s), {duplicate_rows} duplicate(s), {skipped_rows} skipped.",
    )
    return AddressAppendResult(
        output_file=output_file,
        report_file=report_path,
        source_file=source_file,
        target_file=target_file,
        source_rows=len(source_items),
        target_existing_rows=len(target_items),
        accepted_rows=len(accepted_rows),
        duplicate_rows=duplicate_rows,
        skipped_rows=skipped_rows,
        appended_rows=len(accepted_rows),
        target_column=str(target_info["column_letter"]),
        target_start_row=int(target_info["start_row"]),
        target_sheet=str(target_info["sheet_name"]),
        duplicate_samples=tuple(duplicate_samples),
        skipped_samples=tuple(skipped_samples),
    )


def benchmark_address_collection(
    *,
    collected_file: Path,
    benchmark_file: Path,
    normalizer: AddressNormalizer,
    benchmark_column: str | None = "Местоположение",
    collected_column: str | None = "A",
    report_file: Path | None = None,
    sample_limit: int = 100,
    enabled_slots: tuple[str, ...] | None = None,
    slot_order: tuple[str, ...] | None = None,
) -> dict[str, Any]:
    selected_slots = enabled_slots or ("locality", "street", "house", "match_key")
    selected_order = slot_order or DEFAULT_SLOT_ORDER
    benchmark_items = _address_benchmark_items(
        benchmark_file,
        normalizer=normalizer,
        column=benchmark_column or "Местоположение",
        default_column="A",
    )
    collected_items = _address_benchmark_items(
        collected_file,
        normalizer=normalizer,
        column=collected_column or "A",
        default_column="A",
    )
    benchmark_by_key = _items_by_benchmark_key(benchmark_items)
    collected_by_key = _items_by_benchmark_key(collected_items)
    benchmark_keys = set(benchmark_by_key)
    collected_keys = set(collected_by_key)
    matched_keys = benchmark_keys & collected_keys
    missing_keys = benchmark_keys - collected_keys
    extra_keys = collected_keys - benchmark_keys
    benchmark_rows = [_benchmark_item_row(item) for item in benchmark_items]
    collected_rows = [_benchmark_item_row(item) for item in collected_items]
    result = {
        "mode": "address-collection-benchmark",
        "benchmark_file": str(benchmark_file),
        "benchmark_column": benchmark_column or "Местоположение",
        "collected_file": str(collected_file),
        "collected_column": collected_column or "A",
        "benchmark_row_count": len(benchmark_items),
        "benchmark_unique_key_count": len(benchmark_keys),
        "benchmark_duplicate_key_count": max(0, len(benchmark_items) - len(benchmark_keys)),
        "collected_row_count": len(collected_items),
        "collected_unique_key_count": len(collected_keys),
        "collected_duplicate_key_count": max(0, len(collected_items) - len(collected_keys)),
        "matched_unique_key_count": len(matched_keys),
        "missing_unique_key_count": len(missing_keys),
        "extra_unique_key_count": len(extra_keys),
        "coverage_ratio": round(len(matched_keys) / max(1, len(benchmark_keys)), 6),
        "precision_ratio": round(len(matched_keys) / max(1, len(collected_keys)), 6),
        "quality": {
            "coverage_percent": round(len(matched_keys) / max(1, len(benchmark_keys)) * 100, 2),
            "precision_percent": round(len(matched_keys) / max(1, len(collected_keys)) * 100, 2),
            "benchmark_duplicates": max(0, len(benchmark_items) - len(benchmark_keys)),
            "collected_duplicates": max(0, len(collected_items) - len(collected_keys)),
        },
        "address_diagnostics": {
            "benchmark": diagnostic_summary(benchmark_rows),
            "collected": diagnostic_summary(collected_rows),
        },
        "slot_fill": {
            "benchmark": build_slot_fill_summary(
                benchmark_rows,
                enabled_slots=selected_slots,
                slot_order=selected_order,
                slot_columns=SLOT_COLUMN_BY_ID,
                default_slot_order=DEFAULT_SLOT_ORDER,
                get_slot_value=lambda row, slot: row.get(SLOT_COLUMN_BY_ID[slot]),
                get_normalized=lambda row: row.get("Normalized_Address"),
            ),
            "collected": build_slot_fill_summary(
                collected_rows,
                enabled_slots=selected_slots,
                slot_order=selected_order,
                slot_columns=SLOT_COLUMN_BY_ID,
                default_slot_order=DEFAULT_SLOT_ORDER,
                get_slot_value=lambda row, slot: row.get(SLOT_COLUMN_BY_ID[slot]),
                get_normalized=lambda row: row.get("Normalized_Address"),
            ),
        },
        "missing_samples": _benchmark_samples(missing_keys, benchmark_by_key, sample_limit=sample_limit),
        "extra_samples": _benchmark_samples(extra_keys, collected_by_key, sample_limit=sample_limit),
        "artifacts": {
            "outputs": [],
            "reports": [str(report_file)] if report_file else [],
            "staging": [],
            "inputs": [str(benchmark_file), str(collected_file)],
        },
    }
    if report_file:
        report_file.parent.mkdir(parents=True, exist_ok=True)
        report_file.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    return result


@dataclass(frozen=True)
class _CandidateRow:
    row: dict[str, Any]
    record: NormalizedAddress
    candidate: SourceCandidate
    source_kind: str
    quality_score: float


@dataclass(frozen=True)
class _GroupedRows:
    key: str
    rows: tuple[_CandidateRow, ...]


@dataclass(frozen=True)
class AddressMatchCandidate:
    address: str
    raw: str
    source_file: Path
    source_kind: str
    locator: str
    source_row: int | None
    source_column: str | None
    row: dict[str, Any]


def collect_address_match_candidates_from_file(
    file_path: Path,
    *,
    normalizer: AddressNormalizer,
    source_columns: str | None = None,
    enabled_slots: tuple[str, ...] = DEFAULT_SELECTED_SLOTS,
    slot_order: tuple[str, ...] = DEFAULT_SLOT_ORDER,
    omit_microdistrict_when_street_found: bool = True,
    clean_duplicates: bool = True,
    sanitize_incomplete: bool = True,
) -> tuple[list[AddressMatchCandidate], dict[str, Any]]:
    oktmo_key_profile = normalizer_oktmo_key_profile(normalizer)
    candidate_rows, stats, issues = _candidate_rows_from_file(
        file_path,
        normalizer=normalizer,
        source_columns=source_columns,
        reference_columns=None,
        enabled_slots=enabled_slots,
        slot_order=slot_order,
        omit_microdistrict_when_street_found=omit_microdistrict_when_street_found,
        source_kind="source",
        oktmo_key_profile=oktmo_key_profile,
    )
    if not candidate_rows:
        return [], {
            **stats,
            "accepted_candidates": 0,
            "normalized_rows": 0,
            "duplicate_rows": 0,
            "rejected_rows": 0,
            "issues": issues,
        }

    global_context = _dominant_admin_context(candidate_rows)
    groups = _group_candidate_rows(candidate_rows, clean_duplicates=clean_duplicates, global_context=global_context)
    match_candidates: list[AddressMatchCandidate] = []
    rejected_rows = 0
    duplicate_rows = 0
    seen_output_keys: set[str] = set()
    for group in groups:
        duplicate_rows += max(0, len(group.rows) - 1)
        row = _merge_group_rows(
            group,
            enabled_slots=enabled_slots,
            slot_order=slot_order,
            omit_microdistrict_when_street_found=omit_microdistrict_when_street_found,
            global_context=global_context,
        )
        reason = _drop_reason(row) if sanitize_incomplete else None
        if reason:
            rejected_rows += 1
            continue
        normalized = _project_reference_columns(
            row,
            include_slots=True,
            enabled_slots=enabled_slots,
            slot_order=slot_order,
            clean_old_slots=True,
        )
        output_keys = _final_output_dedupe_keys(normalized)
        if clean_duplicates:
            if any(key in seen_output_keys for key in output_keys):
                duplicate_rows += 1
                continue
            seen_output_keys.update(output_keys)
        best = max(group.rows, key=lambda item: item.quality_score)
        source_row, source_column = _candidate_locator_position(best.candidate.locator)
        address = _clean(normalized.get("Normalized_Address")) or _clean(row.get("Normalized_Address")) or ""
        if not address:
            rejected_rows += 1
            continue
        match_candidates.append(
            AddressMatchCandidate(
                address=address,
                raw=_clean(row.get("Raw_Address")) or best.candidate.value,
                source_file=best.candidate.source_file,
                source_kind=best.candidate.source_kind,
                locator=best.candidate.locator,
                source_row=source_row,
                source_column=source_column,
                row=normalized,
            )
        )

    return match_candidates, {
        **stats,
        "accepted_candidates": len(candidate_rows),
        "normalized_rows": len(match_candidates),
        "duplicate_rows": duplicate_rows,
        "rejected_rows": rejected_rows,
        "issues": issues,
        "global_context": global_context,
    }


def _candidate_locator_position(locator: str) -> tuple[int | None, str | None]:
    text = str(locator or "")
    match = re.search(r"!([A-Z]+)(\d+)", text)
    if match:
        return int(match.group(2)) - 1, match.group(1)
    match = re.search(r"!R(\d+)", text)
    if match:
        return int(match.group(1)) - 1, None
    return None, None


def _address_benchmark_items(
    workbook_file: Path,
    *,
    normalizer: AddressNormalizer,
    column: str,
    default_column: str,
) -> list[dict[str, Any]]:
    values = _read_workbook_column_values(workbook_file, column=column, default_column=default_column)
    items: list[dict[str, Any]] = []
    for cell in values:
        text = _clean(cell.get("value"))
        if not text:
            continue
        record = normalizer.normalize(text)
        key = _benchmark_core_key(record)
        if not key:
            continue
        row = _benchmark_record_row(record, text)
        items.append(
            {
                "key": key,
                "address": record.normalized or text,
                "raw": text,
                "sheet": cell.get("sheet"),
                "cell": cell.get("cell"),
                "street": record.parts.street,
                "house": record.parts.house,
                "locality": record.parts.locality,
                "municipality": record.municipality or record.parts.municipality,
                "oktmo_code": record.oktmo_code,
                "row": row,
            }
        )
    return items


def _benchmark_record_row(record: NormalizedAddress, raw: Any) -> dict[str, Any]:
    slots = record.slot_dict()
    row = {
        "Normalized_Address": record.normalized or raw,
        "Raw_Address": raw,
    }
    for slot, column in SLOT_COLUMN_BY_ID.items():
        value = slots.get(slot)
        if isinstance(value, (list, tuple, set, frozenset)):
            value = ", ".join(str(item) for item in value if str(item).strip())
        row[column] = value
    if not row.get("House_Base"):
        row["House_Base"] = record.house_base
    if not row.get("House_Mods"):
        row["House_Mods"] = record.house_mods
    if not row.get("Territory_Key"):
        row["Territory_Key"] = record.territory_key
    if not row.get("Match_Key"):
        row["Match_Key"] = record.match_key
    return row


def _benchmark_item_row(item: dict[str, Any]) -> dict[str, Any]:
    row = item.get("row")
    if isinstance(row, dict):
        return row
    return {
        "Normalized_Address": item.get("address"),
        "Raw_Address": item.get("raw"),
        "Municipality": item.get("municipality"),
        "Locality": item.get("locality"),
        "Street": item.get("street"),
        "House": item.get("house"),
        "OKTMO_Code": item.get("oktmo_code"),
        "Match_Key": item.get("key"),
    }


def _read_workbook_column_values(workbook_file: Path, *, column: str, default_column: str) -> list[dict[str, Any]]:
    workbook = load_workbook(workbook_file, read_only=True, data_only=True)
    values: list[dict[str, Any]] = []
    try:
        for sheet in workbook.worksheets:
            column_index, start_row = _read_column_selection(sheet, column or default_column)
            for row_index in range(start_row, (sheet.max_row or 0) + 1):
                value = sheet.cell(row=row_index, column=column_index).value
                if _clean(value):
                    values.append(
                        {
                            "sheet": sheet.title,
                            "cell": f"{excel_column_letter(column_index - 1)}{row_index}",
                            "value": value,
                        }
                    )
    finally:
        workbook.close()
    return values


def _read_column_selection(sheet: Any, token: str) -> tuple[int, int]:
    text = str(token or "").strip()
    if not text:
        text = "A"
    if text.isdigit():
        return max(1, int(text)), 1
    if re.fullmatch(r"[A-Za-z]{1,3}", text):
        return excel_column_index(text) + 1, 1
    header = _find_header_column(sheet, text)
    if header:
        column_index, header_row = header
        return column_index, header_row + 1
    fallback = str(token or "").strip() or "A"
    if fallback != "A":
        raise ValueError(f"Workbook column {token!r} was not found by letter, number, or header name.")
    return 1, 1


def _address_append_items(
    workbook_file: Path,
    *,
    normalizer: AddressNormalizer,
    column: str,
    default_column: str,
) -> list[dict[str, Any]]:
    values = _read_workbook_column_values(workbook_file, column=column, default_column=default_column)
    items: list[dict[str, Any]] = []
    for cell in values:
        text = _clean(cell.get("value"))
        if not text or _looks_like_address_column_header(text):
            continue
        record = normalizer.normalize(text)
        address = _clean(record.normalized) or text
        key = _append_address_core_key(record, text)
        items.append(
            {
                "key": key,
                "address": address,
                "raw": text,
                "record": record,
                "sheet": cell.get("sheet"),
                "cell": cell.get("cell"),
            }
        )
    return items


def _append_address_core_key(record: NormalizedAddress, raw: Any) -> str:
    street_text = _clean(record.parts.street)
    house_text = _clean(record.parts.house)
    street_words = record.street_words
    street_numbers = set(record.street_numbers)
    house_base = record.house_base
    house_mods = _safe_house_mods(record.house_mods or record.parts.house_mods)

    if street_text:
        _territory, parsed_street_words, parsed_street_numbers, _street_house_base, _street_house_mods = parse_address_components(
            street_text
        )
        street_words = parsed_street_words or street_words
        street_numbers.update(parsed_street_numbers)
    if house_text:
        _house_territory, _house_street_words, _house_street_numbers, parsed_house_base, parsed_house_mods = parse_address_components(
            house_text
        )
        house_base = parsed_house_base or house_base
        house_mods = _safe_house_mods(parsed_house_mods or house_mods)

    house_kind = _norm(record.parts.house_kind) or _house_kind_key(house_text)
    premise = _norm(record.parts.premise)
    if street_words and house_base:
        return "|".join(
            part
            for part in (
                "core",
                street_words,
                ",".join(sorted(street_numbers)),
                house_kind,
                house_base,
                house_mods,
                premise,
            )
            if part
        )

    normalized = _norm(record.normalized or raw)
    return f"text|{normalized}" if normalized else ""


def _looks_like_address_column_header(value: Any) -> bool:
    key = _norm(value)
    if not key:
        return False
    known = {
        "normalized address",
        "address",
        "адрес",
        "адрес объекта",
        "адрес полностью",
        "почтовый адрес",
        "фактический адрес",
        "местоположение",
        "местонахождение",
        "место нахождения",
    }
    return key in known or any(key == _norm(marker) for marker in ADDRESS_HEADER_MARKERS)


def _benchmark_core_key(record: NormalizedAddress) -> str:
    if not (record.street_words and record.house_base):
        return ""
    numbers = ",".join(sorted(record.street_numbers))
    return "|".join(
        part
        for part in (
            "core",
            record.street_words,
            numbers,
            record.house_base,
            record.house_mods,
        )
        if part
    )


def _items_by_benchmark_key(items: list[dict[str, Any]]) -> dict[str, list[dict[str, Any]]]:
    result: dict[str, list[dict[str, Any]]] = {}
    for item in items:
        key = str(item.get("key") or "")
        if key:
            result.setdefault(key, []).append(item)
    return result


def _benchmark_samples(
    keys: set[str],
    items_by_key: dict[str, list[dict[str, Any]]],
    *,
    sample_limit: int,
) -> list[dict[str, Any]]:
    samples: list[dict[str, Any]] = []
    for key in sorted(keys)[:sample_limit]:
        item = items_by_key[key][0]
        samples.append(
            {
                "key": key,
                "address": item.get("address"),
                "raw": item.get("raw"),
                "sheet": item.get("sheet"),
                "cell": item.get("cell"),
            }
        )
    return samples


def assemble_address_workbook(
    *,
    input_dir: Path,
    output_file: Path,
    normalizer: AddressNormalizer,
    options: AddressAssemblyOptions,
    report_file: Path | None = None,
    log: Callable[[str], None] | None = None,
    progress: Callable[[float], None] | None = None,
    cancelled: Callable[[], bool] | None = None,
) -> AddressAssemblyResult:
    action = str(options.action or "generate").strip().lower()
    if action not in {"generate", "update"}:
        raise ValueError(f"Unsupported address assembly action: {options.action!r}")

    source_root = (options.source_root or input_dir).resolve()
    source_files = discover_source_files(source_root)
    target_dir = (options.target_dir or output_file.parent).resolve()
    target_dir.mkdir(parents=True, exist_ok=True)
    target_file = (
        options.target_file.resolve()
        if options.target_file and options.target_file.exists()
        else _discover_existing_address_workbook(target_dir)
        if action == "update"
        else None
    )
    reference_file = options.reference_file.resolve() if options.reference_file and options.reference_file.exists() else target_file

    issues: list[str] = []
    candidate_rows: list[_CandidateRow] = []
    scanned_values = 0
    plausible_values = 0
    oktmo_key_profile = normalizer_oktmo_key_profile(normalizer)

    if reference_file:
        reference_file = reference_file.resolve()
        existing_rows, stats, existing_issues = _candidate_rows_from_file(
            reference_file,
            normalizer=normalizer,
            source_columns=None,
            reference_columns=options.reference_columns or options.target_column,
            enabled_slots=options.enabled_slots,
            slot_order=options.slot_order,
            omit_microdistrict_when_street_found=options.omit_microdistrict_when_street_found,
            source_kind="existing_reference",
            oktmo_key_profile=oktmo_key_profile,
        )
        candidate_rows.extend(existing_rows)
        scanned_values += stats["scanned_values"]
        plausible_values += stats["plausible_values"]
        issues.extend(existing_issues)
        _log(log, f"Existing address workbook loaded for duplicate guard: {reference_file}")
    elif action == "update":
        _log(log, "No existing address workbook found; update will create a new table.")

    processed_files: list[Path] = []
    total_files = max(1, len(source_files))
    for index, file_path in enumerate(source_files, start=1):
        if _cancelled(cancelled):
            raise RuntimeError("Address assembly was cancelled.")
        if reference_file and file_path.resolve() == reference_file:
            continue
        rows, stats, file_issues = _candidate_rows_from_file(
            file_path,
            normalizer=normalizer,
            source_columns=options.source_columns,
            reference_columns=None,
            enabled_slots=options.enabled_slots,
            slot_order=options.slot_order,
            omit_microdistrict_when_street_found=options.omit_microdistrict_when_street_found,
            source_kind="source",
            oktmo_key_profile=oktmo_key_profile,
        )
        candidate_rows.extend(rows)
        scanned_values += stats["scanned_values"]
        plausible_values += stats["plausible_values"]
        issues.extend(file_issues)
        processed_files.append(file_path)
        if rows:
            _log(log, f"  {file_path.name}: {len(rows)} accepted address candidate(s)")
        if progress:
            progress(index / total_files * 0.85)

    if not candidate_rows:
        raise RuntimeError("No address candidates were collected from the selected input sources.")
    issues = list(dict.fromkeys(issues))

    global_context = _dominant_admin_context(candidate_rows)
    if global_context:
        _log(
            log,
            "Dominant administrative context: "
            + ", ".join(f"{key}={value}" for key, value in global_context.items() if value),
        )
    grouped = _group_candidate_rows(candidate_rows, clean_duplicates=options.clean_duplicates, global_context=global_context)
    normalized_rows, rejected_rows, duplicate_rows, staged_rows = _final_rows_from_groups(
        grouped,
        enabled_slots=options.enabled_slots,
        slot_order=options.slot_order,
        omit_microdistrict_when_street_found=options.omit_microdistrict_when_street_found,
        include_slots=options.include_slots,
        sanitize_incomplete=options.sanitize_incomplete,
        global_context=global_context,
    )
    if not normalized_rows:
        raise RuntimeError("All collected address candidates were rejected by the sanitizer.")
    conflict_rows = _conflict_feedback_rows_from_rejected(rejected_rows)

    staging_file = options.staging_file or output_file.with_suffix(".stage.xlsx")
    _write_staging_workbook(
        staging_file,
        staged_rows,
        normalized_rows,
        rejected_rows if options.write_rejected else [],
        conflict_rows if options.write_rejected else [],
    )
    output_file.parent.mkdir(parents=True, exist_ok=True)
    target_info = _write_address_output_workbook(
        output_file,
        normalized_rows,
        target_file=None if options.clean_output_workbook else options.target_file,
        target_column=options.target_column,
        clean_output_workbook=options.clean_output_workbook,
    )

    report_path = report_file or output_file.with_suffix(".summary.json")
    report_path.parent.mkdir(parents=True, exist_ok=True)
    summary = {
        "mode": "address-source-table-assembly",
        "action": action,
        "input_dir": str(input_dir),
        "source_root": str(source_root),
        "staging_file": str(staging_file),
        "output_file": str(output_file),
        "target_dir": str(target_dir),
        "target_file": str(target_file) if target_file else None,
        "source_file_count": len(processed_files),
        "scanned_values": scanned_values,
        "plausible_values": plausible_values,
        "accepted_candidates": len(candidate_rows),
        "normalized_rows": len(normalized_rows),
        "duplicate_rows": duplicate_rows,
        "rejected_rows": len(rejected_rows),
        "conflict_rows": len(conflict_rows),
        "feedback_conflict_rows": len(conflict_rows),
        "appended_rows": target_info["appended_rows"],
        "output_rows": target_info["appended_rows"],
        "address_diagnostics": {
            "normalized": diagnostic_summary(normalized_rows),
            "rejected": diagnostic_summary(rejected_rows),
            "conflicts": diagnostic_summary([row for row in rejected_rows if row.get("Drop_Reason") == "oktmo_scope_conflict"]),
        },
        "slot_fill": build_slot_fill_summary(
            normalized_rows,
            enabled_slots=options.enabled_slots,
            slot_order=options.slot_order,
            slot_columns=SLOT_COLUMN_BY_ID,
            default_slot_order=DEFAULT_SLOT_ORDER,
            get_slot_value=lambda row, slot: row.get(SLOT_COLUMN_BY_ID[slot]),
            get_normalized=lambda row: row.get("Normalized_Address"),
        ),
        "target_sheet": target_info["sheet_name"],
        "target_column": target_info["column_letter"],
        "target_start_row": target_info["start_row"],
        "normalizer": _normalizer_report(normalizer, oktmo_key_profile),
        "oktmo_scope_matches": _oktmo_scope_match_summary(candidate_rows),
        "options": {
            "source_columns": options.source_columns,
            "reference_columns": options.reference_columns,
            "enabled_slots": list(options.enabled_slots),
            "slot_order": list(options.slot_order),
            "omit_microdistrict_when_street_found": bool(options.omit_microdistrict_when_street_found),
            "include_slots": bool(options.include_slots),
            "clean_output_workbook": bool(options.clean_output_workbook),
            "clean_duplicates": bool(options.clean_duplicates),
            "sanitize_incomplete": bool(options.sanitize_incomplete),
            "write_rejected": bool(options.write_rejected),
        },
        "issues": issues[:200],
        "sources": [str(path) for path in processed_files],
        "rejected_samples": rejected_rows[:50],
        "conflict_samples": conflict_rows[:50],
    }
    report_path.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    if progress:
        progress(1.0)

    return AddressAssemblyResult(
        output_file=output_file,
        staging_file=staging_file,
        report_file=report_path,
        source_files=tuple(processed_files),
        issues=tuple(issues),
        scanned_values=scanned_values,
        plausible_values=plausible_values,
        accepted_candidates=len(candidate_rows),
        normalized_rows=len(normalized_rows),
        duplicate_rows=duplicate_rows,
        rejected_rows=len(rejected_rows),
        conflict_rows=len(conflict_rows),
        appended_rows=int(target_info["appended_rows"]),
        target_column=str(target_info["column_letter"]),
        target_start_row=int(target_info["start_row"]),
        target_sheet=str(target_info["sheet_name"]),
    )


def _candidate_rows_from_file(
    file_path: Path,
    *,
    normalizer: AddressNormalizer,
    source_columns: str | None,
    reference_columns: str | None,
    enabled_slots: tuple[str, ...],
    slot_order: tuple[str, ...],
    omit_microdistrict_when_street_found: bool,
    source_kind: str,
    oktmo_key_profile: dict[str, Any] | None = None,
) -> tuple[list[_CandidateRow], dict[str, int], list[str]]:
    if source_kind == "existing_reference":
        candidates, issues = extract_reference_candidates(file_path, reference_columns=reference_columns)
    else:
        raw_candidates: list[SourceCandidate] = []
        issues: list[str] = []
        suffix = file_path.suffix.lower()
        if suffix == ".pdf":
            raw_candidates, issues = extract_source_candidates(file_path, source_columns=source_columns)
            fallback_for_observations = raw_candidates
        else:
            fallback_for_observations = []
        observation_candidates, observation_issues = _extract_registry_style_source_candidates(
            file_path,
            source_columns=source_columns,
            fallback_candidates=fallback_for_observations,
        )
        issues.extend(observation_issues)
        if not observation_candidates or suffix in {".txt", ".csv", ".md", ".markdown"}:
            raw_candidates, raw_issues = extract_source_candidates(file_path, source_columns=source_columns)
            issues.extend(raw_issues)
        structured_candidates, structured_issues = _extract_structured_source_candidates(file_path, source_columns=source_columns)
        if observation_candidates and suffix in {".txt", ".csv", ".md", ".markdown"}:
            primary_candidates = [*observation_candidates, *raw_candidates]
        else:
            primary_candidates = observation_candidates or raw_candidates
        fragment_candidates = _address_fragment_candidates_from_candidates([*primary_candidates, *structured_candidates])
        candidates = _deduplicate_source_candidates(
            [
                *primary_candidates,
                *_contextual_candidates_from_candidates(primary_candidates),
                *structured_candidates,
                *fragment_candidates,
            ]
        )
        issues.extend(structured_issues)

    rows: list[_CandidateRow] = []
    scanned_values = 0
    plausible_values = 0
    for candidate in candidates:
        scanned_values += 1
        if not _plausible_address_text(candidate.value):
            continue
        plausible_values += 1
        record = normalizer.normalize(candidate.value)
        if not _has_collectable_address_core(record):
            continue
        row = _reference_row(
            record,
            enabled_slots=enabled_slots,
            slot_order=slot_order,
            omit_microdistrict_when_street_found=omit_microdistrict_when_street_found,
            candidate=candidate,
            source_kind=source_kind,
        )
        scope_match = _oktmo_scope_candidate_match(row, record, candidate, oktmo_key_profile)
        if scope_match:
            row.update(_oktmo_scope_candidate_diagnostics(scope_match))
        rows.append(
            _CandidateRow(
                row=row,
                record=record,
                candidate=candidate,
                source_kind=source_kind,
                quality_score=_candidate_quality_score(row, record, candidate, source_kind, oktmo_scope_match=scope_match),
            )
        )
    return rows, {"scanned_values": scanned_values, "plausible_values": plausible_values}, issues


def _contextual_candidates_from_candidates(candidates: list[SourceCandidate]) -> list[SourceCandidate]:
    result: list[SourceCandidate] = []
    result.extend(_contextual_xlsx_row_candidates(candidates))
    result.extend(_contextual_sequential_candidates(candidates))
    return result


def _address_fragment_candidates_from_candidates(candidates: list[SourceCandidate]) -> list[SourceCandidate]:
    result: list[SourceCandidate] = []
    for candidate in candidates:
        for index, fragment in enumerate(_extract_clean_address_fragments(candidate.value), start=1):
            result.append(
                SourceCandidate(
                    value=fragment,
                    source_file=candidate.source_file,
                    source_kind=f"{candidate.source_kind}_address_fragment",
                    locator=f"{candidate.locator}:F{index}",
                )
            )
    return _deduplicate_source_candidates(result)


def _extract_clean_address_fragments(value: Any) -> list[str]:
    text = _clean(_normalize_spaced_address_marker_dots(value))
    if not text:
        return []
    if len(text) < 8 or len(text) > 1200:
        return []
    marker = (
        r"ул\.?|улица|пер\.?|переулок|проспект|пр-т\.?|просп\.?|пр-кт\.?|"
        r"проезд|пр-д\.?|бульвар|б-р\.?|набережная|наб\.?|площадь|пл\.?|"
        r"шоссе|ш\.?|тупик|аллея|линия"
    )
    house = (
        r"(?:д\.?|дом|зд\.?|здание|стр\.?|строение)?\s*"
        r"\d{1,4}[а-яa-z]?\s*(?:[-/]\s*[а-яa-z0-9]+)?"
        r"(?:\s*,?\s*(?:к\.?|корп\.?|корпус|с\.?|стр\.?|строение|лит\.?а?|литера)\s*[а-яa-z0-9]+)?"
    )
    patterns = (
        rf"(?P<street>\b(?:{marker})\s*[А-ЯЁA-Zа-яёa-z0-9 .'\-]{{2,80}}?)\s*,?\s+(?P<house>{house})(?=$|[),;–—-]|\s*,\s*\d)",
        rf"(?P<street>\b[А-ЯЁA-Zа-яёa-z0-9 .'\-]{{2,80}}\s+(?:тракт|тракта|шоссе|проезд|бульвар|переулок|проспект|набережная|площадь))\s*,?\s+(?P<house>{house})(?=$|[),;–—-]|\s*,\s*\d)",
    )
    fragments: list[str] = []
    for pattern in patterns:
        for match in re.finditer(pattern, text, flags=re.IGNORECASE):
            if _fragment_is_parenthesized_org_origin(text, match.start(), match.end()):
                continue
            raw_house = str(match.group("house") or "").strip()
            if _fragment_house_is_ordinal_street_number(text, match.end(), raw_house):
                continue
            street = _normalize_street_slot(match.group("street"))
            house_value = _normalize_house_slot(raw_house)
            if not street or not house_value:
                continue
            if _strict_street_slot_drop_reason(street) or not _looks_like_ideal_house_slot(house_value):
                continue
            fragment = f"{street}, {house_value}"
            if _contains_org_context(fragment) or _looks_like_route_fragment(fragment):
                continue
            _append_unique(fragments, fragment)
    if len(fragments) > 8:
        return fragments[-8:]
    return fragments


def _fragment_is_parenthesized_org_origin(text: str, start: int, end: int) -> bool:
    left_paren = text.rfind("(", 0, start)
    right_paren = text.find(")", start, max(end, start) + 40)
    if left_paren < 0:
        return False
    before = text[max(0, left_paren - 120) : left_paren]
    inside = text[left_paren : right_paren + 1] if right_paren >= 0 else text[left_paren:end]
    after = text[right_paren + 1 : right_paren + 80] if right_paren >= 0 else text[end : end + 80]
    if not _contains_org_context(before + " " + inside):
        return False
    return bool(re.search(r"\s[-–—]\s", after) or re.search(r"\b(?:с/о|дк|жк|кп|мкр)\b", after, flags=re.IGNORECASE))


def _fragment_house_is_ordinal_street_number(text: str, house_end: int, raw_house: str | None) -> bool:
    house = str(raw_house or "").strip().lower().replace("ё", "е")
    if re.search(r"\d+\s*[-–—]?\s*(?:я|ая|й|ой)\b", house):
        return True
    if not re.fullmatch(r"(?:д\.?|дом)?\s*\d{1,3}[а-яa-z]?", house):
        return False
    tail = str(text or "")[house_end : house_end + 20].lower().replace("ё", "е")
    return bool(re.match(r"\s*[-–—]\s*(?:я|ая|й|ой)\b", tail))


def _extract_registry_style_source_candidates(
    file_path: Path,
    *,
    source_columns: str | None,
    fallback_candidates: list[SourceCandidate],
) -> tuple[list[SourceCandidate], list[str]]:
    suffix = file_path.suffix.lower()
    try:
        if suffix == ".xlsx":
            return _extract_xlsx_observation_candidates(file_path, source_columns=source_columns), []
        if suffix in {".txt", ".csv", ".md", ".markdown"}:
            return _extract_text_observation_candidates(file_path), []
        if suffix == ".docx":
            return _extract_docx_observation_candidates(file_path), []
        if suffix == ".odt":
            return _extract_odt_observation_candidates(file_path), []
        if suffix == ".pdf":
            return _extract_block_observation_candidates(file_path, fallback_candidates, source_kind="pdf_observation"), []
    except Exception as exc:
        return [], [f"{file_path.name}: registry-style address observation intake failed: {exc}"]
    return [], []


def _extract_xlsx_observation_candidates(file_path: Path, *, source_columns: str | None) -> list[SourceCandidate]:
    workbook = pd.read_excel(file_path, sheet_name=None, header=None, dtype=object)
    selected = _source_column_filter(source_columns)
    result: list[SourceCandidate] = []
    for sheet_name, frame in workbook.items():
        if frame.empty:
            continue
        frame = frame.dropna(axis=0, how="all")
        if frame.empty:
            continue
        header_index = _detect_observation_header_row(frame)
        if header_index is None:
            continue
        headers, data_start = _build_observation_headers(frame, header_index)
        context_lines: list[str] = []
        for row_index in range(data_start, len(frame)):
            values = [_clean_excel_value(frame.iat[row_index, column_index]) for column_index in range(len(frame.columns))]
            selected_values = _values_for_selected_columns(values, selected)
            if not any(selected_values):
                continue
            non_empty = [value for value in selected_values if value]
            row_text = _clean(" | ".join(non_empty)) or ""
            if _looks_like_admin_context(row_text) and not _has_street_or_house_text(row_text):
                _append_context_line(context_lines, row_text)
                continue
            payload = _payload_from_headers(headers, values, selected_columns=selected)
            candidate = _address_candidate_from_payload(payload, context_lines=context_lines)
            if not candidate:
                continue
            result.append(
                SourceCandidate(
                    value=candidate,
                    source_file=file_path,
                    source_kind="xlsx_observation_row",
                    locator=f"{sheet_name}!R{row_index + 1}",
                )
            )
    return result


def _extract_text_observation_candidates(file_path: Path) -> list[SourceCandidate]:
    text = _read_source_text(file_path)
    result = _extract_delimited_text_observation_candidates(file_path, text)
    result.extend(_extract_plain_text_observation_candidates(file_path, text, source_kind=f"{file_path.suffix.lower().lstrip('.') or 'txt'}_observation"))
    return _deduplicate_source_candidates(result)


def _extract_docx_observation_candidates(file_path: Path) -> list[SourceCandidate]:
    with zipfile.ZipFile(file_path) as archive:
        root = ET.fromstring(archive.read("word/document.xml"))
    result: list[SourceCandidate] = []
    for table_index, table in enumerate(root.findall(".//w:tbl", WORD_NS), start=1):
        rows = [
            [_clean_xml_text(_docx_node_text(cell)) for cell in row.findall("./w:tc", WORD_NS)]
            for row in table.findall("./w:tr", WORD_NS)
        ]
        result.extend(
            _observation_candidates_from_cell_rows(
                file_path,
                rows,
                source_kind="docx_observation_row",
                locator_prefix=f"T{table_index}",
            )
        )
    paragraphs = [
        _clean_xml_text(_docx_node_text(paragraph))
        for paragraph in root.findall(".//w:body/w:p", WORD_NS)
    ]
    result.extend(
        _plain_text_candidates_from_blocks(
            file_path,
            [paragraph for paragraph in paragraphs if paragraph],
            source_kind="docx_observation",
            locator_prefix="P",
        )
    )
    return _deduplicate_source_candidates(result)


def _extract_odt_observation_candidates(file_path: Path) -> list[SourceCandidate]:
    with zipfile.ZipFile(file_path) as archive:
        root = ET.fromstring(archive.read("content.xml"))
    result: list[SourceCandidate] = []
    for table_index, table in enumerate(root.findall(".//table:table", ODT_NS), start=1):
        rows = [_odt_row_values(row) for row in table.findall("./table:table-row", ODT_NS)]
        result.extend(
            _observation_candidates_from_cell_rows(
                file_path,
                rows,
                source_kind="odt_observation_row",
                locator_prefix=f"T{table_index}",
            )
        )
    body = root.find("office:body/office:text", ODT_NS)
    paragraphs: list[str] = []
    if body is not None:
        for child in list(body):
            if _local_xml_name(child.tag) in {"p", "h"}:
                text = _clean_xml_text(_xml_node_text(child))
                if text:
                    paragraphs.append(text)
    result.extend(
        _plain_text_candidates_from_blocks(
            file_path,
            paragraphs,
            source_kind="odt_observation",
            locator_prefix="P",
        )
    )
    return _deduplicate_source_candidates(result)


def _extract_block_observation_candidates(
    file_path: Path,
    candidates: list[SourceCandidate],
    *,
    source_kind: str,
) -> list[SourceCandidate]:
    blocks = [candidate.value for candidate in candidates if _clean(candidate.value)]
    return _plain_text_candidates_from_blocks(file_path, blocks, source_kind=source_kind, locator_prefix="B")


def _extract_delimited_text_observation_candidates(file_path: Path, text: str) -> list[SourceCandidate]:
    lines = [(index, line.strip()) for index, line in enumerate(_source_lines(text), start=1) if line.strip()]
    if not lines:
        return []
    parsed_rows: list[tuple[int, list[str | None]]] = []
    for line_number, line in lines:
        cells = _parse_observation_table_line(line, suffix=file_path.suffix.lower())
        if cells and len(cells) >= 2:
            parsed_rows.append((line_number, cells))
    if not parsed_rows:
        return []
    header_position: int | None = None
    for index, (_line_number, cells) in enumerate(parsed_rows[:50]):
        if _score_observation_header_values([cell or "" for cell in cells]) >= 4:
            header_position = index
            break
    if header_position is None:
        return []
    headers = _make_observation_headers(parsed_rows[header_position][1])
    result: list[SourceCandidate] = []
    context_lines: list[str] = []
    for line_number, cells in parsed_rows[header_position + 1 :]:
        if _is_repeated_observation_header(headers, cells):
            continue
        payload = _payload_from_headers(headers, cells)
        row_text = _clean(" | ".join(str(value) for value in cells if value)) or ""
        if _looks_like_admin_context(row_text) and not _has_street_or_house_text(row_text):
            _append_context_line(context_lines, row_text)
            continue
        candidate = _address_candidate_from_payload(payload, context_lines=context_lines)
        if candidate:
            result.append(
                SourceCandidate(
                    value=candidate,
                    source_file=file_path,
                    source_kind=f"{file_path.suffix.lower().lstrip('.') or 'txt'}_observation_row",
                    locator=f"L{line_number}",
                )
            )
    return result


def _extract_plain_text_observation_candidates(file_path: Path, text: str, *, source_kind: str) -> list[SourceCandidate]:
    return _plain_text_candidates_from_blocks(file_path, _split_observation_blocks(text), source_kind=source_kind, locator_prefix="B")


def _plain_text_candidates_from_blocks(
    file_path: Path,
    blocks: list[str],
    *,
    source_kind: str,
    locator_prefix: str,
) -> list[SourceCandidate]:
    result: list[SourceCandidate] = []
    context_lines: list[str] = []
    for index, block in enumerate(blocks, start=1):
        text = _clean(block)
        if not text:
            continue
        payload = _parse_observation_kv_payload(text)
        if payload:
            candidate = _address_candidate_from_payload(payload, context_lines=context_lines)
            if candidate:
                result.append(
                    SourceCandidate(
                        value=candidate,
                        source_file=file_path,
                        source_kind=f"{source_kind}_record",
                        locator=f"{locator_prefix}{index}",
                    )
                )
                continue
        if _looks_like_admin_context(text) and not _has_street_or_house_text(text):
            _append_context_line(context_lines, text)
            continue
        if _strong_plain_address_text(text):
            candidate = _contextual_address_value(context_lines, [text]) or text
            result.append(
                SourceCandidate(
                    value=candidate,
                    source_file=file_path,
                    source_kind=f"{source_kind}_paragraph",
                    locator=f"{locator_prefix}{index}",
                )
            )
    return result


def _observation_candidates_from_cell_rows(
    file_path: Path,
    rows: list[list[str | None]],
    *,
    source_kind: str,
    locator_prefix: str,
) -> list[SourceCandidate]:
    result: list[SourceCandidate] = []
    headers: list[str] = []
    context_lines: list[str] = []
    for row_index, cells in enumerate(rows, start=1):
        values = [_clean(value) for value in cells]
        if not any(values):
            continue
        if _score_observation_header_values([value or "" for value in values]) >= 4:
            headers = _make_observation_headers(values)
            continue
        row_text = _clean(" | ".join(value for value in values if value)) or ""
        if _looks_like_admin_context(row_text) and not _has_street_or_house_text(row_text):
            _append_context_line(context_lines, row_text)
            continue
        if headers:
            payload = _payload_from_headers(headers, values)
            candidate = _address_candidate_from_payload(payload, context_lines=context_lines)
        else:
            candidate = _contextual_address_value(context_lines, [row_text]) if _strong_plain_address_text(row_text) else None
        if not candidate:
            continue
        result.append(
            SourceCandidate(
                value=candidate,
                source_file=file_path,
                source_kind=source_kind,
                locator=f"{locator_prefix}:R{row_index}",
            )
        )
    return result


def _detect_observation_header_row(frame: pd.DataFrame) -> int | None:
    best_index: int | None = None
    best_score = 0
    for row_index in range(min(len(frame.index), 35)):
        values = [_clean_excel_value(value) for value in frame.iloc[row_index].tolist()]
        score = _score_observation_header_values([value or "" for value in values])
        if score > best_score:
            best_score = score
            best_index = row_index
    return best_index if best_score >= 4 else None


def _build_observation_headers(frame: pd.DataFrame, header_index: int) -> tuple[list[str], int]:
    header_values = [_clean_excel_value(value) for value in frame.iloc[header_index].tolist()]
    continuation_rows: list[list[str | None]] = []
    for row_index in range(header_index + 1, min(len(frame.index), header_index + 4)):
        values = [_clean_excel_value(value) for value in frame.iloc[row_index].tolist()]
        if not _is_observation_header_continuation(values, header_values):
            break
        continuation_rows.append(values)
    combined: list[str | None] = []
    last_group_header: str | None = None
    for column_index, base_value in enumerate(header_values):
        if base_value:
            last_group_header = base_value
        parts: list[str] = []
        if base_value:
            parts.append(base_value)
        continuation_values = [
            row[column_index]
            for row in continuation_rows
            if column_index < len(row) and row[column_index]
        ]
        if not base_value and continuation_values and last_group_header:
            parts.append(last_group_header)
        for value in continuation_values:
            if value and value not in parts:
                parts.append(value)
        combined.append(" ".join(parts) if parts else None)
    return _make_observation_headers(combined), header_index + 1 + len(continuation_rows)


def _is_observation_header_continuation(
    values: list[str | None],
    header_values: list[str | None] | None = None,
) -> bool:
    non_empty = [value for value in values if value]
    if not non_empty:
        return False
    if _looks_like_structural_subheader(values, header_values):
        return True
    joined = _norm(" | ".join(str(value) for value in non_empty))
    if _has_street_or_house_text(joined) or "местоположение" in joined or "адрес" in joined:
        return False
    headerish = sum(1 for value in non_empty if _observation_header_slot(value) or re.fullmatch(r"20\d{2}", str(value)))
    return headerish >= max(2, len(non_empty) // 2)


def _looks_like_structural_subheader(
    values: list[str | None],
    header_values: list[str | None] | None,
) -> bool:
    """Recognize a second header row by its shape rather than by its wording.

    The slot test above only accepts sub-headers whose words map to a known
    address slot or a year, so a table headed "Проектная мощность, мест" over
    "по корпусам | при обучении в одну смену" kept the numbers and lost the words
    describing them — and the address column beside them lost its parent name too.

    Two things identify the row without any vocabulary: it carries no data
    values, and it fills the gaps a horizontally merged parent header left behind.
    """
    if not header_values:
        return False
    non_empty = [(index, str(value)) for index, value in enumerate(values) if value]
    if not non_empty:
        return False

    fills_a_merge_gap = False
    for index, _ in non_empty:
        if index >= len(header_values) or header_values[index]:
            continue
        # A gap only counts when a header actually opened to the left of it.
        if any(header_values[left] for left in range(index)):
            fills_a_merge_gap = True
            break
    if not fills_a_merge_gap:
        return False

    return not any(_looks_like_table_data_value(value) for _, value in non_empty)


def _looks_like_table_data_value(value: str) -> bool:
    """True for anything that reads as a measurement rather than a column name."""
    text = re.sub(r"\s+", " ", str(value or "")).strip()
    if not text:
        return False
    bare_number = text.replace(",", ".").replace(" ", "")
    if re.fullmatch(r"-?\d+(?:\.\d+)?", bare_number) and not re.fullmatch(r"(?:19|20)\d{2}", bare_number):
        return True
    if re.search(r"\d+[.,]\d+", text):
        return True
    if looks_like_address_text(text):
        return True
    # Column names stay short; a filled cell runs on.
    return len(text) > 120


def _score_observation_header_values(values: list[str]) -> int:
    non_empty = [value for value in values if _clean(value)]
    if len(non_empty) < 2:
        return 0
    joined = _header_norm(" | ".join(non_empty))
    score = 0
    for value in non_empty:
        if _observation_header_slot(value):
            score += 2
    if any(marker in joined for marker in ("адрес", "местоположение", "местонахождение")):
        score += 4
    if any(marker in joined for marker in ("улица", "номер дома", "дом")):
        score += 2
    if any(marker in joined for marker in ("населенный пункт", "город", "муницип")):
        score += 1
    if len(non_empty) >= 5:
        score += 1
    return score


def _make_observation_headers(values: list[object]) -> list[str]:
    headers: list[str] = []
    seen: dict[str, int] = {}
    for index, value in enumerate(values, start=1):
        header = _clean_excel_value(value) or f"column_{index}"
        header = re.sub(r"\s+", " ", header).strip()
        count = seen.get(header, 0) + 1
        seen[header] = count
        if count > 1:
            header = f"{header}_{count}"
        headers.append(header)
    return headers


def _payload_from_headers(
    headers: list[str],
    values: list[Any],
    *,
    selected_columns: set[int] | None = None,
) -> dict[str, str | None]:
    payload: dict[str, str | None] = {}
    for index, header in enumerate(headers):
        if selected_columns is not None and index not in selected_columns:
            continue
        value = _clean_excel_value(values[index]) if index < len(values) else None
        payload[header] = value
    return payload


def _address_candidate_from_payload(payload: dict[str, str | None], *, context_lines: list[str]) -> str | None:
    payload = {str(key): _clean(value) for key, value in payload.items() if _clean(value)}
    if not payload:
        return None
    slots = _payload_address_slots(payload)
    address = _pick_payload_address_value(payload)
    if address:
        address = _strip_non_address_tail(address)
        address = _complete_address_from_payload_slots(address, slots)
        return _add_payload_context_to_address(address, slots, context_lines)
    street = slots.get("Street")
    house = slots.get("House")
    if street and house:
        return _address_from_slots(slots, context_lines=context_lines)
    fallback = _pick_address_like_payload_value(payload)
    if fallback:
        fallback = _complete_address_from_payload_slots(fallback, slots)
        return _add_payload_context_to_address(fallback, slots, context_lines)
    return None


def _payload_address_slots(payload: dict[str, str | None]) -> dict[str, str]:
    slots: dict[str, str] = {}
    for header, value in payload.items():
        text = _clean(value)
        if not text:
            continue
        slot = _observation_header_slot(header)
        if not slot:
            continue
        if slot == "Street":
            formatted = _format_structured_slot("Street", text)
        elif slot == "House":
            formatted = _format_structured_slot("House", text)
        else:
            formatted = text
        if formatted and slot not in slots:
            slots[slot] = formatted
    return slots


def _pick_payload_address_value(payload: dict[str, str | None]) -> str | None:
    candidates: list[tuple[int, int, str]] = []
    for header, value in payload.items():
        text = _clean(value)
        if not text or _placeholder_text(text):
            continue
        header_key = _header_norm(header)
        if not _header_matches(header_key, ADDRESS_HEADER_MARKERS):
            continue
        score = 20
        if "местоположение" in header_key:
            score += 4
        if "кадастр" in header_key:
            score -= 8
        territory, street_words, _street_numbers, house_base, house_mods = parse_address_components(text)
        if looks_like_address_text(text) or explicit_settlement_hint(text):
            score += 8
        if street_words or territory:
            score += 6
        if house_base:
            score += 10
        if house_mods:
            score += 4
        if _looks_like_contaminated_table_address(text):
            score -= 20
        candidates.append((score, len(text), text))
    if not candidates:
        return None
    return max(candidates, key=lambda item: (item[0], item[1]))[2]


def _pick_address_like_payload_value(payload: dict[str, str | None]) -> str | None:
    candidates: list[tuple[int, int, str]] = []
    for header, value in payload.items():
        text = _clean(value)
        if not text or _placeholder_text(text):
            continue
        header_key = _header_norm(header)
        if any(marker in header_key for marker in ("планируемый срок", "примеч", "количество", "потребность", "мощность")):
            continue
        score = 0
        if _header_matches(header_key, ADDRESS_HEADER_MARKERS):
            score += 10
        if _plausible_address_text(text):
            score += 7
        if explicit_settlement_hint(text):
            score += 3
        if score:
            candidates.append((score, len(text), text))
    if not candidates:
        return None
    return max(candidates, key=lambda item: (item[0], item[1]))[2]


def _complete_address_from_payload_slots(address: str | None, slots: dict[str, str]) -> str | None:
    text = _clean(address)
    if not text:
        return None
    _territory, street_words, _street_numbers, house_base, _house_mods = parse_address_components(text)
    if not street_words and slots.get("Street"):
        _append = slots["Street"]
        text = f"{text.rstrip(' ,')}, {_append}"
        _territory, street_words, _street_numbers, house_base, _house_mods = parse_address_components(text)
    if not house_base and slots.get("House"):
        text = f"{text.rstrip(' ,')}, {slots['House']}"
    return text


def _add_payload_context_to_address(address: str | None, slots: dict[str, str], context_lines: list[str]) -> str | None:
    text = _clean(address)
    if not text:
        return None
    parts: list[str] = []
    for value in context_lines[-3:]:
        _append_unique(parts, value)
    for slot in ("Postal_Index", "Subject", "Municipality", "Locality", "Territory", "Microdistrict"):
        value = slots.get(slot)
        if value and _norm(value) not in _norm(text):
            _append_unique(parts, value)
    _append_unique(parts, text)
    return ", ".join(parts) or text


def _address_from_slots(slots: dict[str, str], *, context_lines: list[str]) -> str | None:
    parts: list[str] = []
    for value in context_lines[-3:]:
        _append_unique(parts, value)
    for slot in ("Postal_Index", "Subject", "Municipality", "Locality", "Territory", "Microdistrict", "Street", "House"):
        _append_unique(parts, slots.get(slot))
    return ", ".join(parts) or None


def _observation_header_slot(value: Any) -> str | None:
    key = _header_norm(value)
    if not key:
        return None
    if _header_matches(key, ADDRESS_HEADER_MARKERS) and key not in {"улица", "street", "дом", "номер дома", "house", "house number"}:
        return "Address"
    if _header_matches(key, STREET_HEADER_MARKERS):
        return "Street"
    if _header_matches(key, HOUSE_HEADER_MARKERS):
        return "House"
    if _header_matches(key, POSTAL_HEADER_MARKERS):
        return "Postal_Index"
    if _header_matches(key, LOCALITY_HEADER_MARKERS):
        return "Locality"
    if _header_matches(key, MUNICIPALITY_HEADER_MARKERS):
        return "Municipality"
    if _header_matches(key, SUBJECT_HEADER_MARKERS):
        return "Subject"
    if _header_matches(key, MICRODISTRICT_HEADER_MARKERS):
        return "Microdistrict"
    if _header_matches(key, TERRITORY_HEADER_MARKERS):
        return "Territory"
    if _header_matches(key, OKTMO_HEADER_MARKERS):
        return "OKTMO_Code" if "код" in key else "OKTMO_Name"
    return None


def _header_matches(header_key: str, markers: tuple[str, ...]) -> bool:
    return any(_header_norm(marker) in header_key or header_key in _header_norm(marker) for marker in markers)


def _strong_plain_address_text(value: Any) -> bool:
    text = _clean(value)
    if not text or len(text) > 1200:
        return False
    if _looks_like_contaminated_table_address(text):
        return False
    territory, street_words, _street_numbers, house_base, _house_mods = parse_address_components(text)
    return bool(house_base and (street_words or territory or explicit_settlement_hint(text)))


def _parse_observation_kv_payload(text: str) -> dict[str, str] | None:
    if ":" not in text:
        return None
    parts = re.split(r"\s+\|\s+|\n+", str(text))
    payload: dict[str, str] = {}
    for part in parts:
        if ":" not in part:
            continue
        key, _sep, value = part.partition(":")
        key = key.strip()
        value = value.strip()
        if key and value and len(key) <= 120:
            payload[key] = value
    if len(payload) < 2:
        return None
    joined_keys = _header_norm(" | ".join(payload.keys()))
    if not any(marker in joined_keys for marker in ("адрес", "местоположение", "местонахождение", "октмо", "улица", "дом")):
        return None
    return payload


def _parse_observation_table_line(line: str, *, suffix: str) -> list[str | None]:
    text = line.strip().lstrip("\ufeff")
    if not text:
        return []
    if text.startswith("|") and "|" in text[1:]:
        cells = [cell.strip() for cell in text.strip("|").split("|")]
    elif "\t" in text:
        cells = next(csv.reader([text], delimiter="\t"))
    elif text.count(";") >= 2:
        cells = next(csv.reader([text], delimiter=";"))
    elif suffix == ".csv" and text.count(",") >= 2:
        cells = next(csv.reader([text], delimiter=","))
    else:
        return []
    if _is_markdown_separator_cells(cells):
        return []
    return [_clean_excel_value(cell) for cell in cells]


def _is_markdown_separator_cells(cells: list[str | None]) -> bool:
    meaningful = [re.sub(r"\s+", "", str(cell or "")) for cell in cells if str(cell or "").strip()]
    return bool(meaningful) and all(re.fullmatch(r":?-{3,}:?", cell) or re.fullmatch(r":?={3,}:?", cell) for cell in meaningful)


def _is_repeated_observation_header(headers: list[str], cells: list[str | None]) -> bool:
    candidate = _make_observation_headers(cells)
    matches = sum(1 for left, right in zip(headers, candidate) if _header_norm(left) == _header_norm(right))
    return matches >= max(3, min(len(headers), len(candidate)) // 2)


def _strip_non_address_tail(value: str | None) -> str | None:
    text = _clean(value)
    if not text:
        return None
    text = re.sub(r"\s*\(?\b\d{2}:\d{2}:\d{1,7}:\d+\b\)?", "", text)
    text = re.split(r"\s+\b(?:тел\.?|телефон|e[\s-]*mail|email|http|https|www)\b\s*[:.]?", text, maxsplit=1, flags=re.IGNORECASE)[0]
    return re.sub(r"\s+", " ", text).strip(" ,;") or None


def _read_source_text(path: Path) -> str:
    raw = path.read_bytes()
    if raw.startswith((b"\xff\xfe", b"\xfe\xff")):
        return raw.decode("utf-16")
    if b"\x00" in raw[:200]:
        try:
            return raw.decode("utf-16-le")
        except UnicodeDecodeError:
            pass
    for encoding in ("utf-8-sig", "utf-8", "cp1251", "cp866"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


def _source_lines(text: str) -> list[str]:
    return str(text or "").replace("\r\n", "\n").replace("\r", "\n").split("\n")


def _split_observation_blocks(text: str) -> list[str]:
    blocks: list[str] = []
    for raw_block in re.split(r"\n\s*\n+", str(text or "").replace("\r\n", "\n").replace("\r", "\n")):
        lines = [re.sub(r"\s+", " ", line).strip() for line in raw_block.split("\n") if line.strip()]
        if not lines:
            continue
        if len(lines) > 1 and sum(1 for line in lines if ":" in line) >= 2:
            blocks.append(" | ".join(lines))
        else:
            blocks.extend(lines)
    return blocks


def _docx_node_text(node: ET.Element) -> str:
    fragments: list[str] = []
    for child in node.iter():
        name = _local_xml_name(child.tag)
        if name == "t":
            fragments.append(child.text or "")
        elif name in {"tab", "br", "cr"}:
            fragments.append(" ")
    return re.sub(r"\s+", " ", "".join(fragments)).strip()


def _odt_row_values(row: ET.Element) -> list[str | None]:
    cells: list[str | None] = []
    for cell in row.findall("./table:table-cell", ODT_NS):
        value = _clean_xml_text(_xml_node_text(cell))
        try:
            repeat = max(1, min(int(cell.attrib.get(ODT_REPEAT_ATTR, "1") or "1"), 80))
        except ValueError:
            repeat = 1
        cells.extend([value] * repeat)
    return cells


def _xml_node_text(node: ET.Element) -> str:
    fragments: list[str] = []
    for child in node.iter():
        if child.text:
            fragments.append(child.text)
        if _local_xml_name(child.tag) in {"tab", "line-break"}:
            fragments.append(" ")
        if child.tail:
            fragments.append(child.tail)
    return re.sub(r"\s+", " ", "".join(fragments)).strip()


def _local_xml_name(tag: str) -> str:
    return str(tag).rsplit("}", 1)[-1]


def _clean_xml_text(value: Any) -> str | None:
    return _clean_excel_value(value)


def _clean_excel_value(value: Any) -> str | None:
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = re.sub(r"\s+", " ", str(value).replace("\xa0", " ")).strip()
    if not text or text.lower() == "nan":
        return None
    return text


def _header_norm(value: Any) -> str:
    text = str(value or "").lower().replace("ё", "е")
    text = re.sub(r"[^0-9a-zа-я]+", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def _placeholder_text(value: Any) -> bool:
    text = _header_norm(value)
    return text in {"", "x", "х", "-", "нет", "н д", "н/д"}


def _extract_structured_source_candidates(
    file_path: Path,
    *,
    source_columns: str | None,
) -> tuple[list[SourceCandidate], list[str]]:
    try:
        if file_path.suffix.lower() == ".xlsx":
            return _extract_structured_xlsx_candidates(file_path, source_columns=source_columns), []
    except Exception as exc:
        return [], [f"{file_path.name}: structured context assembly failed: {exc}"]
    return [], []


def _extract_structured_xlsx_candidates(file_path: Path, *, source_columns: str | None) -> list[SourceCandidate]:
    workbook = pd.read_excel(file_path, sheet_name=None, header=None, dtype=object)
    selected = _source_column_filter(source_columns)
    result: list[SourceCandidate] = []
    for sheet_name, frame in workbook.items():
        if frame.empty:
            continue
        context_lines: list[str] = []
        header_state = _StructuredHeaderState()
        for row_index in range(len(frame)):
            raw_values = [_clean(frame.iat[row_index, col_index]) for col_index in range(len(frame.columns))]
            selected_raw_values = _values_for_selected_columns(raw_values, selected)
            values = [value for value in selected_raw_values if value]
            if not values:
                continue
            if _update_structured_headers(header_state, selected_raw_values):
                continue
            row_text = _clean(" | ".join(values)) or ""
            if _looks_like_admin_context(row_text) and not _has_street_or_house_text(row_text):
                _append_context_line(context_lines, row_text)
                continue
            assembled = _structured_row_address(header_state, raw_values, context_lines, selected)
            if not assembled:
                continue
            result.append(
                SourceCandidate(
                    value=assembled,
                    source_file=file_path,
                    source_kind="xlsx_structured_row",
                    locator=f"{sheet_name}!R{row_index + 1}",
                )
            )
    return result


class _StructuredHeaderState:
    def __init__(self) -> None:
        self.headers: dict[int, str] = {}
        self.header_row: int | None = None


def _update_structured_headers(state: _StructuredHeaderState, values: list[str | None]) -> bool:
    labels: dict[int, str] = {}
    hits = 0
    for index, value in enumerate(values):
        text = _clean(value)
        if not text:
            continue
        slot = _structured_header_slot(text)
        if slot:
            hits += 1
            labels[index] = slot
    if hits >= 2:
        state.headers = labels
        return True
    return False


def _structured_header_slot(value: Any) -> str | None:
    key = _norm(value)
    compact = re.sub(r"\s+", "", key)
    if compact in {"улица", "наименованиеулицы", "адресулица", "street"} or "улица" in key:
        return "Street"
    if compact in {"дом", "номердома", "домномер", "house", "housenumber"} or "номер дома" in key:
        return "House"
    if compact in {"индекс", "почтовыииндекс", "postalindex"}:
        return "Postal_Index"
    if compact in {"город", "населенныипункт", "населенныйпункт", "locality"}:
        return "Locality"
    if "муницип" in key or "городской округ" in key:
        return "Municipality"
    if compact in {"регион", "субъект", "субъектрф"} or "область" in key:
        return "Subject"
    if "микрорайон" in key or compact in {"мкр", "квартал"}:
        return "Microdistrict"
    if "территория" in key or compact in {"территория", "тер"}:
        return "Territory"
    if "октмо" in key:
        return "OKTMO_Code" if "код" in key else "OKTMO_Name"
    return None


def _looks_like_structured_header_row(values: list[str]) -> bool:
    return sum(1 for value in values if _structured_header_slot(value)) >= 2


def _structured_row_address(
    state: _StructuredHeaderState,
    values: list[str | None],
    context_lines: list[str],
    selected_columns: set[int] | None,
) -> str | None:
    if not state.headers:
        return None
    slot_values: OrderedDict[str, str] = OrderedDict()
    fallback_values: list[str] = []
    for index, value in enumerate(values):
        text = _clean(value)
        if not text:
            continue
        if selected_columns is not None and index not in selected_columns:
            continue
        slot = state.headers.get(index)
        if slot:
            slot_values[slot] = _format_structured_slot(slot, text)
        elif _plausible_address_text(text) or _looks_like_admin_context(text):
            _append_unique(fallback_values, text)

    street = _clean(slot_values.get("Street"))
    house = _clean(slot_values.get("House"))
    if not street and not fallback_values:
        return None
    if street and not house:
        return None
    if street and house:
        parts: list[str] = []
        for value in context_lines[-3:]:
            _append_unique(parts, value)
        for slot in ("Postal_Index", "Subject", "Municipality", "Locality", "Territory", "Microdistrict", "Street", "House"):
            _append_unique(parts, slot_values.get(slot))
        return ", ".join(parts) or None
    parts = []
    for value in [*context_lines[-3:], *fallback_values]:
        _append_unique(parts, value)
    return ", ".join(parts) or None


def _format_structured_slot(slot: str, value: Any) -> str | None:
    text = _clean(value)
    if not text:
        return None
    if slot == "Street" and not STREET_MARKER_RE.search(text):
        return f"ул. {text}"
    if slot == "House":
        if _has_house_text(text):
            return text
        if _looks_like_standalone_house_cell(text):
            return f"д. {text}"
    return text


def _source_column_filter(value: str | None) -> set[int] | None:
    tokens = [part.strip() for part in str(value or "").split(",") if part.strip()]
    if not tokens:
        return None
    result: set[int] = set()
    for token in tokens:
        if re.fullmatch(r"[A-Za-z]+", token):
            result.add(excel_column_index(token))
        elif token.isdigit():
            result.add(max(0, int(token) - 1))
    return result or None


def _values_for_selected_columns(values: list[str | None], selected_columns: set[int] | None) -> list[str | None]:
    if selected_columns is None:
        return values
    return [value if index in selected_columns else None for index, value in enumerate(values)]


def _contextual_xlsx_row_candidates(candidates: list[SourceCandidate]) -> list[SourceCandidate]:
    rows: OrderedDict[tuple[str, int], list[tuple[int, SourceCandidate]]] = OrderedDict()
    for candidate in candidates:
        if candidate.source_kind != "xlsx_cell":
            continue
        match = re.match(r"^(.+)!([A-Z]+)(\d+)$", candidate.locator)
        if not match:
            continue
        sheet_name, column_letter, row_number = match.group(1), match.group(2), int(match.group(3))
        rows.setdefault((sheet_name, row_number), []).append((excel_column_index(column_letter), candidate))

    contextual: list[SourceCandidate] = []
    context_lines: list[str] = []
    for (sheet_name, row_number), cells in rows.items():
        ordered = [item for _column, item in sorted(cells, key=lambda pair: pair[0])]
        values = [_clean(item.value) for item in ordered]
        values = [value for value in values if value]
        if not values:
            continue
        if _looks_like_structured_header_row(values):
            continue
        row_text = _clean(" | ".join(values)) or ""
        if _looks_like_admin_context(row_text) and not _has_street_or_house_text(row_text):
            _append_context_line(context_lines, row_text)
            continue
        if not _row_has_address_fragment(values):
            continue
        combined = _contextual_address_value(context_lines, values)
        if not combined:
            continue
        contextual.append(
            SourceCandidate(
                value=combined,
                source_file=ordered[0].source_file,
                source_kind="xlsx_context_row",
                locator=f"{sheet_name}!R{row_number}",
            )
        )
    return contextual


def _contextual_sequential_candidates(candidates: list[SourceCandidate]) -> list[SourceCandidate]:
    contextual: list[SourceCandidate] = []
    context_lines: list[str] = []
    for candidate in candidates:
        if candidate.source_kind == "xlsx_cell":
            continue
        text = _clean(candidate.value)
        if not text:
            continue
        if _looks_like_admin_context(text) and not _has_street_or_house_text(text):
            _append_context_line(context_lines, text)
            continue
        if not context_lines or not _has_street_or_house_text(text):
            continue
        combined = _contextual_address_value(context_lines, [text])
        if not combined or _norm(combined) == _norm(text):
            continue
        contextual.append(
            SourceCandidate(
                value=combined,
                source_file=candidate.source_file,
                source_kind=f"{candidate.source_kind}_with_context",
                locator=f"{candidate.locator}+context",
            )
        )
    return contextual


def _deduplicate_source_candidates(candidates: list[SourceCandidate]) -> list[SourceCandidate]:
    seen: set[tuple[str, str, str]] = set()
    result: list[SourceCandidate] = []
    for candidate in candidates:
        key = (_norm(candidate.value), candidate.source_kind, candidate.locator)
        if not key[0] or key in seen:
            continue
        seen.add(key)
        result.append(candidate)
    return result


def _append_context_line(context_lines: list[str], value: str) -> None:
    text = _clean(value)
    if not text or len(text) > 350:
        return
    if _looks_like_contaminated_table_address(text):
        return
    key = _norm(text)
    if any(_norm(existing) == key for existing in context_lines):
        return
    context_lines.append(text)
    del context_lines[:-4]


def _contextual_address_value(context_lines: list[str], values: list[str]) -> str | None:
    address_parts = _addressish_row_values(values)
    if not address_parts:
        return None
    parts: list[str] = []
    for value in [*context_lines[-3:], *address_parts]:
        _append_unique(parts, value)
    if len(parts) <= len(address_parts):
        return ", ".join(parts) or None
    return ", ".join(parts)


def _addressish_row_values(values: list[str]) -> list[str]:
    has_street = any(_has_street_text(value) for value in values)
    has_house = any(_has_house_text(value) or _looks_like_standalone_house_cell(value) for value in values)
    result: list[str] = []
    for value in values:
        text = _clean(value)
        if not text:
            continue
        if _looks_like_contaminated_table_address(text):
            continue
        if (
            _plausible_address_text(text)
            or _looks_like_admin_context(text)
            or _has_street_text(text)
            or _has_house_text(text)
            or (has_street and _looks_like_standalone_house_cell(text))
            or _valid_postal_index(text)
        ):
            _append_unique(result, text)
    if not result and has_street and has_house:
        result = [_clean(value) for value in values if _clean(value)][:12]  # type: ignore[list-item]
    return [value for value in result if value]


def _row_has_address_fragment(values: list[str]) -> bool:
    return bool(_addressish_row_values(values))


def _looks_like_admin_context(value: Any) -> bool:
    text = _clean(value)
    if not text:
        return False
    if len(text) > 500 or _looks_like_contaminated_table_address(text):
        return False
    return bool(ADMIN_CONTEXT_RE.search(text) or re.search(r"\b\d{11}\b", text))


def _has_street_or_house_text(value: Any) -> bool:
    return _has_street_text(value) or _has_house_text(value)


def _has_street_text(value: Any) -> bool:
    text = str(value or "")
    if STREET_MARKER_RE.search(text):
        return True
    _territory, street_words, _street_numbers, _house_base, _house_mods = parse_address_components(text)
    return bool(street_words)


def _has_house_text(value: Any) -> bool:
    text = str(value or "")
    if HOUSE_MARKER_RE.search(text):
        return True
    _territory, _street_words, _street_numbers, house_base, _house_mods = parse_address_components(text)
    return bool(house_base)


def _looks_like_standalone_house_cell(value: Any) -> bool:
    text = _clean(value) or ""
    return bool(re.fullmatch(r"\d{1,4}[а-яa-z]?(?:\s*(?:[/\\-]|к\.?|корп\.?|стр\.?|с\.?)\s*\d+[а-яa-z]?)?", text, flags=re.IGNORECASE))


def _group_candidate_rows(
    rows: list[_CandidateRow],
    *,
    clean_duplicates: bool,
    global_context: dict[str, Any] | None = None,
) -> list[_GroupedRows]:
    if not clean_duplicates:
        return [_GroupedRows(f"raw:{index:08d}", (row,)) for index, row in enumerate(rows)]
    groups: OrderedDict[str, list[_CandidateRow]] = OrderedDict()
    for row in sorted(rows, key=lambda item: (-item.quality_score, str(item.candidate.source_file), item.candidate.locator)):
        key = _assembly_key(row, global_context=global_context)
        if not key:
            key = f"raw:{len(groups):08d}"
        groups.setdefault(key, []).append(row)
    return [_GroupedRows(key, tuple(group_rows)) for key, group_rows in groups.items()]


def _final_rows_from_groups(
    groups: list[_GroupedRows],
    *,
    enabled_slots: tuple[str, ...],
    slot_order: tuple[str, ...],
    omit_microdistrict_when_street_found: bool,
    include_slots: bool,
    sanitize_incomplete: bool,
    global_context: dict[str, Any] | None = None,
) -> tuple[list[dict[str, Any]], list[dict[str, Any]], int, list[dict[str, Any]]]:
    normalized_rows: list[dict[str, Any]] = []
    rejected_rows: list[dict[str, Any]] = []
    staged_rows: list[dict[str, Any]] = []
    seen_output_keys: dict[str, dict[str, Any]] = {}
    duplicates = 0
    for group in groups:
        duplicates += max(0, len(group.rows) - 1)
        if any(item.source_kind == "existing_reference" for item in group.rows):
            row = _merge_group_rows(
                group,
                enabled_slots=enabled_slots,
                slot_order=slot_order,
                omit_microdistrict_when_street_found=omit_microdistrict_when_street_found,
                global_context=global_context,
            )
            staged_rows.append(_stage_row(group, row, status="existing_reference", reason="duplicate_guard"))
            continue
        row = _merge_group_rows(
            group,
            enabled_slots=enabled_slots,
            slot_order=slot_order,
            omit_microdistrict_when_street_found=omit_microdistrict_when_street_found,
            global_context=global_context,
        )
        reason = _drop_reason(row) if sanitize_incomplete else None
        if reason:
            rejected = dict(row)
            rejected["Drop_Reason"] = reason
            if reason == "oktmo_scope_conflict":
                rejected["Disposition"] = "held_out_of_clean_stack"
                rejected["Review_Action"] = "manual_scope_review"
            rejected_rows.append(_project_diagnostic_columns(rejected))
            staged_rows.append(_stage_row(group, rejected, status="rejected", reason=reason))
            continue
        normalized = _project_reference_columns(
            row,
            include_slots=include_slots,
            enabled_slots=enabled_slots,
            slot_order=slot_order,
            clean_old_slots=True,
        )
        output_keys = _final_output_dedupe_keys(normalized)
        existing_output = next((seen_output_keys[key] for key in output_keys if key in seen_output_keys), None)
        if existing_output is not None:
            duplicates += 1
            _merge_projected_duplicate_row(existing_output, normalized)
            staged_rows.append(_stage_row(group, row, status="duplicate", reason="final_output_duplicate"))
            continue
        for output_key in output_keys:
            seen_output_keys[output_key] = normalized
        normalized_rows.append(normalized)
        staged_rows.append(_stage_row(group, row, status="ready", reason=None))
    normalized_rows.sort(key=lambda item: _sort_key(item, enabled_slots, slot_order))
    rejected_rows.sort(key=lambda item: str(item.get("Raw_Address") or item.get("Normalized_Address") or "").casefold())
    staged_rows.sort(key=lambda item: (_natural_key(item.get("Proposed_Address")), str(item.get("Assembly_Status") or "")))
    return normalized_rows, rejected_rows, duplicates, staged_rows


def _final_output_dedupe_keys(row: dict[str, Any]) -> tuple[str, ...]:
    keys: list[str] = []
    core = _final_output_core_key(row)
    if core:
        scope = _final_output_scope_key(row)
        keys.append(f"core|{scope}|{core}" if scope else f"core|{core}")
    address = _norm(row.get("Normalized_Address"))
    if address:
        keys.append(f"address|{address}")
    return tuple(dict.fromkeys(keys))


def _final_output_core_key(row: dict[str, Any]) -> str:
    house_text = _normalize_house_slot(row.get("House"))
    if not house_text:
        return ""
    _house_territory, _house_street_words, _house_street_numbers, house_base, parsed_house_mods = parse_address_components(house_text)
    if not house_base:
        return ""
    locator = ""
    street_text = _normalize_street_slot(row.get("Street"))
    if street_text:
        _territory, street_words, street_numbers, _street_house_base, _street_house_mods = parse_address_components(street_text)
        if street_words:
            locator = "|".join(part for part in ("street", street_words, ",".join(sorted(street_numbers))) if part)
        else:
            locator = f"street-text|{_norm(street_text)}"
    if not locator:
        microdistrict = _normalize_microdistrict_slot(row.get("Microdistrict"))
        if microdistrict:
            locator = f"microdistrict|{_norm(microdistrict)}"
    if not locator:
        territory = _clean(row.get("Territory"))
        if territory:
            locator = f"territory|{_norm(territory)}"
    if not locator:
        return ""
    house_mods = parsed_house_mods or _safe_house_mods(row.get("House_Mods") or row.get("House_Modifiers"))
    return "|".join(
        part
        for part in (
            locator,
            _house_kind_key(house_text),
            house_base,
            house_mods,
            _norm(row.get("Premise")),
        )
        if part
    )


def _final_output_scope_key(row: dict[str, Any]) -> str:
    digits = re.sub(r"\D+", "", str(row.get("OKTMO_Code") or ""))
    if digits:
        return f"oktmo:{digits}"
    scope = _norm(" | ".join(str(row.get(column) or "") for column in ("Subject", "Municipality", "Locality")))
    return f"context:{scope}" if scope else ""


def _conflict_feedback_rows_from_rejected(rejected_rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    result: list[dict[str, Any]] = []
    seen: set[str] = set()
    for row in rejected_rows:
        if row.get("Drop_Reason") != "oktmo_scope_conflict":
            continue
        raw = _clean(row.get("Raw_Address"))
        if not raw:
            continue
        key = _norm(raw)
        if key in seen:
            continue
        seen.add(key)
        feedback = {
            "Raw_Address": raw,
            "Drop_Reason": row.get("Drop_Reason"),
            "Disposition": "held_out_of_clean_stack",
            "Review_Action": "manual_scope_review",
            "Feedback_Mode": "raw_conflict",
            "Feedback_Note": "OKTMO scope conflict; original text is kept for review outside the clean address stack.",
            "Rejected_Normalized_Address": row.get("Normalized_Address"),
            "Rejected_OKTMO_Code": row.get("OKTMO_Code"),
            "Rejected_Municipality": row.get("Municipality"),
            "Rejected_Locality": row.get("Locality"),
            "Rejected_Street": row.get("Street"),
            "Rejected_House": row.get("House"),
        }
        for column in (
            "Sources",
            "First_Source",
            "Candidate_Count",
            "Source_Count",
            "OKTMO_Scope_Match",
            "OKTMO_Scope_Code",
            "OKTMO_Scope_Source",
            "OKTMO_Scope_Weight",
            "OKTMO_Scope_Valid",
            "OKTMO_Scope_Context",
        ):
            if column in row:
                feedback[column] = row.get(column)
        result.append(feedback)
    return result


def _merge_projected_duplicate_row(target: dict[str, Any], duplicate: dict[str, Any]) -> None:
    target["Source_Count"] = int(target.get("Source_Count") or 1) + int(duplicate.get("Source_Count") or 1)
    target["Raw_Examples"] = _append_unique_examples(
        target.get("Raw_Examples"),
        duplicate.get("Raw_Examples") or duplicate.get("Raw_Address"),
    )
    target["Sources"] = _append_unique_examples(
        target.get("Sources"),
        duplicate.get("Sources") or duplicate.get("First_Source"),
        limit=30,
        separator="; ",
    )


def _merge_group_rows(
    group: _GroupedRows,
    *,
    enabled_slots: tuple[str, ...],
    slot_order: tuple[str, ...],
    omit_microdistrict_when_street_found: bool,
    global_context: dict[str, Any] | None = None,
) -> dict[str, Any]:
    best = max(group.rows, key=lambda item: item.quality_score)
    merged = dict(best.row)
    merged["Candidate_Count"] = len(group.rows)
    merged["Best_Slot_Score"] = round(best.quality_score, 3)

    slot_columns = list(dict.fromkeys([*_ordered_slot_columns(DEFAULT_ENABLED_SLOTS, DEFAULT_SLOT_ORDER), "Match_Key"]))
    for column in slot_columns:
        if column == "Match_Key":
            continue
        value = _choose_slot_value(column, group.rows)
        if value not in (None, ""):
            merged[column] = value
        elif column in merged:
            merged[column] = None

    if global_context:
        for column in ("Federal_District", "Parent_Subject", "Subject", "Municipality", "Locality", "OKTMO_Code", "OKTMO_Name"):
            if not _clean(merged.get(column)) and _clean(global_context.get(column)):
                merged[column] = global_context.get(column)

    _clean_invalid_slot_values(merged)
    merged["Source_Count"] = sum(int(item.row.get("Source_Count") or 1) for item in group.rows)
    merged["Raw_Examples"] = _merge_text_examples(item.row.get("Raw_Address") for item in group.rows)
    merged["Sources"] = _merge_text_examples((item.row.get("Sources") for item in group.rows), limit=30, separator="; ")
    merged["First_Source"] = str(best.candidate.source_file)
    merged["Normalized_Address"] = (
        _construct_address_from_row(
            merged,
            enabled_slots,
            slot_order,
            omit_microdistrict_when_street_found=omit_microdistrict_when_street_found,
        )
        or merged.get("Normalized_Address")
    )
    merged["Match_Key"] = group.key
    return merged


def _dominant_admin_context(rows: list[_CandidateRow]) -> dict[str, Any]:
    context: dict[str, Any] = {}
    for column in ("Federal_District", "Parent_Subject", "Subject", "Municipality", "Locality", "OKTMO_Code", "OKTMO_Name"):
        votes: dict[str, dict[str, Any]] = {}
        for item in rows:
            value = _normalize_slot_candidate_value(column, item.row.get(column))
            if not value:
                continue
            key = _slot_vote_key(column, value)
            if not key:
                continue
            vote = votes.setdefault(key, {"value": value, "count": 0, "score": 0.0})
            vote["count"] += 1
            vote["score"] += item.quality_score + _slot_value_bonus(column, value)
        if not votes:
            continue
        ranked = sorted(votes.values(), key=lambda item: (item["count"], item["score"], len(str(item["value"]))), reverse=True)
        best = ranked[0]
        second_count = int(ranked[1]["count"]) if len(ranked) > 1 else 0
        total = sum(int(item["count"]) for item in ranked)
        if len(ranked) == 1 or (
            int(best["count"]) >= 2
            and int(best["count"]) >= max(2, second_count * 2)
            and int(best["count"]) / max(1, total) >= 0.45
        ):
            context[column] = best["value"]
    return context


def _stage_row(group: _GroupedRows, row: dict[str, Any], *, status: str, reason: str | None) -> dict[str, Any]:
    staged: OrderedDict[str, Any] = OrderedDict()
    staged["Assembly_Status"] = status
    staged["Drop_Reason"] = reason
    staged["Proposed_Address"] = row.get("Normalized_Address")
    staged["Match_Key"] = row.get("Match_Key") or group.key
    staged["Candidate_Count"] = row.get("Candidate_Count") or len(group.rows)
    staged["Source_Count"] = row.get("Source_Count")
    staged["Best_Slot_Score"] = row.get("Best_Slot_Score")
    staged["OKTMO_Scope_Match"] = row.get("OKTMO_Scope_Match")
    staged["OKTMO_Scope_Code"] = row.get("OKTMO_Scope_Code")
    staged["OKTMO_Scope_Source"] = row.get("OKTMO_Scope_Source")
    staged["OKTMO_Scope_Weight"] = row.get("OKTMO_Scope_Weight")
    staged["OKTMO_Scope_Valid"] = row.get("OKTMO_Scope_Valid")
    staged["OKTMO_Scope_Context"] = row.get("OKTMO_Scope_Context")
    staged["First_Source"] = row.get("First_Source")
    staged["Sources"] = row.get("Sources")
    staged["Raw_Examples"] = row.get("Raw_Examples")
    for column in STAGING_CANDIDATE_COLUMNS:
        if column in row:
            staged[column] = row.get(column)
    for column in STAGING_CANDIDATE_COLUMNS:
        summary = _slot_candidate_summary(column, group.rows)
        if summary:
            staged[f"{column}_Candidates"] = summary
    return dict(staged)


def _slot_candidate_summary(column: str, rows: tuple[_CandidateRow, ...], *, limit: int = 8) -> str:
    votes: dict[str, dict[str, Any]] = {}
    for item in rows:
        value = _normalize_slot_candidate_value(column, item.row.get(column), allow_dirty=True)
        if not value:
            continue
        key = _slot_vote_key(column, value)
        if not key:
            continue
        vote = votes.setdefault(key, {"value": value, "count": 0, "score": 0.0})
        vote["count"] += 1
        vote["score"] += item.quality_score + _slot_value_bonus(column, value) - _slot_value_penalty(column, value)
    if not votes:
        return ""
    ranked = sorted(votes.values(), key=lambda item: (item["count"], item["score"], len(str(item["value"]))), reverse=True)
    parts = [f"{item['value']} [{item['count']}]" for item in ranked[:limit]]
    if len(ranked) > limit:
        parts.append(f"+{len(ranked) - limit}")
    return " | ".join(parts)


def _choose_slot_value(column: str, rows: tuple[_CandidateRow, ...]) -> Any:
    votes: dict[str, dict[str, Any]] = {}
    allow_dirty = column in {"Street", "House"}
    for item in rows:
        value = item.row.get(column)
        text = _normalize_slot_candidate_value(column, value, allow_dirty=allow_dirty)
        if not text:
            continue
        key = _slot_vote_key(column, text)
        if not key:
            continue
        vote = votes.setdefault(key, {"value": text, "count": 0, "score": 0.0, "best_len": 0})
        vote["count"] += 1
        value_score = item.quality_score + _slot_value_bonus(column, text) - _slot_value_penalty(column, text)
        if value_score > vote["score"] or (value_score == vote["score"] and len(text) > vote["best_len"]):
            vote["value"] = text
            vote["score"] = value_score
            vote["best_len"] = len(text)
    if not votes:
        return None
    best_vote = max(votes.values(), key=lambda vote: (vote["count"], vote["score"], vote["best_len"]))
    return best_vote["value"]


def _slot_value_bonus(column: str, value: str) -> float:
    text = _norm(value)
    bonus = 0.0
    if column in {"Subject", "Municipality", "Locality", "OKTMO_Code", "OKTMO_Name", "Federal_District"}:
        bonus += 8.0
    if column == "Postal_Index" and _valid_postal_index(value):
        bonus += 8.0
    if column == "House" and re.search(r"\bд\.?\s*\d", text):
        bonus += 6.0
    if column == "Street" and ADDRESS_HINT_RE.search(value):
        bonus += 5.0
    return bonus


def _slot_value_penalty(column: str, value: str) -> float:
    if column == "Street" and _strict_street_slot_drop_reason(value):
        return 80.0
    if column == "House" and not _looks_like_ideal_house_slot(value):
        return 60.0
    if column in {"Municipality", "Locality", "Microdistrict"} and _strict_geo_slot_drop_reason(column, value):
        return 70.0
    if column == "Territory" and (_contains_org_context(value) or _looks_like_route_fragment(value) or _has_unbalanced_brackets(value)):
        return 70.0
    return 0.0


def _normalize_slot_candidate_value(column: str, value: Any, *, allow_dirty: bool = False) -> str | None:
    text = _clean(value)
    if not text:
        return None
    if column == "Postal_Index":
        return text if _valid_postal_index(text) else None
    if column == "Street":
        normalized = _normalize_street_slot(text)
        if not normalized:
            return None
        if not allow_dirty and _strict_street_slot_drop_reason(normalized):
            return None
        return normalized
    if column == "House":
        normalized = _normalize_house_slot(text)
        if not normalized:
            return None
        if not allow_dirty and not _looks_like_ideal_house_slot(normalized):
            return None
        return normalized
    if column == "Municipality":
        normalized = _normalize_municipality_slot(text)
        if not normalized:
            return None
        if not allow_dirty and _strict_geo_slot_drop_reason(column, normalized):
            return None
        return normalized
    if column == "Locality":
        normalized = _normalize_locality_slot(text)
        if not normalized:
            return None
        if not allow_dirty and _strict_geo_slot_drop_reason(column, normalized):
            return None
        return normalized
    if column == "Microdistrict":
        normalized = _normalize_microdistrict_slot(text)
        if not normalized:
            return None
        if not allow_dirty and _strict_geo_slot_drop_reason(column, normalized):
            return None
        return normalized
    if column == "Territory":
        if not allow_dirty and (
            _contains_org_context(text) or _looks_like_route_fragment(text) or _has_unbalanced_brackets(text)
        ):
            return None
        return text
    return text


def _slot_vote_key(column: str, value: Any) -> str:
    text = _clean(value)
    if not text:
        return ""
    if column == "Street":
        _territory, street_words, street_numbers, _house_base, _house_mods = parse_address_components(text)
        return "|".join(part for part in ("street", street_words, ",".join(sorted(street_numbers))) if part)
    if column == "House":
        _territory, _street_words, _street_numbers, house_base, house_mods = parse_address_components(text)
        return "|".join(part for part in ("house", _house_kind_key(text), house_base, _safe_house_mods(house_mods)) if part)
    if column in {"Municipality", "Locality", "Territory", "Microdistrict", "Subject", "Parent_Subject", "Autonomous_Okrug"}:
        return f"{column.lower()}|{_geo_core_key(text) or _norm(text)}"
    return _norm(text)


def _clean_invalid_slot_values(row: dict[str, Any]) -> None:
    postal_index = _clean(row.get("Postal_Index"))
    if postal_index and not _valid_postal_index(postal_index):
        row["Postal_Index"] = None
    street = _clean(row.get("Street"))
    if street:
        row["Street"] = _normalize_street_slot(street)
    house = _clean(row.get("House"))
    if house:
        row["House"] = _normalize_house_slot(house)
    municipality = _clean(row.get("Municipality"))
    if municipality:
        row["Municipality"] = _normalize_municipality_slot(municipality)
    locality = _clean(row.get("Locality"))
    if locality:
        row["Locality"] = _normalize_locality_slot(locality)
    microdistrict = _clean(row.get("Microdistrict"))
    if microdistrict:
        row["Microdistrict"] = _normalize_microdistrict_slot(microdistrict)


def _valid_postal_index(value: Any) -> bool:
    text = str(value or "").strip()
    return bool(re.fullmatch(r"[1-9]\d{5}", text))


def _normalize_street_slot(value: Any) -> str | None:
    text = _clean(_normalize_spaced_address_marker_dots(value))
    if not text:
        return None
    text = re.sub(r"^[\s\"'«»(]+|[\s\"'«»);:]+$", "", text)
    text = re.sub(
        r"\b(ул|пер|пр|пр-т|пр-кт|просп|пр-д|б-р|наб|пл)\.\s*(?=[А-ЯЁA-Zа-яёa-z0-9])",
        lambda match: f"{match.group(1)}. ",
        text,
        flags=re.IGNORECASE,
    )
    bare_pr_as_proezd = _bare_pr_should_be_proezd(text)
    replacements = (
        (r"^ул\.?\s+", "ул. "),
        (r"^улица\s+", "ул. "),
        (r"^пер\.?\s+", "пер. "),
        (r"^переулок\s+", "пер. "),
        (r"^проспект\s+", "пр-кт "),
        (r"^просп\.?\s+", "пр-кт "),
        (r"^пр-т\.?\s+", "пр-кт "),
        (r"^пр\.\s+", "пр-д " if bare_pr_as_proezd else "пр-кт "),
        (r"^пр-кт\.?\s+", "пр-кт "),
        (r"^проезд\s+", "пр-д "),
        (r"^пр-д\.?\s+", "пр-д "),
        (r"^бульвар\s+", "б-р "),
        (r"^набережная\s+", "наб. "),
        (r"^площадь\s+", "пл. "),
    )
    for pattern, replacement in replacements:
        text = re.sub(pattern, replacement, text, flags=re.IGNORECASE)
    # "прт" is a mistyped "пр-т"; the avenue name after it is left as written.
    text = re.sub(r"(?i)\b(?:ул\.?|улица)\s+прт\s+(?=[А-ЯЁA-Zа-яёa-z])", "пр-кт ", text)
    text = _strip_trailing_locality_suffix_from_street(text)
    text = re.sub(r"(?i)\bсовсетов\b", "Советов", text)
    text = re.sub(r"(?i)\bпрт\s+(?=[А-ЯЁA-Zа-яёa-z])", "пр-кт ", text)
    text = _normalize_street_name_typos(text)
    text = _collapse_repeated_street_type_markers(text)
    text = _strip_residential_complex_tail_from_street(text)
    text = _normalize_street_range_dashes(text)
    text = re.sub(r"\bул\.?\s+", "ул. ", text, flags=re.IGNORECASE)
    text = re.sub(r"\bпер\.?\s+", "пер. ", text, flags=re.IGNORECASE)
    text = re.sub(r"\bпр\.\s+", "пр-д " if bare_pr_as_proezd else "пр-кт ", text, flags=re.IGNORECASE)
    text = re.sub(r"\bпр-кт\.?\s+", "пр-кт ", text, flags=re.IGNORECASE)
    text = re.sub(r"\bпр-д\.?\s+", "пр-д ", text, flags=re.IGNORECASE)
    text = re.sub(r"\s+", " ", text).strip(" ,")
    if _street_slot_has_explicit_locator(text):
        return text
    if _looks_like_plain_street_name(text):
        return f"ул. {text}"
    return text or None


def _bare_pr_should_be_proezd(text: str) -> bool:
    match = re.match(r"(?i)^\s*пр\.\s+(.+)$", str(text or "").strip())
    if not match:
        return False
    return _norm(match.group(1)) in BARE_PR_PROEZD_STREET_NAMES


def _normalize_street_name_typos(value: str) -> str:
    replacements = {
        "совесткая": "Советская",
        "советска": "Советская",
        "украинска": "Украинская",
    }
    text = value
    for source, replacement in replacements.items():
        text = re.sub(rf"(?i)\b{source}\b", replacement, text)
    return text


def _collapse_repeated_street_type_markers(value: str) -> str:
    text = str(value or "")
    marker_groups = (
        (r"(?:ул\.?|улица)", "ул. "),
        (r"(?:пер\.?|переулок)", "пер. "),
        (r"(?:пр-кт\.?|пр-т\.?|просп\.?|проспект)", "пр-кт "),
        (r"(?:пр-д\.?|проезд)", "пр-д "),
        (r"(?:б-р\.?|бульвар)", "б-р "),
        (r"(?:наб\.?|набережная)", "наб. "),
        (r"(?:пл\.?|площадь)", "пл. "),
    )
    for marker_pattern, replacement in marker_groups:
        pattern = rf"(?i)^\s*{marker_pattern}\s+{marker_pattern}\s+"
        previous = None
        while previous != text:
            previous = text
            text = re.sub(pattern, replacement, text)
    return re.sub(r"\.\s+\.", ".", text)


def _strip_residential_complex_tail_from_street(value: str) -> str:
    text = str(value or "").strip()
    if not re.match(
        r"(?i)^\s*(?:ул\.?|улица|пер\.?|переулок|проспект|пр-т\.?|пр-кт\.?|просп\.?|пр-д\.?|проезд|"
        r"б-р\.?|бульвар|наб\.?|набережная|пл\.?|площадь)\s+",
        text,
    ):
        return text
    stripped = re.sub(
        r"(?i)\s+жк\s+(?:[\"«][^\"»]+[\"»]?|[А-ЯЁA-Zа-яёa-z0-9 .'-]{2,60})\s*$",
        "",
        text,
    ).strip(" ,")
    if stripped and not _looks_like_empty_street_locator_segment(stripped):
        return stripped
    return text


def _normalize_street_range_dashes(value: str) -> str:
    text = str(value or "")
    text = re.sub(r"(?i)\bСалтыкова\s*[-–—]\s*Щедрина\b", "Салтыкова-Щедрина", text)
    text = re.sub(r"\s+[-–—]\s+", " – ", text)
    return text


def _strip_trailing_locality_suffix_from_street(value: str) -> str:
    text = str(value or "").strip()
    if not re.match(
        r"(?i)^\s*(?:ул\.?|улица|пер\.?|переулок|проспект|пр-т\.?|пр-кт\.?|просп\.?|пр-д\.?|проезд|"
        r"б-р\.?|бульвар|наб\.?|набережная|пл\.?|площадь)\s+",
        text,
    ):
        return text
    return re.sub(
        r"(?i)\s+(?:г\.?|город)\s+[А-ЯЁA-Zа-яёa-z0-9 .'-]+$",
        "",
        text,
    ).strip(" ,")


def _looks_like_empty_street_locator_segment(value: Any) -> bool:
    return bool(
        re.fullmatch(
            r"(?i)\s*(?:ул\.?|улица|пер\.?|переулок|проспект|пр-т\.?|пр-кт\.?|просп\.?|пр-д\.?|проезд|"
            r"б-р\.?|бульвар|наб\.?|набережная|пл\.?|площадь)\s*",
            str(value or "").strip(" ,"),
        )
    )


def _normalize_house_slot(value: Any) -> str | None:
    text = _clean(value)
    if not text:
        return None
    text = _strip_premise_tail(text)
    text = re.sub(r"(?i)(\d)\s*-\s*([а-яa-z])\b", r"\1\2", text)
    if re.match(r"^\d{1,4}[а-яa-z]?(?:\b|[\s,/\\-])", text, flags=re.IGNORECASE):
        text = f"д. {text}"
    text = re.sub(r"^(?:дом|д)\.?\s*", "д. ", text, flags=re.IGNORECASE)
    text = re.sub(r"^(?:здание|зд)\.?\s*", "зд. ", text, flags=re.IGNORECASE)
    text = re.sub(r"^(?:строение|стр)\.?\s*", "стр. ", text, flags=re.IGNORECASE)
    text = re.sub(r"\bкорпус\s*", "к. ", text, flags=re.IGNORECASE)
    text = re.sub(r"\bстроение\s*", "стр. ", text, flags=re.IGNORECASE)
    text = re.sub(r"\s+", " ", text)
    return text.strip(" ,")


def _strip_premise_tail(value: Any) -> str:
    return re.sub(
        r"(?i)\s*[,;]?\s*(?:помещ\.?|помещение|пом\.?|офис|оф\.?)\s*\d+[а-яa-z]?\b.*$",
        "",
        str(value or ""),
    ).strip(" ,")


def _normalize_municipality_slot(value: Any) -> str | None:
    text = _clean(value)
    if not text:
        return None
    text = re.sub(r"\s+", " ", text).strip(" ,;")
    text = re.sub(r"^(?:г\.?|город)\s+", "город ", text, flags=re.IGNORECASE)
    text = re.sub(r"\bмуниципальное\s+образование\s+", "", text, flags=re.IGNORECASE)
    text = re.sub(r"\bмуниципальный\s+район\b", "муниципальный район", text, flags=re.IGNORECASE)
    text = re.sub(r"\bмуниципальный\s+округ\b", "муниципальный округ", text, flags=re.IGNORECASE)
    text = re.sub(r"\bгородской\s+округ\b", "городской округ", text, flags=re.IGNORECASE)
    text = re.sub(r"\s+", " ", text).strip(" ,;")
    return text or None


def _normalize_locality_slot(value: Any) -> str | None:
    text = _clean(value)
    if not text:
        return None
    text = re.sub(r"\s+", " ", text).strip(" ,;")
    replacements = (
        (r"^город\s+", "г. "),
        (r"^г\.?\s+", "г. "),
        (r"^село\s+", "с. "),
        (r"^с\.?\s+", "с. "),
        (r"^деревня\s+", "д. "),
        (r"^д\.?\s+", "д. "),
        (r"^пос[её]лок\s+городского\s+типа\s+", "пгт "),
        (r"^пос[её]лок\s+", "п. "),
        (r"^пос\.?\s+", "п. "),
        (r"^п\.?\s+", "п. "),
        (r"^рабочий\s+пос[её]лок\s+", "р. п. "),
        (r"^р\.?\s*п\.?\s+", "р. п. "),
    )
    for pattern, replacement in replacements:
        text = re.sub(pattern, replacement, text, flags=re.IGNORECASE)
    text = re.sub(r"\s+", " ", text).strip(" ,;")
    return text or None


def _normalize_microdistrict_slot(value: Any) -> str | None:
    text = _clean(value)
    if not text:
        return None
    text = re.sub(r"\s+", " ", text).strip(" ,;")
    text = re.sub(r"^микрорайон\s+", "мкр. ", text, flags=re.IGNORECASE)
    text = re.sub(r"^мкр\.?\s+", "мкр. ", text, flags=re.IGNORECASE)
    text = re.sub(r"^квартал\s+", "квартал ", text, flags=re.IGNORECASE)
    text = re.sub(r"\s+", " ", text).strip(" ,;")
    return text or None


def _candidate_quality_score(
    row: dict[str, Any],
    record: NormalizedAddress,
    candidate: SourceCandidate,
    source_kind: str,
    *,
    oktmo_scope_match: dict[str, Any] | None = None,
) -> float:
    score = 0.0
    slots = record.slot_dict()
    score += 35.0 if record.house_base else 0.0
    score += 25.0 if record.street_words or slots.get("territory") else 0.0
    score += 14.0 if record.oktmo_code else 0.0
    score += 8.0 if record.postal_index else 0.0
    score += min(len(str(record.normalized or "")), 140) * 0.05
    score += sum(1 for value in slots.values() if value) * 0.8
    if source_kind == "existing_reference":
        score += 12.0
    if "reference" in candidate.source_kind or "address" in candidate.source_kind:
        score += 6.0
    raw_len = len(str(candidate.value or ""))
    if raw_len > 450:
        score -= min((raw_len - 450) / 20.0, 20.0)
    if _looks_like_contaminated_table_address(candidate.value):
        score -= 35.0
    if row.get("Normalized_Address") and _norm(row.get("Normalized_Address")) == _norm(record.raw):
        score += 2.0
    if oktmo_scope_match:
        score += float(oktmo_scope_match.get("score_bonus") or 0.0)
    return score


def _oktmo_scope_candidate_match(
    row: dict[str, Any],
    record: NormalizedAddress,
    candidate: SourceCandidate,
    key_profile: dict[str, Any] | None,
) -> dict[str, Any]:
    weighted_keys = list((key_profile or {}).get("weighted_extra_keys") or [])
    if not weighted_keys:
        return {}

    searchable = _oktmo_scope_search_text(row, record, candidate)
    if not searchable:
        return {}
    has_context = _oktmo_scope_has_address_or_object_context(row, record, candidate)
    record_code = re.sub(r"\D+", "", str(record.oktmo_code or row.get("OKTMO_Code") or ""))
    matches: list[dict[str, Any]] = []
    seen: set[tuple[str, str, str]] = set()
    for item in weighted_keys:
        if str(item.get("source") or "") == "subject":
            continue
        weight = float(item.get("weight") or 0.0)
        if weight <= 0:
            continue
        value = str(item.get("value") or "").strip()
        code = re.sub(r"\D+", "", str(item.get("code") or ""))
        matched_by_code = bool(code and record_code and code == record_code)
        if not matched_by_code and not _oktmo_scope_term_matches(searchable, value):
            continue
        context_required = bool(item.get("context_required"))
        valid = matched_by_code or not context_required or has_context
        key = (_norm(value), str(item.get("role") or ""), code)
        if key in seen:
            continue
        seen.add(key)
        matches.append(
            {
                "value": value,
                "weight": weight,
                "role": str(item.get("role") or ""),
                "source": str(item.get("source") or ""),
                "context_required": context_required,
                "valid": valid,
                "code": code,
                "matched_by_code": matched_by_code,
            }
        )
    if not matches:
        return {}

    valid_matches = [item for item in matches if item["valid"]]
    ranked = sorted(valid_matches or matches, key=lambda item: (float(item["weight"]), len(str(item["value"]))), reverse=True)
    best = ranked[0]
    valid = bool(valid_matches)
    score_bonus = min(float(best["weight"]) / 10.0, 12.0) if valid else -4.0
    return {
        "valid": valid,
        "has_context": has_context,
        "best": best,
        "matches": matches,
        "score_bonus": score_bonus,
    }


def _oktmo_scope_candidate_diagnostics(scope_match: dict[str, Any]) -> dict[str, Any]:
    best = dict(scope_match.get("best") or {})
    matches = list(scope_match.get("matches") or [])
    values: list[str] = []
    for item in sorted(matches, key=lambda row: (bool(row.get("valid")), float(row.get("weight") or 0), len(str(row.get("value") or ""))), reverse=True):
        suffix = "" if item.get("valid") else "?"
        values.append(f"{item.get('value')}{suffix}:{item.get('role')}")
        if len(values) >= 8:
            break
    return {
        "OKTMO_Scope_Match": "; ".join(values),
        "OKTMO_Scope_Code": best.get("code"),
        "OKTMO_Scope_Source": best.get("source"),
        "OKTMO_Scope_Weight": best.get("weight"),
        "OKTMO_Scope_Valid": bool(scope_match.get("valid")),
        "OKTMO_Scope_Context": bool(scope_match.get("has_context")),
    }


def _oktmo_scope_match_summary(rows: list[_CandidateRow]) -> dict[str, Any]:
    counts: Counter[str] = Counter()
    valid = 0
    invalid = 0
    for item in rows:
        role = str(item.row.get("OKTMO_Scope_Match") or "").strip()
        if not role:
            continue
        if bool(item.row.get("OKTMO_Scope_Valid")):
            valid += 1
        else:
            invalid += 1
        first_role = role.split(":", 1)[-1].split(";", 1)[0].strip() or "unknown"
        counts[first_role] += 1
    return {
        "valid": valid,
        "invalid": invalid,
        "by_role": dict(counts),
    }


def _oktmo_scope_search_text(row: dict[str, Any], record: NormalizedAddress, candidate: SourceCandidate) -> str:
    values = [
        candidate.value,
        record.raw,
        record.normalized,
        row.get("Normalized_Address"),
        row.get("Municipality"),
        row.get("Locality"),
        row.get("Territory"),
        row.get("Microdistrict"),
        row.get("Street"),
        row.get("OKTMO_Name"),
    ]
    return _norm(" | ".join(str(value or "") for value in values if _clean(value)))


def _oktmo_scope_term_matches(normalized_text: str, term: str) -> bool:
    key = _norm(term)
    if not key:
        return False
    return re.search(rf"(?<![0-9a-zа-я]){re.escape(key)}(?![0-9a-zа-я])", normalized_text) is not None


def _oktmo_scope_has_address_or_object_context(
    row: dict[str, Any],
    record: NormalizedAddress,
    candidate: SourceCandidate,
) -> bool:
    if record.house_base:
        return True
    if any(_clean(row.get(column)) for column in ("Street", "House", "Territory", "Microdistrict", "Premise")):
        return True
    text = " | ".join(
        str(value or "")
        for value in (candidate.value, record.raw, record.normalized, row.get("Normalized_Address"))
        if _clean(value)
    )
    return bool(
        STREET_MARKER_RE.search(text)
        or HOUSE_MARKER_RE.search(text)
        or _looks_like_microdistrict_text(text)
        or _contains_org_context(text)
    )


def _assembly_key(row: _CandidateRow, *, global_context: dict[str, Any] | None = None) -> str:
    slot_core = _slot_address_core_key(row)
    if slot_core:
        scope = _address_scope_key(row, global_context=global_context)
        return f"core|{scope}|{slot_core}" if scope else f"core|{slot_core}"
    if row.record.street_words and row.record.house_base:
        numbers = ",".join(sorted(row.record.street_numbers))
        core = "|".join(
            part
            for part in (
                row.record.street_words,
                numbers,
                row.record.house_base,
                row.record.house_mods,
            )
            if part
        )
        scope = _address_scope_key(row, global_context=global_context)
        return f"core|{scope}|{core}" if scope else f"core|{core}"
    if row.record.match_key:
        return row.record.match_key
    value = row.row.get("Match_Key") or row.row.get("Normalized_Address") or row.row.get("Raw_Address")
    return _norm(value)


def _slot_address_core_key(row: _CandidateRow) -> str:
    street_text = _clean(row.row.get("Street"))
    house_text = _clean(row.row.get("House"))
    if not street_text or not house_text:
        return ""
    _territory, street_words, street_numbers, _street_house_base, _street_house_mods = parse_address_components(street_text)
    _house_territory, _house_street_words, _house_street_numbers, house_base, parsed_house_mods = parse_address_components(house_text)
    if not street_words or not house_base:
        return ""
    house_mods = parsed_house_mods or _safe_house_mods(row.row.get("House_Mods"))
    house_kind = _house_kind_key(house_text)
    premise = _norm(row.row.get("Premise"))
    return "|".join(
        part
        for part in (
            street_words,
            ",".join(sorted(street_numbers)),
            house_kind,
            house_base,
            house_mods,
            premise,
        )
        if part
    )


def _house_kind_key(value: Any) -> str:
    text = str(value or "").strip().lower().replace("ё", "е")
    match = re.match(r"^(д|дом|зд|здание|стр|строение)\.?", text)
    if not match:
        return ""
    kind = match.group(1)
    if kind.startswith("зд") or kind == "здание":
        return "зд"
    if kind.startswith("стр") or kind == "строение":
        return "стр"
    return "д"


def _safe_house_mods(value: Any) -> str:
    text = _norm(value)
    if text in {"г", "город", "с", "село", "д", "деревня", "п", "пос", "поселок", "посёлок"}:
        return ""
    return text


def _address_scope_key(row: _CandidateRow, *, global_context: dict[str, Any] | None = None) -> str:
    if global_context and not _row_scope_conflicts_with_context(row.row, global_context):
        context_oktmo = re.sub(r"\D+", "", str(global_context.get("OKTMO_Code") or ""))
        if context_oktmo:
            return f"oktmo:{context_oktmo}"
        context_scope = _norm(" | ".join(str(global_context.get(column) or "") for column in ("Municipality", "Locality")))
        if context_scope:
            return f"context:{context_scope}"
    digits = re.sub(r"\D+", "", str(row.row.get("OKTMO_Code") or row.record.oktmo_code or ""))
    if digits:
        return f"oktmo:{digits}"
    row_scope = _norm(" | ".join(str(row.row.get(column) or "") for column in ("Municipality", "Locality", "Territory")))
    if row_scope:
        return f"row:{row_scope}"
    return _norm(row.record.territory_key)


def _row_scope_conflicts_with_context(row: dict[str, Any], global_context: dict[str, Any]) -> bool:
    context_oktmo = re.sub(r"\D+", "", str(global_context.get("OKTMO_Code") or ""))
    row_oktmo = re.sub(r"\D+", "", str(row.get("OKTMO_Code") or ""))
    if context_oktmo and row_oktmo and row_oktmo != context_oktmo:
        return not (row_oktmo.startswith(context_oktmo) or context_oktmo.startswith(row_oktmo))
    for column in ("Municipality", "Locality"):
        row_value = _clean(row.get(column))
        context_value = _clean(global_context.get(column))
        if not row_value or not context_value:
            continue
        if _geo_core_key(row_value) != _geo_core_key(context_value):
            return True
    return False


def _geo_core_key(value: Any) -> str:
    text = _norm(value)
    text = re.sub(
        r"\b(?:муниципальное образование|городской округ|муниципальный округ|муниципальный район|"
        r"городское поселение|сельское поселение|город|село|деревня|поселок|поселок городского типа|"
        r"рабочий поселок|район|округ|пгт|рп|г|с|д|п)\b",
        " ",
        text,
    )
    return re.sub(r"\s+", " ", text).strip()


def _construct_address_from_row(
    row: dict[str, Any],
    enabled_slots: tuple[str, ...],
    slot_order: tuple[str, ...],
    *,
    omit_microdistrict_when_street_found: bool,
) -> str | None:
    values: list[str] = []
    enabled = set(enabled_slots)
    ordered = [slot for slot in slot_order if slot in enabled and slot in SLOT_COLUMN_BY_ID]
    ordered.extend(slot for slot in DEFAULT_SLOT_ORDER if slot in enabled and slot not in ordered)
    for slot in ordered:
        if slot == "microdistrict" and omit_microdistrict_when_street_found and row.get("Street"):
            continue
        column = SLOT_COLUMN_BY_ID.get(slot)
        if column:
            _append_unique(values, row.get(column))
    return ", ".join(values) or None


def _has_collectable_address_core(record: NormalizedAddress) -> bool:
    if not record.has_address_core:
        return False
    raw = str(record.raw or "")
    if len(raw) > 1200:
        return False
    if not (record.house_base or record.parts.street or record.parts.territory or record.parts.locality or record.oktmo_code):
        return False
    return True


def _plausible_address_text(value: Any) -> bool:
    text = _clean(_normalize_spaced_address_marker_dots(value))
    if not text or len(text) < 6 or len(text) > 2500:
        return False
    lowered = text.lower()
    if lowered in {"адрес", "местоположение", "адрес объекта", "местоположение объекта"}:
        return False
    if re.fullmatch(r"[\d\s.,;:()/\\-]+", text):
        return False
    return bool(ADDRESS_HINT_RE.search(text) or looks_like_address_text(text) or re.search(r"\b\d{6}\b", text))


def _drop_reason(row: dict[str, Any]) -> str | None:
    address = _clean(row.get("Normalized_Address")) or ""
    raw = _clean(row.get("Raw_Address")) or address
    if not address:
        return "empty_address"
    if len(address) > 700:
        return "address_too_long"
    if _looks_like_contaminated_table_address(raw) or _looks_like_contaminated_table_address(address):
        return "contaminated_table_dump"
    if _is_subject_only_address(address):
        return "subject_only_address"
    street = _clean(row.get("Street"))
    territory = _clean(row.get("Territory"))
    microdistrict = _clean(row.get("Microdistrict"))
    house = _clean(row.get("House"))
    house_base = _clean(row.get("House_Base")) or _house_base_from_house(house)
    if _is_too_general_location(address, street=street, territory=territory, microdistrict=microdistrict, house_base=house_base):
        return "too_general_location"
    if not house_base:
        return "street_without_house" if street or territory or microdistrict else "no_house"
    if _is_invalid_house_base(house_base):
        return "invalid_house_number"
    if not (street or territory or microdistrict):
        return "house_without_street_or_territory"
    scope_reason = _oktmo_scope_conflict_drop_reason(row)
    if scope_reason:
        return scope_reason
    strict_reason = _strict_address_slot_drop_reason(row, raw=raw)
    if strict_reason:
        return strict_reason
    return None


def _oktmo_scope_conflict_drop_reason(row: dict[str, Any]) -> str | None:
    if not bool(row.get("OKTMO_Scope_Valid")):
        return None
    scope_code = re.sub(r"\D+", "", str(row.get("OKTMO_Scope_Code") or ""))
    row_code = re.sub(r"\D+", "", str(row.get("OKTMO_Code") or ""))
    if not scope_code or not row_code:
        return None
    if _oktmo_codes_are_compatible(scope_code, row_code):
        return None
    return "oktmo_scope_conflict"


def _oktmo_codes_are_compatible(left: str, right: str) -> bool:
    if left == right:
        return True
    left8 = left[:8] if len(left) >= 8 else left
    right8 = right[:8] if len(right) >= 8 else right
    return bool(left8 and right8 and left8 == right8)


def _is_too_general_location(
    address: str,
    *,
    street: str | None,
    territory: str | None,
    microdistrict: str | None,
    house_base: str | None,
) -> bool:
    if street or territory or microdistrict or house_base:
        return False
    text = _clean(address) or ""
    if not text:
        return False
    if STREET_MARKER_RE.search(text) or _looks_like_microdistrict_text(text) or HOUSE_MARKER_RE.search(text):
        return False
    key = _norm(text)
    if not key:
        return False
    general_markers = (
        "федеральный округ",
        "область",
        "край",
        "республика",
        "автономный округ",
        "муниципальный округ",
        "муниципальный район",
        "городской округ",
        "город ",
        "г ",
    )
    return any(marker in key for marker in general_markers)


def _strict_address_slot_drop_reason(row: dict[str, Any], *, raw: str) -> str | None:
    street = _clean(row.get("Street"))
    house = _clean(row.get("House"))
    address = _clean(row.get("Normalized_Address")) or ""
    if _source_context_is_route(row) and (
        _contains_org_context(raw)
        or _contains_org_context(street)
        or _looks_like_route_fragment(raw)
        or _looks_like_route_fragment(street)
    ):
        return "route_source_fragment"
    if _contains_org_context(address):
        return "object_name_in_address"
    if _looks_like_route_fragment(address):
        return "route_fragment_in_address"
    for column in ("Municipality", "Locality", "Territory", "Microdistrict"):
        reason = _strict_geo_slot_drop_reason(column, row.get(column))
        if reason:
            return reason
    if street:
        street_reason = _strict_street_slot_drop_reason(street)
        if street_reason:
            return street_reason
    if house and not _looks_like_ideal_house_slot(house):
        return "invalid_house_slot"
    if _has_unbalanced_brackets(address) or _has_unbalanced_brackets(street):
        return "unbalanced_address_fragment"
    return None


def _strict_geo_slot_drop_reason(column: str, value: Any) -> str | None:
    text = _clean(value)
    if not text:
        return None
    if _contains_org_context(text):
        return f"object_name_in_{column.lower()}_slot"
    if _looks_like_route_fragment(text):
        return f"route_fragment_in_{column.lower()}_slot"
    if _has_unbalanced_brackets(text):
        return f"unbalanced_{column.lower()}_slot"
    if column == "Municipality" and not _looks_like_ideal_municipality_slot(text):
        return "invalid_municipality_slot"
    if column == "Locality" and not _looks_like_ideal_locality_slot(text):
        return "invalid_locality_slot"
    return None


def _strict_street_slot_drop_reason(value: Any) -> str | None:
    street = _clean(value)
    if not street:
        return None
    if _contains_org_context(street):
        return "object_name_in_street_slot"
    if _looks_like_route_fragment(street):
        return "route_fragment_in_street_slot"
    if _has_unbalanced_brackets(street):
        return "unbalanced_street_fragment"
    if not (_street_slot_has_explicit_locator(street) or _looks_like_plain_street_name(street)):
        return "invalid_street_slot"
    return None


def _source_context_is_route(row: dict[str, Any]) -> bool:
    text = _norm(" | ".join(str(row.get(column) or "") for column in ("Sources", "First_Source", "Raw_Examples")))
    return any(marker in text for marker in ROUTE_CONTEXT_MARKERS)


def _contains_org_context(value: Any) -> bool:
    text = _norm(value)
    if not text:
        return False
    return any(re.search(rf"\b{re.escape(marker)}\b", text) for marker in ORG_CONTEXT_MARKERS)


def _looks_like_route_fragment(value: Any) -> bool:
    text = _norm(value)
    raw = str(value or "")
    if not text:
        return False
    if any(marker in text for marker in ROUTE_CONTEXT_MARKERS):
        return True
    if re.search(r"\s[-–—]\s", raw) and (_contains_org_context(raw) or re.search(r"\b(?:с/о|дк|жк)\b", raw, flags=re.IGNORECASE)):
        return True
    if re.search(r"^\s*\d+\.\s+", raw) and _contains_org_context(raw):
        return True
    return False


def _street_slot_has_explicit_locator(value: Any) -> bool:
    text = _clean(value) or ""
    if not text:
        return False
    if STREET_PREFIX_RE.search(text):
        return True
    if STREET_SUFFIX_RE.search(text) and not _contains_org_context(text):
        return True
    return False


def _looks_like_ideal_municipality_slot(value: Any) -> bool:
    text = _clean(value) or ""
    key = _norm(text)
    if not text or not key:
        return False
    if STREET_MARKER_RE.search(text) or _looks_like_microdistrict_text(text):
        return False
    return bool(
        "муниципаль" in key
        or "городской округ" in key
        or re.search(r"\b(?:район|округ)\b", key)
        or re.match(r"^(?:город|г)\s+[а-яa-z0-9 .'-]+$", key)
    )


def _looks_like_ideal_locality_slot(value: Any) -> bool:
    text = _clean(value) or ""
    if not text:
        return False
    if STREET_MARKER_RE.search(text) or _looks_like_microdistrict_text(text):
        return False
    return bool(
        re.match(
            r"(?i)^\s*(?:г\.?|город|с\.?|село|д\.?|деревня|п\.?|пос\.?|пос[её]лок|пгт|рп)\s+"
            r"[а-яёa-z0-9 .'-]+\s*$",
            text,
        )
    )


def _looks_like_microdistrict_text(value: Any) -> bool:
    return bool(re.search(r"\b(?:мкр\.?|микрорайон|квартал|кв-л\.?)\b", str(value or ""), flags=re.IGNORECASE))


def _looks_like_plain_street_name(value: Any) -> bool:
    text = _clean(value) or ""
    key = _norm(text)
    if not text or not key:
        return False
    if _contains_org_context(text) or _looks_like_route_fragment(text) or _looks_like_admin_context(text):
        return False
    if re.search(r"[()|;:]", text):
        return False
    words = [part for part in key.split() if part]
    if not 1 <= len(words) <= 4:
        return False
    if all(part.isdigit() for part in words):
        return False
    return bool(re.search(r"[а-яa-z]", key))


def _looks_like_ideal_house_slot(value: Any) -> bool:
    text = _clean(value) or ""
    if not text:
        return False
    if _has_unbalanced_brackets(text) or "|" in text:
        return False
    return bool(
        re.fullmatch(
            r"(?i)(?:д|зд|стр)\.\s*\d{1,4}[а-яa-z]?"
            r"(?:\s*,\s*(?:(?:к|корп|стр|с|лит)\.?\s*\d{0,4}[а-яa-z]?))*",
            text,
        )
    )


def _has_unbalanced_brackets(value: Any) -> bool:
    text = str(value or "")
    return text.count("(") != text.count(")") or text.count("«") != text.count("»")


def _looks_like_contaminated_table_address(value: Any) -> bool:
    raw = str(value or "")
    text = _norm(raw)
    if not text:
        return False
    if "|" not in raw and "column_" not in raw.lower():
        return False
    markers = (
        "планируемый срок",
        "примечение",
        "примечание",
        "площадь учебная",
        "количество обучающихся",
        "требуемая учебная площадь",
        "дополнительная потребность",
        "кадастровый номер",
        "column",
    )
    return any(marker in text for marker in markers)


def _is_subject_only_address(value: Any) -> bool:
    text = _norm(value)
    return text in {
        "тюменская область",
        "амурская область",
        "иркутская область",
        "кемеровская область",
        "красноярский край",
        "приморский край",
    }


def _house_base_from_house(value: Any) -> str | None:
    match = re.search(r"\b(?:д\.?|дом)?\s*(\d+[а-яa-z]?)\b", str(value or ""), flags=re.IGNORECASE)
    return match.group(1).lower() if match else None


def _is_invalid_house_base(value: Any) -> bool:
    text = str(value or "").strip().lower()
    return bool(re.fullmatch(r"0+[а-яa-z]?", text, flags=re.IGNORECASE))


def _sort_key(row: dict[str, Any], enabled_slots: tuple[str, ...], slot_order: tuple[str, ...]) -> tuple[Any, ...]:
    enabled = set(enabled_slots)
    values: list[Any] = []
    for slot in slot_order:
        if slot not in enabled:
            continue
        column = SLOT_COLUMN_BY_ID.get(slot)
        values.append(_natural_key(row.get(column) if column else None))
    values.append(_natural_key(row.get("Normalized_Address")))
    return tuple(values)


def _natural_key(value: Any) -> tuple[Any, ...]:
    text = _norm(value)
    if not text:
        return (1, "")
    parts: list[Any] = [0]
    for part in re.split(r"(\d+)", text):
        if not part:
            continue
        parts.append((0, int(part)) if part.isdigit() else (1, part))
    return tuple(parts)


def _project_diagnostic_columns(row: dict[str, Any]) -> dict[str, Any]:
    base = {
        "Drop_Reason": row.get("Drop_Reason"),
        "Disposition": row.get("Disposition"),
        "Review_Action": row.get("Review_Action"),
        "Normalized_Address": row.get("Normalized_Address"),
        "Raw_Address": row.get("Raw_Address"),
    }
    for column in _ordered_slot_columns(DEFAULT_ENABLED_SLOTS, DEFAULT_SLOT_ORDER):
        if column in row:
            base[column] = row.get(column)
    for column in (
        "Candidate_Count",
        "Source_Count",
        "Sources",
        "First_Source",
        "Best_Slot_Score",
        "OKTMO_Scope_Match",
        "OKTMO_Scope_Code",
        "OKTMO_Scope_Source",
        "OKTMO_Scope_Weight",
        "OKTMO_Scope_Valid",
        "OKTMO_Scope_Context",
    ):
        base[column] = row.get(column)
    return base


def _write_staging_workbook(
    staging_file: Path,
    staged_rows: list[dict[str, Any]],
    normalized_rows: list[dict[str, Any]],
    rejected_rows: list[dict[str, Any]],
    conflict_rows: list[dict[str, Any]] | None = None,
) -> None:
    staging_file.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    try:
        sheet = workbook.active
        sheet.title = "Slot_Candidates"
        _write_dict_sheet(sheet, staged_rows)
        sanitized_sheet = workbook.create_sheet("Sanitized")
        _write_dict_sheet(sanitized_sheet, normalized_rows)
        if rejected_rows:
            rejected_sheet = workbook.create_sheet("Rejected")
            _write_dict_sheet(rejected_sheet, rejected_rows)
        if conflict_rows:
            conflict_sheet = workbook.create_sheet("OKTMO_Conflicts")
            _write_dict_sheet(conflict_sheet, conflict_rows)
        workbook.save(staging_file)
    finally:
        workbook.close()


def _write_address_output_workbook(
    output_file: Path,
    normalized_rows: list[dict[str, Any]],
    *,
    target_file: Path | None,
    target_column: str | None,
    clean_output_workbook: bool = False,
) -> dict[str, Any]:
    if clean_output_workbook:
        workbook = Workbook()
        workbook.active.title = "Addresses"
    elif target_file and target_file.exists():
        workbook = load_workbook(target_file)
    else:
        workbook = Workbook()
    try:
        sheet = workbook.worksheets[0]
        if clean_output_workbook:
            column_index = 1
            sheet.cell(row=1, column=column_index).value = _clean_output_address_header(target_column)
            start_row = 2
        else:
            column_index, header_row = _resolve_target_column(sheet, target_column or "A")
            start_row = _append_start_row(sheet, column_index, header_row)
        for offset, row in enumerate(normalized_rows):
            sheet.cell(row=start_row + offset, column=column_index).value = row.get("Normalized_Address")
        if clean_output_workbook:
            _style_clean_output_sheet(sheet)
        output_file.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(output_file)
        return {
            "appended_rows": len(normalized_rows),
            "written_rows": len(normalized_rows),
            "sheet_name": sheet.title,
            "column_letter": excel_column_letter(column_index - 1),
            "start_row": start_row,
        }
    finally:
        workbook.close()


def _clean_output_address_header(target_column: str | None) -> str:
    return "Normalized_Address"


def _style_clean_output_sheet(sheet: Any) -> None:
    header_fill = PatternFill(fill_type="solid", fgColor="1F2933")
    header_font = Font(name="Segoe UI", size=9, bold=True, color="FFFFFF")
    body_font = Font(name="Segoe UI", size=9)
    alignment = Alignment(vertical="top", wrap_text=True)
    max_row = sheet.max_row or 1
    max_column = sheet.max_column or 1
    for column_index in range(1, max_column + 1):
        header_cell = sheet.cell(row=1, column=column_index)
        header_cell.font = header_font
        header_cell.fill = header_fill
        header_cell.alignment = alignment
        max_len = len(str(header_cell.value or ""))
        for row_index in range(2, max_row + 1):
            cell = sheet.cell(row=row_index, column=column_index)
            cell.font = body_font
            cell.alignment = alignment
            max_len = max(max_len, len(str(cell.value or "")))
        letter = excel_column_letter(column_index - 1)
        sheet.column_dimensions[letter].width = max(18, min(80, max_len + 2))
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions


def _discover_existing_address_workbook(folder: Path) -> Path | None:
    candidates: list[Path] = []
    if folder.exists():
        candidates.extend(
            path
            for pattern in ("*.xlsx",)
            for path in folder.glob(pattern)
            if path.is_file() and not path.name.startswith("~$")
        )
    if not candidates:
        return None
    return max(candidates, key=lambda path: path.stat().st_mtime)


def _resolve_target_column(sheet: Any, token: str) -> tuple[int, int | None]:
    text = str(token or "").strip()
    if not text:
        return 1, None
    if text.isdigit():
        return max(1, int(text)), None
    if re.fullmatch(r"[A-Za-z]{1,3}", text):
        return excel_column_index(text) + 1, None
    header = _find_header_column(sheet, text)
    if header:
        return header
    if sheet.max_row <= 1 and sheet.max_column <= 1 and not _clean(sheet.cell(row=1, column=1).value):
        sheet.cell(row=1, column=1).value = text
        return 1, 1
    raise ValueError(f"Target column {token!r} was not found by letter, number, or header name.")


def _find_header_column(sheet: Any, token: str) -> tuple[int, int] | None:
    needle = _norm(token)
    if not needle:
        return None
    max_rows = min(sheet.max_row or 0, 20)
    max_cols = min(sheet.max_column or 0, 200)
    for row_index in range(1, max_rows + 1):
        for column_index in range(1, max_cols + 1):
            key = _norm(sheet.cell(row=row_index, column=column_index).value)
            if key and (key == needle or needle in key or key in needle):
                return column_index, row_index
    return None


def _append_start_row(sheet: Any, column_index: int, header_row: int | None) -> int:
    last = 0
    for row_index in range(sheet.max_row or 1, 0, -1):
        if _clean(sheet.cell(row=row_index, column=column_index).value):
            last = row_index
            break
    if header_row is not None:
        return max(last + 1, header_row + 1)
    return max(1, last + 1)


def _write_dict_sheet(sheet: Any, rows: list[dict[str, Any]]) -> None:
    if not rows:
        return
    headers: list[str] = []
    for row in rows:
        for key in row:
            if key not in headers:
                headers.append(key)
    for column_index, header in enumerate(headers, start=1):
        sheet.cell(row=1, column=column_index).value = header
    for row_index, row in enumerate(rows, start=2):
        for column_index, header in enumerate(headers, start=1):
            sheet.cell(row=row_index, column=column_index).value = _excel_cell_value(row.get(header))


def _excel_cell_value(value: Any) -> Any:
    if isinstance(value, (list, tuple, set, frozenset)):
        value = " | ".join(str(item) for item in value if item not in (None, ""))
    if isinstance(value, str) and len(value) > 32000:
        return value[:31980] + " ...[truncated]"
    return value


def _unique_sheet_name(workbook: Any, base: str) -> str:
    if base not in workbook.sheetnames:
        return base
    for index in range(2, 1000):
        candidate = f"{base}_{index}"
        if candidate not in workbook.sheetnames:
            return candidate
    raise RuntimeError(f"Could not allocate a unique sheet name for {base}.")


def _merge_text_examples(values: Any, *, limit: int = 20, separator: str = " | ") -> str:
    result: str | None = None
    for value in values:
        result = _append_unique_examples(result, value, limit=limit, separator=separator)
    return result or ""


def _append_unique(values: list[str], value: Any) -> None:
    text = _clean(value)
    if not text:
        return
    key = _norm(text)
    if not key or any(_norm(existing) == key for existing in values):
        return
    values.append(text)


def _clean(value: Any) -> str | None:
    if value is None:
        return None
    text = re.sub(r"\s+", " ", str(value).replace("\xa0", " ")).strip(" ,")
    if not text or text.lower() == "nan":
        return None
    return text


def _norm(value: Any) -> str:
    text = str(value or "").lower().replace("ё", "е")
    text = re.sub(r"[^0-9a-zа-я]+", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def _log(log: Callable[[str], None] | None, message: str) -> None:
    if log:
        log(message)


def _cancelled(cancelled: Callable[[], bool] | None) -> bool:
    return bool(cancelled and cancelled())
