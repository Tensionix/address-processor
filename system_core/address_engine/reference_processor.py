from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Any
import json
import re

import pandas as pd
from openpyxl import load_workbook

from .address_diagnostics import diagnostic_summary
from .audion_address_core import looks_like_address_text
from .normalizer_runtime import normalizer_oktmo_key_profile, normalizer_runtime_report
from .reference_builder import DEFAULT_ENABLED_SLOTS, DEFAULT_SELECTED_SLOTS, DEFAULT_SLOT_ORDER, SLOT_COLUMN_BY_ID
from .slot_fill import build_slot_fill_summary
from .source_intake import detect_reference_column, excel_column_index, excel_column_letter
from .universal_address import AddressNormalizer, NormalizedAddress


@dataclass(frozen=True)
class ReferenceProcessingOptions:
    action: str = "sort"
    reference_columns: str | None = None
    enabled_slots: tuple[str, ...] = DEFAULT_SELECTED_SLOTS
    slot_order: tuple[str, ...] = DEFAULT_SLOT_ORDER
    remove_source_column: bool = True
    slot_header_mode: str = "full"


@dataclass(frozen=True)
class ReferenceProcessingResult:
    output_file: Path
    report_file: Path | None
    input_file: Path
    sheet_name: str
    reference_column: str
    data_rows: int
    sorted_rows: int
    deconstructed_rows: int
    slot_columns: tuple[str, ...]


@dataclass(frozen=True)
class _ReferenceColumn:
    index: int
    data_start_row: int
    header_row: int | None
    label: str
    source: str


@dataclass(frozen=True)
class _ProcessedRow:
    values: list[Any]
    record: NormalizedAddress | None
    sort_key: tuple[Any, ...]


def process_reference_workbook(
    *,
    input_file: Path,
    output_file: Path,
    normalizer: AddressNormalizer,
    options: ReferenceProcessingOptions,
    report_file: Path | None = None,
) -> ReferenceProcessingResult:
    action = str(options.action or "sort").strip().lower()
    if action not in {"sort", "deconstruct", "sort_deconstruct"}:
        raise ValueError(f"Unsupported reference processing action: {options.action!r}")

    frame = pd.read_excel(input_file, sheet_name=0, header=None, dtype=object)
    if frame.empty:
        raise RuntimeError(f"Workbook is empty: {input_file}")
    reference_column = _resolve_reference_column(options.reference_columns, frame)

    workbook = load_workbook(input_file)
    try:
        sheet = workbook.worksheets[0]
        sheet_name = sheet.title
        max_column = sheet.max_column
        data_start = max(1, reference_column.data_start_row)
        source_rows = [
            [sheet.cell(row=row_index, column=column_index).value for column_index in range(1, max_column + 1)]
            for row_index in range(data_start, sheet.max_row + 1)
        ]

        processed_rows = [
            _process_row(
                row_values,
                reference_index=reference_column.index,
                normalizer=normalizer,
                slot_order=options.slot_order,
                enabled_slots=options.enabled_slots,
                original_index=index,
            )
            for index, row_values in enumerate(source_rows)
        ]
        if action in {"sort", "sort_deconstruct"}:
            processed_rows.sort(key=lambda item: item.sort_key)

        for offset, processed in enumerate(processed_rows):
            row_index = data_start + offset
            for column_index, value in enumerate(processed.values, start=1):
                sheet.cell(row=row_index, column=column_index).value = value

        slot_ids: tuple[str, ...] = ()
        if action in {"deconstruct", "sort_deconstruct"}:
            slot_ids = _ordered_enabled_slots(options.enabled_slots, options.slot_order)
            _write_slot_columns(
                sheet,
                processed_rows,
                reference_column=reference_column,
                slot_ids=slot_ids,
                remove_source_column=options.remove_source_column,
                slot_header_mode=options.slot_header_mode,
            )

        output_file.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(output_file)
    finally:
        workbook.close()

    if report_file:
        oktmo_key_profile = normalizer_oktmo_key_profile(normalizer)
        report_file.parent.mkdir(parents=True, exist_ok=True)
        summary = {
            "mode": "address-reference-processing",
            "action": action,
            "input_file": str(input_file),
            "output_file": str(output_file),
            "sheet_name": sheet_name,
            "reference_column": excel_column_letter(reference_column.index),
            "data_rows": len(source_rows),
            "sorted_rows": len(source_rows) if action in {"sort", "sort_deconstruct"} else 0,
            "deconstructed_rows": len(source_rows) if action in {"deconstruct", "sort_deconstruct"} else 0,
            "slot_columns": [_slot_header(slot, options.slot_header_mode) for slot in slot_ids if slot in SLOT_COLUMN_BY_ID],
            "slot_fill": build_slot_fill_summary(
                processed_rows,
                enabled_slots=options.enabled_slots,
                slot_order=options.slot_order,
                slot_columns=SLOT_COLUMN_BY_ID,
                default_slot_order=DEFAULT_SLOT_ORDER,
                slot_labels={slot: _slot_header(slot, options.slot_header_mode) for slot in SLOT_COLUMN_BY_ID},
                get_slot_value=_processed_slot_value,
                get_normalized=lambda processed: processed.record.normalized if processed.record else None,
                is_empty_row=lambda processed: processed.record is None,
            ),
            "address_diagnostics": diagnostic_summary(processed.record for processed in processed_rows if processed.record),
            "normalizer": normalizer_runtime_report(normalizer, oktmo_key_profile),
            "options": {
                "reference_columns": options.reference_columns,
                "enabled_slots": list(options.enabled_slots),
                "slot_order": list(options.slot_order),
                "remove_source_column": bool(options.remove_source_column),
                "slot_header_mode": options.slot_header_mode,
            },
        }
        report_file.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")

    return ReferenceProcessingResult(
        output_file=output_file,
        report_file=report_file,
        input_file=input_file,
        sheet_name=sheet_name,
        reference_column=excel_column_letter(reference_column.index),
        data_rows=len(source_rows),
        sorted_rows=len(source_rows) if action in {"sort", "sort_deconstruct"} else 0,
        deconstructed_rows=len(source_rows) if action in {"deconstruct", "sort_deconstruct"} else 0,
        slot_columns=tuple(_slot_header(slot, options.slot_header_mode) for slot in slot_ids if slot in SLOT_COLUMN_BY_ID),
    )


def _resolve_reference_column(value: str | None, frame: pd.DataFrame) -> _ReferenceColumn:
    token = str(value or "").strip()
    if token:
        index, header_row, source = _resolve_reference_column_token(token, frame)
        data_start_row = (header_row + 2) if header_row is not None else _data_start_for_manual_column(frame, index)
        return _ReferenceColumn(
            index=index,
            data_start_row=data_start_row,
            header_row=(data_start_row - 1 if data_start_row > 1 else None),
            label=f"Sheet1!{excel_column_letter(index)}",
            source=source,
        )

    choice = detect_reference_column(frame)
    if not choice:
        raise RuntimeError("Could not auto-detect a reference address column.")
    return _ReferenceColumn(
        index=choice.column_index,
        data_start_row=choice.data_start_row + 1,
        header_row=choice.data_start_row if choice.data_start_row > 0 else None,
        label=choice.label,
        source=choice.source,
    )


def _resolve_reference_column_token(token: str, frame: pd.DataFrame) -> tuple[int, int | None, str]:
    text = str(token).strip()
    if text.isdigit():
        index = int(text) - 1
        return _validate_column_index(index, frame, token), None, "manual_number"
    if re.fullmatch(r"[A-Za-z]{1,3}", text):
        index = excel_column_index(text)
        if 0 <= index < len(frame.columns):
            return index, None, "manual_letter"
    header = _find_header_column(text, frame)
    if header:
        return header[0], header[1], "manual_header"
    raise ValueError(f"Reference column {token!r} was not found by letter, number, or header name.")


def _find_header_column(token: str, frame: pd.DataFrame) -> tuple[int, int] | None:
    needle = _header_key(token)
    if not needle:
        return None
    for row_index in range(min(20, len(frame))):
        for column_index in range(len(frame.columns)):
            key = _header_key(frame.iat[row_index, column_index])
            if key and (key == needle or needle in key or key in needle):
                return column_index, row_index
    return None


def _validate_column_index(index: int, frame: pd.DataFrame, token: object) -> int:
    if index < 0 or index >= len(frame.columns):
        max_col = excel_column_letter(len(frame.columns) - 1) if len(frame.columns) else "A"
        raise ValueError(f"Column {token!r} is outside workbook width A:{max_col}.")
    return index


def _data_start_for_manual_column(frame: pd.DataFrame, index: int) -> int:
    first_value = _clean(frame.iat[0, index]) if len(frame) else None
    return 2 if _looks_like_header(first_value) else 1


def _looks_like_header(value: str | None) -> bool:
    key = _header_key(value)
    return bool(
        key
        and any(
            marker in key
            for marker in (
                "address",
                "adres",
                "normalized",
                "reference",
                "groundtruth",
                "эталон",
                "адрес",
                "нормализ",
            )
        )
        and not looks_like_address_text(str(value or ""))
    )


def _process_row(
    row_values: list[Any],
    *,
    reference_index: int,
    normalizer: AddressNormalizer,
    slot_order: tuple[str, ...],
    enabled_slots: tuple[str, ...],
    original_index: int,
) -> _ProcessedRow:
    raw = row_values[reference_index] if reference_index < len(row_values) else None
    text = _clean(raw)
    record = normalizer.normalize(text) if text else None
    return _ProcessedRow(
        values=row_values,
        record=record,
        sort_key=_sort_key(record, slot_order, enabled_slots, original_index),
    )


def _sort_key(
    record: NormalizedAddress | None,
    slot_order: tuple[str, ...],
    enabled_slots: tuple[str, ...],
    original_index: int,
) -> tuple[Any, ...]:
    if not record:
        return (1, original_index)
    slots = record.slot_dict()
    values: list[Any] = [0]
    for slot in _ordered_sort_slots(enabled_slots, slot_order):
        values.append(_natural_key(_slot_value(slots, slot)))
    values.append(original_index)
    return tuple(values)


def _ordered_sort_slots(enabled_slots: tuple[str, ...], slot_order: tuple[str, ...]) -> tuple[str, ...]:
    enabled = set(enabled_slots)
    result: list[str] = []
    for slot in slot_order:
        if slot in DEFAULT_SLOT_ORDER and slot in enabled and slot not in result:
            result.append(slot)
    for slot in DEFAULT_SLOT_ORDER:
        if slot in enabled and slot not in result:
            result.append(slot)
    return tuple(result)


def _ordered_enabled_slots(enabled_slots: tuple[str, ...], slot_order: tuple[str, ...]) -> tuple[str, ...]:
    ordered = list(_ordered_sort_slots(enabled_slots, slot_order))
    ordered.extend(slot for slot in enabled_slots if slot not in ordered and slot in SLOT_COLUMN_BY_ID)
    return tuple(slot for slot in ordered if slot in SLOT_COLUMN_BY_ID)


def _write_slot_columns(
    sheet: Any,
    rows: list[_ProcessedRow],
    *,
    reference_column: _ReferenceColumn,
    slot_ids: tuple[str, ...],
    remove_source_column: bool,
    slot_header_mode: str,
) -> None:
    if not slot_ids:
        return
    insert_at = reference_column.index + 2
    data_start = max(1, reference_column.data_start_row)
    sheet.insert_cols(insert_at, amount=len(slot_ids))
    header_row = reference_column.header_row
    if header_row is None:
        sheet.insert_rows(data_start, amount=1)
        header_row = data_start
        data_start += 1
    for offset, slot in enumerate(slot_ids):
        sheet.cell(row=header_row, column=insert_at + offset).value = _slot_header(slot, slot_header_mode)
    for offset, processed in enumerate(rows):
        row_index = data_start + offset
        slots = processed.record.slot_dict() if processed.record else {}
        for slot_offset, slot in enumerate(slot_ids):
            sheet.cell(row=row_index, column=insert_at + slot_offset).value = _slot_value(slots, slot)
    if remove_source_column:
        sheet.delete_cols(reference_column.index + 1, amount=1)


def _processed_slot_value(processed: _ProcessedRow, slot: str) -> Any:
    slots = processed.record.slot_dict() if processed.record else {}
    return _slot_value(slots, slot)


def _slot_value(slots: dict[str, Any], slot: str) -> str | None:
    value = slots.get(slot)
    if isinstance(value, (list, tuple, set, frozenset)):
        return ", ".join(str(item) for item in value if str(item).strip())
    return _clean(value)


SHORT_SLOT_HEADERS = {
    "postal_index": "Index",
    "federal_district": "Fed_District",
    "parent_subject": "Parent",
    "subject": "Subject",
    "autonomous_okrug": "AO",
    "municipality": "Municipality",
    "locality": "Locality",
    "territory": "Territory",
    "microdistrict": "Mkr_Qtr",
    "street": "Street",
    "house": "House",
    "premise": "Premise",
    "oktmo_code": "OKTMO",
    "oktmo_name": "OKTMO_Name",
    "house_base": "House_No",
    "house_mods": "House_Mods",
    "house_modifiers": "House_Parts",
    "street_numbers": "Street_Nums",
    "territory_key": "Territory_Key",
    "match_key": "Match_Key",
}


def _slot_header(slot: str, mode: str) -> str:
    if str(mode or "").strip().lower() == "short":
        return SHORT_SLOT_HEADERS.get(slot, SLOT_COLUMN_BY_ID[slot])
    return SLOT_COLUMN_BY_ID[slot]


def _natural_key(value: Any) -> tuple[Any, ...]:
    text = _norm(value)
    if not text:
        return (1, "")
    parts: list[Any] = [0]
    for part in re.split(r"(\d+)", text):
        if not part:
            continue
        parts.append(int(part) if part.isdigit() else part)
    return tuple(parts)


def _header_key(value: Any) -> str:
    text = str(value or "").lower().replace("ё", "е")
    return re.sub(r"[^0-9a-zа-я]+", "", text)


def _norm(value: Any) -> str:
    text = str(value or "").lower().replace("ё", "е")
    text = re.sub(r"[^0-9a-zа-я]+", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def _clean(value: Any) -> str | None:
    if value is None:
        return None
    text = re.sub(r"\s+", " ", str(value).replace("\xa0", " ")).strip(" ,")
    if not text or text.lower() == "nan":
        return None
    return text
